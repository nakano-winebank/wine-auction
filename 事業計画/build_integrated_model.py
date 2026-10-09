# -*- coding: utf-8 -*-
"""ティエリーマルクス・ブラッスリー 事業収支 統合モデル
   席数・席効率・回転率・客単価 → 六本木単店PL → 5か年連結 → 資金まで一気通貫。"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule

OUT="/home/user/wine-auction/事業計画/ティエリーマルクス・ブラッスリー_事業収支統合モデル_202610.xlsx"
JP="Meiryo"; BLUE,BLACK,GREEN,RED="0000FF","000000","008000","C00000"
HDR=PatternFill("solid",fgColor="1F3864"); SUB=PatternFill("solid",fgColor="D9E2F3")
TOT=PatternFill("solid",fgColor="F2F2F2"); KEY=PatternFill("solid",fgColor="FFFF00")
INP=PatternFill("solid",fgColor="EAF1FB"); HIT=PatternFill("solid",fgColor="FFF2CC")
thin=Side(style="thin",color="BFBFBF"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
MM='#,##0.0;(#,##0.0);-'; MM0='#,##0;(#,##0);-'; YEN='#,##0;(#,##0);-'
NUM='#,##0.0;(#,##0.0);-'; NUM2='0.00;(0.00);-'; PCT='0.0%;(0.0%);-'; PCT2='0.00%;(0.00%);-'
YEARS=["FY2027","FY2028","FY2029","FY2030","FY2031"]
YC=["C","D","E","F","G"]; PC=["B","C","D","E","F"]

wb=openpyxl.Workbook()
def S(ws,cell,*,bold=False,color=BLACK,fmt=None,fill=None,align=None,size=10,border=False,wrap=False):
    c=ws[cell]; c.font=Font(name=JP,bold=bold,color=color,size=size)
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    if align or wrap: c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if border: c.border=BOX
    return c
def band(ws,row,text,last,fill=HDR,color="FFFFFF",size=11):
    ws.cell(row=row,column=1,value=text)
    for col in range(1,openpyxl.utils.column_index_from_string(last)+1):
        c=ws.cell(row=row,column=col); c.fill=fill
        c.font=Font(name=JP,bold=True,color=color,size=size)
def sec(ws,row,text,last): band(ws,row,text,last,fill=SUB,color="000000",size=10)

# ══════════ 2. 売上前提 ══════════
ws2=wb.create_sheet("売上前提"); ws2.sheet_view.showGridLines=False
ws2.column_dimensions["A"].width=22; ws2.column_dimensions["B"].width=10
for c in "CDEFGHIJKL": ws2.column_dimensions[c].width=12
ws2.column_dimensions["M"].width=3; ws2.column_dimensions["N"].width=50
band(ws2,1,"売上前提　／　席数 × 席効率 × 回転率 × 客単価 からの積上げ","L")
S(ws2,"A2",size=9,color="7F7F7F")
ws2["A2"]="水色セル＝入力値。ここを変えると単店PLだけでなく5か年連結・資金計画まで全て再計算されます。"
sec(ws2,4,"■ 基本条件","L")
for i,(lab,v,f,note) in enumerate([
    ("営業日数（日／年）",363.0,NUM,"年363日営業"),
    ("立上り率　1年目",0.90,PCT,"新店初年度は保守的に90%"),
    ("立上り率　2年目",0.98,PCT,""),
    ("立上り率　3年目以降（成熟）",1.00,PCT,"満年度＝この水準"),
    ("新店の開業年度 稼働月数",7.5,NUM,"2号店以降。年2店を上期・下期に分散する前提の平均"),
    ("1号店の稼働月数（FY2027）",12.0,NUM,"2027年4月開業＝FY2027は通年稼働")]):
    r=5+i
    S(ws2,f"A{r}"); ws2[f"A{r}"]=lab
    S(ws2,f"B{r}",color=BLUE,fmt=f,border=True,align="center",fill=INP); ws2[f"B{r}"]=v
    S(ws2,f"N{r}",size=9,color="7F7F7F"); ws2[f"N{r}"]=note
R_DAYS,R_RP1,R_RP2,R_RP3,R_MON_NEW,R_MON_TKY=5,6,7,8,9,10
sec(ws2,10+2,"■ 席数（席）","L")
for i,(lab,v) in enumerate([("カウンター",14),("ダイニング",24),("個室",25)]):
    r=13+i
    S(ws2,f"A{r}"); ws2[f"A{r}"]=lab
    S(ws2,f"B{r}",color=BLUE,fmt=MM0,border=True,align="center",fill=INP); ws2[f"B{r}"]=v
S(ws2,"A16",bold=True); ws2["A16"]="店内　小計"
S(ws2,"B16",bold=True,fmt=MM0,border=True,align="center",fill=TOT); ws2["B16"]="=SUM(B13:B15)"
S(ws2,"A17"); ws2["A17"]="テラス"
S(ws2,"B17",color=BLUE,fmt=MM0,border=True,align="center",fill=INP); ws2["B17"]=52
S(ws2,"A18",bold=True); ws2["A18"]="合　計"
S(ws2,"B18",bold=True,fmt=MM0,border=True,align="center",fill=KEY); ws2["B18"]="=B16+B17"
S(ws2,"N13",size=9,color="7F7F7F")
ws2["N13"]="2026年10月時点：店内63席＋テラス52席＝115席。面積55.5坪"
R_IN_SEATS,R_TE_SEATS,R_ALL_SEATS=16,17,18
sec(ws2,20,"■ 時間帯 × エリア 別の売上前提","L")
for c,t in [("A","時間帯"),("B","エリア"),("C","席数"),("D","席効率"),("E","回転率\n(回/日)"),
            ("F","客単価\n料理(円)"),("G","客単価\n飲料(円)"),("H","客単価\n計(円)"),
            ("I","1日\n客数(人)"),("J","日商(円)"),("K","年間売上\n(百万円)"),("L","構成比")]:
    S(ws2,f"{c}21",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True); ws2[f"{c}21"]=t
ws2.row_dimensions[21].height=34
GRID=[("ランチ","店内",f"$B${R_IN_SEATS}",0.75,1.15,3000,1000),
      ("ランチ","テラス",f"$B${R_TE_SEATS}",0.50,1.00,3000,1000),
      ("カフェ（アイドルタイム）","店内",f"$B${R_IN_SEATS}",0.75,0.50,1500,1000),
      ("カフェ（アイドルタイム）","テラス",f"$B${R_TE_SEATS}",0.50,0.50,1500,1000),
      ("ディナー","店内",f"$B${R_IN_SEATS}",0.75,0.90,9000,6000),
      ("ディナー","テラス",f"$B${R_TE_SEATS}",0.50,0.50,5000,5000)]
R_G0=22; R_G1=R_G0+5; R_GT=R_G0+6
for i,(jk,ar,sref,eff,tn,ry,ir) in enumerate(GRID):
    r=R_G0+i
    S(ws2,f"A{r}",border=True,size=9); ws2[f"A{r}"]=jk
    S(ws2,f"B{r}",border=True,size=9,align="center"); ws2[f"B{r}"]=ar
    S(ws2,f"C{r}",color=GREEN,fmt=MM0,border=True,align="center"); ws2[f"C{r}"]=f"={sref}"
    S(ws2,f"D{r}",color=BLUE,fmt=PCT,border=True,align="center",fill=INP); ws2[f"D{r}"]=eff
    S(ws2,f"E{r}",color=BLUE,fmt=NUM2,border=True,align="center",fill=INP); ws2[f"E{r}"]=tn
    S(ws2,f"F{r}",color=BLUE,fmt=YEN,border=True,align="center",fill=INP); ws2[f"F{r}"]=ry
    S(ws2,f"G{r}",color=BLUE,fmt=YEN,border=True,align="center",fill=INP); ws2[f"G{r}"]=ir
    S(ws2,f"H{r}",fmt=YEN,border=True,align="center"); ws2[f"H{r}"]=f"=F{r}+G{r}"
    S(ws2,f"I{r}",fmt=NUM,border=True,align="center"); ws2[f"I{r}"]=f"=C{r}*D{r}*E{r}"
    S(ws2,f"J{r}",fmt=YEN,border=True); ws2[f"J{r}"]=f"=I{r}*H{r}"
    S(ws2,f"K{r}",fmt=MM,border=True); ws2[f"K{r}"]=f"=J{r}*$B${R_DAYS}/1000000"
    S(ws2,f"L{r}",fmt=PCT,border=True,align="center"); ws2[f"L{r}"]=f"=IF($K${R_GT}=0,0,K{r}/$K${R_GT})"
S(ws2,f"A{R_GT}",bold=True,border=True); ws2[f"A{R_GT}"]="合　計"
for c in "BCDEFGH": S(ws2,f"{c}{R_GT}",border=True,fill=TOT)
for c,f_,fm,fl in [("I",f"=SUM(I{R_G0}:I{R_G1})",NUM,TOT),("J",f"=SUM(J{R_G0}:J{R_G1})",YEN,TOT),
                   ("K",f"=SUM(K{R_G0}:K{R_G1})",MM,KEY),("L",f"=SUM(L{R_G0}:L{R_G1})",PCT,TOT)]:
    S(ws2,f"{c}{R_GT}",bold=True,fmt=fm,border=True,align="center" if c!="J" else None,fill=fl)
    ws2[f"{c}{R_GT}"]=f_
S(ws2,f"N{R_G0}",size=9,color=RED)
ws2[f"N{R_G0}"]="★ この表の年間売上合計が、5か年計画の『標準店 満年度売上』になります"
R_RY_SALES,R_IN_SALES=R_GT+2,R_GT+3
sec(ws2,R_GT+1,"■ 料理・飲料別の年間売上（原価計算用）","L")
S(ws2,f"A{R_RY_SALES}"); ws2[f"A{R_RY_SALES}"]="料理売上"
S(ws2,f"B{R_RY_SALES}",fmt=MM,border=True,align="center")
ws2[f"B{R_RY_SALES}"]=f"=SUMPRODUCT(I{R_G0}:I{R_G1},F{R_G0}:F{R_G1})*$B${R_DAYS}/1000000"
S(ws2,f"A{R_IN_SALES}"); ws2[f"A{R_IN_SALES}"]="飲料売上"
S(ws2,f"B{R_IN_SALES}",fmt=MM,border=True,align="center")
ws2[f"B{R_IN_SALES}"]=f"=SUMPRODUCT(I{R_G0}:I{R_G1},G{R_G0}:G{R_G1})*$B${R_DAYS}/1000000"
S(ws2,f"A{R_IN_SALES+1}",bold=True); ws2[f"A{R_IN_SALES+1}"]="合　計"
S(ws2,f"B{R_IN_SALES+1}",bold=True,fmt=MM,border=True,align="center",fill=TOT)
ws2[f"B{R_IN_SALES+1}"]=f"=B{R_RY_SALES}+B{R_IN_SALES}"
R_MIX_RY,R_MIX_IN=R_IN_SALES+2,R_IN_SALES+3
S(ws2,f"A{R_MIX_RY}"); ws2[f"A{R_MIX_RY}"]="料理 構成比"
S(ws2,f"B{R_MIX_RY}",fmt=PCT,border=True,align="center")
ws2[f"B{R_MIX_RY}"]=f"=IF($B${R_IN_SALES+1}=0,0,B{R_RY_SALES}/$B${R_IN_SALES+1})"
S(ws2,f"A{R_MIX_IN}"); ws2[f"A{R_MIX_IN}"]="飲料 構成比"
S(ws2,f"B{R_MIX_IN}",fmt=PCT,border=True,align="center")
ws2[f"B{R_MIX_IN}"]=f"=IF($B${R_IN_SALES+1}=0,0,B{R_IN_SALES}/$B${R_IN_SALES+1})"
R_REF0=R_MIX_IN+2
sec(ws2,R_REF0-1,"■ 参考指標","L")
for i,(lab,f_,fm,note) in enumerate([
    ("日商（円）",f"=J{R_GT}",YEN,""),("月商（百万円）",f"=K{R_GT}/12",MM,""),
    ("年商（百万円・満年度）",f"=K{R_GT}",MM,"＝標準店 満年度売上"),
    ("1日あたり総客数（人）",f"=I{R_GT}",NUM,""),
    ("平均客単価（円）",f"=IF(I{R_GT}=0,0,J{R_GT}/I{R_GT})",YEN,""),
    ("延べ席回転率（回/日）",f"=IF($B${R_ALL_SEATS}=0,0,I{R_GT}/$B${R_ALL_SEATS})",NUM2,""),
    ("坪数（坪）",55.5,NUM,"183.59㎡"),
    ("坪売上（百万円/坪・年）",f"=IF(B{R_REF0+6}=0,0,K{R_GT}/B{R_REF0+6})",MM,""),
    ("1席あたり年間売上（百万円）",f"=IF($B${R_ALL_SEATS}=0,0,K{R_GT}/$B${R_ALL_SEATS})",MM,"席数変更の効果")]):
    r=R_REF0+i
    S(ws2,f"A{r}"); ws2[f"A{r}"]=lab
    isin=not isinstance(f_,str)
    S(ws2,f"B{r}",color=BLUE if isin else BLACK,fmt=fm,border=True,align="center",fill=INP if isin else None)
    ws2[f"B{r}"]=f_
    S(ws2,f"N{r}",size=9,color="7F7F7F"); ws2[f"N{r}"]=note
R_RAMPROW=R_REF0+11; R_SALESROW=R_RAMPROW+1
sec(ws2,R_RAMPROW-2,"■ 立上りを織り込んだ年間売上（百万円）","L")
for c,t in zip("BCDE",["満年度","1年目","2年目","3年目以降"]):
    S(ws2,f"{c}{R_RAMPROW-1}",bold=True,border=True,fill=TOT,align="center"); ws2[f"{c}{R_RAMPROW-1}"]=t
S(ws2,f"A{R_RAMPROW-1}",bold=True,border=True,fill=TOT); ws2[f"A{R_RAMPROW-1}"]="項　目"
S(ws2,f"A{R_RAMPROW}"); ws2[f"A{R_RAMPROW}"]="立上り率"
S(ws2,f"A{R_SALESROW}",bold=True); ws2[f"A{R_SALESROW}"]="年間売上高"
for c,ramp in zip("BCDE",[f"$B${R_RP3}",f"$B${R_RP1}",f"$B${R_RP2}",f"$B${R_RP3}"]):
    S(ws2,f"{c}{R_RAMPROW}",fmt=PCT,border=True,align="center"); ws2[f"{c}{R_RAMPROW}"]=f"={ramp}"
    S(ws2,f"{c}{R_SALESROW}",bold=True,fmt=MM,border=True,align="center",fill=KEY)
    ws2[f"{c}{R_SALESROW}"]=f"=$K${R_GT}*{c}{R_RAMPROW}"
print("売上前提 OK  R_GT=%d RAMP=%d SALES=%d" % (R_GT,R_RAMPROW,R_SALESROW))

# ══════════ 3. 経費前提 ══════════
ws3=wb.create_sheet("経費前提"); ws3.sheet_view.showGridLines=False
ws3.column_dimensions["A"].width=36
for c in "BCD": ws3.column_dimensions[c].width=15
for c in "EF": ws3.column_dimensions[c].width=15
ws3.column_dimensions["G"].width=3; ws3.column_dimensions["H"].width=56
band(ws3,1,"経費前提・投資前提（水色セル＝入力値）","H")
S(ws3,"A2",size=9,color="7F7F7F")
ws3["A2"]="単店PLと5か年連結の両方がこのシートを参照します。経費率は2026年8月3日提出P&L準拠。"
sec(ws3,4,"■ 売上原価の計算方法","H")
S(ws3,"A5"); ws3["A5"]="計算方法（1＝売上比一括／2＝料理・飲料別）"
S(ws3,"B5",color=BLUE,fmt=MM0,border=True,align="center",fill=INP); ws3["B5"]=2
S(ws3,"H5",size=9,color="7F7F7F")
ws3["H5"]="2を選ぶと料理・飲料それぞれの原価率で計算（直輸入ワインの効果を分けて見る場合）"
R_METHOD=5
sec(ws3,7,"■ 経費率（対売上高）","H")
for c,t in zip("ABCD",["費　目","提出P&L水準\n（満年度・1年目）","2年目","3年目以降"]):
    S(ws3,f"{c}8",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True); ws3[f"{c}8"]=t
ws3.row_dimensions[8].height=32
RATES=[("売上原価（方法1：売上比一括）",0.295,0.290,0.285,"方法1のときに使用"),
       ("　料理原価率（方法2）",0.300,0.300,0.300,"方法2のときに使用"),
       ("　飲料原価率（方法2）",0.300,0.300,0.300,"同上。直輸入ワインで下げる余地あり"),
       ("実効売上原価率（方法に連動）",None,None,None,"計算結果。以降のPLはこの率を使用"),
       ("人件費",0.345,0.335,0.325,"本部集中化・多能工化。人員26名"),
       ("業務委託費・消耗品他",0.024,0.022,0.020,"本部集中購買・規格統一"),
       ("広告宣伝費・販売促進費",0.014,0.014,0.014,"据置"),
       ("その他経費",0.005,0.005,0.005,"据置"),
       ("水道光熱費",0.040,0.038,0.036,"高効率厨房機器・営業時間最適化"),
       ("ロイヤリティ",0.040,0.040,0.040,"契約条件（ブラッスリー4.0%）"),
       ("保険・租税公課",0.007,0.007,0.007,"据置"),
       ("修繕維持費（FFE）",0.025,0.025,0.025,"据置")]
R_R0=9
for i,(lab,a,b,c_,note) in enumerate(RATES):
    r=R_R0+i; sub=lab.startswith("　"); calc=(a is None)
    S(ws3,f"A{r}",border=True,bold=calc,size=9 if sub else 10,color="7F7F7F" if sub else BLACK)
    ws3[f"A{r}"]=lab
    for col,v in zip("BCD",[a,b,c_]):
        if calc:
            S(ws3,f"{col}{r}",bold=True,fmt=PCT2,border=True,align="center",fill=TOT)
            ws3[f"{col}{r}"]=(f"=IF($B${R_METHOD}=1,{col}{R_R0},"
                              f"売上前提!$B${R_MIX_RY}*{col}{R_R0+1}+売上前提!$B${R_MIX_IN}*{col}{R_R0+2})")
        else:
            S(ws3,f"{col}{r}",color=BLUE,fmt=PCT2 if sub else PCT,border=True,align="center",fill=INP)
            ws3[f"{col}{r}"]=v
    S(ws3,f"H{r}",size=9,color=RED if calc else "7F7F7F"); ws3[f"H{r}"]=note
R_COST1,R_RY_R,R_IN_R,R_CEFF=R_R0,R_R0+1,R_R0+2,R_R0+3            # 9,10,11,12
R_JIN,R_ITAKU,R_KOK,R_SON,R_SUI,R_ROY,R_HOK,R_SHU=[R_R0+4+i for i in range(8)]  # 13..20
R_VSUM=R_R0+12                                                     # 21
RATEROWS=[R_CEFF,R_JIN,R_ITAKU,R_KOK,R_SON,R_SUI,R_ROY,R_HOK,R_SHU]
S(ws3,f"A{R_VSUM}",bold=True,border=True); ws3[f"A{R_VSUM}"]="変動費　計（対売上）"
for col in "BCD":
    S(ws3,f"{col}{R_VSUM}",bold=True,fmt=PCT,border=True,align="center",fill=TOT)
    ws3[f"{col}{R_VSUM}"]=f"={col}{R_CEFF}+SUM({col}{R_JIN}:{col}{R_SHU})"
sec(ws3,23,"■ 固定費・家賃条件（全店共通）","H")
for i,(lab,v,f_,note) in enumerate([
    ("地代家賃（固定）月額",2.22,MM,"月222万円。席数を変えても面積は同じため不変"),
    ("地代家賃（固定）年額",None,MM,"＝月額×12"),
    ("歩合家賃率",0.07,PCT,"閾値超過分に対して7%"),
    ("歩合家賃 閾値（年商）",331.4,MM,"提出P&L時点の水準から逆算"),
    ("人員数（名・参考）",26,MM0,"人件費は売上比で計算")]):
    r=24+i
    S(ws3,f"A{r}"); ws3[f"A{r}"]=lab
    isin=v is not None
    S(ws3,f"B{r}",color=BLUE if isin else BLACK,fmt=f_,border=True,align="center",fill=INP if isin else None)
    ws3[f"B{r}"]=v if isin else "=B24*12"
    S(ws3,f"H{r}",size=9,color="7F7F7F"); ws3[f"H{r}"]=note
R_RENT_M,R_RENT_Y,R_PCTR,R_THR,R_STAFF=24,25,26,27,28
sec(ws3,30,"■ 初期投資（1店舗あたり・百万円）","H")
S(ws3,"A31",bold=True,border=True,fill=TOT); ws3["A31"]="項　目"
for c,t in zip("BCD",["1号店\n(六本木)","2〜3号店\n(札幌・大阪)","4号店以降"]):
    S(ws3,f"{c}31",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True); ws3[f"{c}31"]=t
ws3.row_dimensions[31].height=28
R_INV0=32
for i,(lab,a,b,c3,note) in enumerate([
    ("契約金",10.0,10.0,5.0,"契約金 3店舗3,000万円／以下7店舗3,500万円"),
    ("事業費（設計・内装施工・OSE/FF&E）",194.25,194.25,194.25,"55.5坪×350万"),
    ("開業準備金（広告・採用等）",4.0,4.0,4.0,""),
    ("保証金",39.0,39.0,39.0,"非償却・退去時返還対象"),
    ("ワイン・美術品在庫",50.0,10.0,10.0,"1号店は美術品含む。2号店以降はワイン1,000万円のみ（非償却）")]):
    r=R_INV0+i
    S(ws3,f"A{r}",border=True); ws3[f"A{r}"]=lab
    for col,v in zip("BCD",[a,b,c3]):
        S(ws3,f"{col}{r}",color=BLUE,fmt=NUM,border=True,align="center",fill=INP); ws3[f"{col}{r}"]=v
    S(ws3,f"H{r}",size=9,color="7F7F7F"); ws3[f"H{r}"]=note
R_INVTOT,R_DEPBASE,R_DEPYR,R_DEPRE=37,38,39,40
S(ws3,f"A{R_INVTOT}",bold=True,border=True); ws3[f"A{R_INVTOT}"]="初期投資　合計"
S(ws3,f"A{R_DEPBASE}",bold=True,border=True); ws3[f"A{R_DEPBASE}"]="償却対象資産（＝合計－保証金－ワイン・美術品）"
S(ws3,f"A{R_DEPYR}",border=True); ws3[f"A{R_DEPYR}"]="償却年数（年）"
S(ws3,f"A{R_DEPRE}",bold=True,border=True); ws3[f"A{R_DEPRE}"]="年間減価償却費／1店（通年）"
for col in "BCD":
    S(ws3,f"{col}{R_INVTOT}",bold=True,fmt=NUM,border=True,align="center",fill=KEY)
    ws3[f"{col}{R_INVTOT}"]=f"=SUM({col}{R_INV0}:{col}{R_INV0+4})"
    S(ws3,f"{col}{R_DEPBASE}",bold=True,fmt=NUM,border=True,align="center",fill=TOT)
    ws3[f"{col}{R_DEPBASE}"]=f"={col}{R_INVTOT}-{col}{R_INV0+3}-{col}{R_INV0+4}"
    S(ws3,f"{col}{R_DEPYR}",color=BLUE,fmt=NUM,border=True,align="center",fill=INP); ws3[f"{col}{R_DEPYR}"]=10.0
    S(ws3,f"{col}{R_DEPRE}",bold=True,fmt=NUM,border=True,align="center",fill=TOT)
    ws3[f"{col}{R_DEPRE}"]=f"={col}{R_DEPBASE}/{col}{R_DEPYR}"
S(ws3,f"H{R_DEPBASE}",size=9,color=RED)
ws3[f"H{R_DEPBASE}"]="ワイン在庫は棚卸資産、美術品は取得価額100万円以上なら原則非償却のため除外"
sec(ws3,42,"■ 分割払いスキーム（百万円）","H")
R_DEF,R_CASH,R_PAY=45,46,47
for i,(lab,v,note) in enumerate([
    ("分割対象：内装工事費",100.0,"月833千円 × 120回（10年）"),
    ("分割対象：保証金",38.9,"月463千円 × 84回（7年）"),
    ("分割対象額　計",None,""),
    ("開業時 現金支出／1店",None,"＝初期投資合計 － 分割対象額"),
    ("年間分割弁済額／1店",15.6,"月額合計1,296千円 × 12か月")]):
    r=43+i
    S(ws3,f"A{r}",bold=(v is None)); ws3[f"A{r}"]=lab
    if lab=="分割対象額　計":
        S(ws3,f"B{r}",bold=True,fmt=NUM,border=True,align="center",fill=TOT); ws3[f"B{r}"]="=B43+B44"
    elif lab=="開業時 現金支出／1店":
        for col in "BCD":
            S(ws3,f"{col}{r}",bold=True,fmt=NUM,border=True,align="center",fill=TOT)
            ws3[f"{col}{r}"]=f"={col}{R_INVTOT}-$B${R_DEF}"
    else:
        S(ws3,f"B{r}",color=BLUE,fmt=NUM,border=True,align="center",fill=INP); ws3[f"B{r}"]=v
    S(ws3,f"H{r}",size=9,color="7F7F7F"); ws3[f"H{r}"]=note
sec(ws3,49,"■ 出店計画・本部費（連結）","H")
R_OPEN,R_STORES,R_HQ=51,52,53
S(ws3,"A50",bold=True,border=True,fill=TOT); ws3["A50"]="項　目"
for i,y in enumerate(YEARS):
    S(ws3,f"{PC[i]}50",bold=True,border=True,fill=TOT,align="center"); ws3[f"{PC[i]}50"]=y
S(ws3,f"A{R_OPEN}"); ws3[f"A{R_OPEN}"]="期中出店数（店）"
S(ws3,f"A{R_STORES}",bold=True); ws3[f"A{R_STORES}"]="期末店舗数（店）"
S(ws3,f"A{R_HQ}"); ws3[f"A{R_HQ}"]="本部費（管理・人事・購買・マーケ）"
for i,cl in enumerate(PC):
    S(ws3,f"{cl}{R_OPEN}",color=BLUE,fmt=MM0,border=True,align="center",fill=INP)
    ws3[f"{cl}{R_OPEN}"]=[1,2,2,2,2][i]
    S(ws3,f"{cl}{R_STORES}",bold=True,fmt=MM0,border=True,align="center",fill=TOT)
    ws3[f"{cl}{R_STORES}"]=f"={cl}{R_OPEN}" if i==0 else f"={PC[i-1]}{R_STORES}+{cl}{R_OPEN}"
    S(ws3,f"{cl}{R_HQ}",color=BLUE,fmt=NUM,border=True,align="center",fill=INP)
    ws3[f"{cl}{R_HQ}"]=[15.0,28.0,40.0,48.0,55.0][i]
S(ws3,f"H{R_OPEN}",size=9,color="7F7F7F")
ws3[f"H{R_OPEN}"]="FY2027 六本木／FY2028 札幌・大阪／FY2029 首都圏・他政令市／FY2030以降 リゾート含め年2店"
print("経費前提 OK  CEFF=%d VSUM=%d" % (R_CEFF,R_VSUM))

PL_ORDER=[("売上原価","var",0),("人件費","var",1),("業務委託費・消耗品他","var",2),
          ("部門利益","sub_dept",None),("広告宣伝費・販売促進費","var",3),("その他経費","var",4),
          ("水道光熱費","var",5),("GOP（店舗営業総利益）","sub_gop",None),
          ("ロイヤリティ","var",6),("保険・租税公課","var",7),
          ("地代家賃（固定）","rent_fix",None),("歩合家賃","rent_pct",None),
          ("修繕維持費（FFE）","var",8)]

# ══════════ 4. 単店PL ══════════
ws4=wb.create_sheet("単店PL"); ws4.sheet_view.showGridLines=False
ws4.column_dimensions["A"].width=30
for c in "BCDE": ws4.column_dimensions[c].width=14
ws4.column_dimensions["F"].width=3; ws4.column_dimensions["G"].width=42
band(ws4,1,"六本木ミッドタウン店（1号店）　単店PL（百万円）","E")
S(ws4,"A2",size=9,color="7F7F7F")
ws4["A2"]="売上は「売上前提」、経費率・減価償却は「経費前提」に連動。科目は2026年8月3日提出P&Lのまま。"
COLS=["B","C","D","E"]; RATECOL=["B","B","C","D"]
for c,t in zip(COLS,["満年度\n（提出P&L水準）","1年目","2年目","3年目以降"]):
    S(ws4,f"{c}4",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True); ws4[f"{c}4"]=t
S(ws4,"A4",bold=True,border=True,fill=TOT); ws4["A4"]="科　目"
ws4.row_dimensions[4].height=30
r=5
S(ws4,f"A{r}"); ws4[f"A{r}"]="立上り率"
for c in COLS:
    S(ws4,f"{c}{r}",color=GREEN,fmt=PCT,border=True,align="center"); ws4[f"{c}{r}"]=f"=売上前提!{c}{R_RAMPROW}"
r+=1
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="売上高"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,color=GREEN,fmt=MM,border=True); ws4[f"{c}{r}"]=f"=売上前提!{c}${R_SALESROW}"
R_REV=r; r+=1
for lab,ref in [("　うち 料理売上",R_MIX_RY),("　うち 飲料売上",R_MIX_IN)]:
    S(ws4,f"A{r}",size=9,color="7F7F7F",border=True); ws4[f"A{r}"]=lab
    for c in COLS:
        S(ws4,f"{c}{r}",size=9,color="7F7F7F",fmt=MM,border=True)
        ws4[f"{c}{r}"]=f"={c}${R_REV}*売上前提!$B${ref}"
    r+=1
PL=[]
def var_row(label,idx):
    global r
    S(ws4,f"A{r}",border=True); ws4[f"A{r}"]=label
    for c,rc in zip(COLS,RATECOL):
        S(ws4,f"{c}{r}",fmt=MM,border=True)
        ws4[f"{c}{r}"]=f"={c}${R_REV}*経費前提!${rc}${RATEROWS[idx]}"
    PL.append((label,r)); r+=1; return r-1
R_CGS=var_row("売上原価",0); R_JINK=var_row("人件費",1); R_ITK=var_row("業務委託費・消耗品他",2)
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="部門利益"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=TOT)
    ws4[f"{c}{r}"]=f"={c}{R_REV}-{c}{R_CGS}-{c}{R_JINK}-{c}{R_ITK}"
R_BUMON=r; PL.append(("部門利益",r)); r+=1
R_KK=var_row("広告宣伝費・販売促進費",3); R_SN=var_row("その他経費",4); R_SU=var_row("水道光熱費",5)
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="GOP（店舗営業総利益）"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=TOT)
    ws4[f"{c}{r}"]=f"={c}{R_BUMON}-{c}{R_KK}-{c}{R_SN}-{c}{R_SU}"
R_GOP=r; PL.append(("GOP（店舗営業総利益）",r)); r+=1
R_RO=var_row("ロイヤリティ",6); R_HK=var_row("保険・租税公課",7)
S(ws4,f"A{r}",border=True); ws4[f"A{r}"]="地代家賃（固定）"
for c in COLS:
    S(ws4,f"{c}{r}",color=GREEN,fmt=MM,border=True); ws4[f"{c}{r}"]=f"=経費前提!$B${R_RENT_Y}"
R_RENT=r; PL.append(("地代家賃（固定）",r)); r+=1
S(ws4,f"A{r}",border=True); ws4[f"A{r}"]="歩合家賃"
for c in COLS:
    S(ws4,f"{c}{r}",fmt=MM,border=True)
    ws4[f"{c}{r}"]=f"=MAX(0,{c}{R_REV}-経費前提!$B${R_THR})*経費前提!$B${R_PCTR}"
R_PRENT=r; PL.append(("歩合家賃",r)); r+=1
R_SH=var_row("修繕維持費（FFE）",8)
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="EBITDA"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=KEY)
    ws4[f"{c}{r}"]=f"={c}{R_GOP}-{c}{R_RO}-{c}{R_HK}-{c}{R_RENT}-{c}{R_PRENT}-{c}{R_SH}"
R_EBITDA=r; PL.append(("EBITDA",r)); r+=1
S(ws4,f"A{r}",border=True); ws4[f"A{r}"]="減価償却費"
for c in COLS:
    S(ws4,f"{c}{r}",color=GREEN,fmt=MM,border=True); ws4[f"{c}{r}"]=f"=経費前提!$B${R_DEPRE}"
R_DEP=r; PL.append(("減価償却費",r)); r+=1
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="営業利益"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=KEY); ws4[f"{c}{r}"]=f"={c}{R_EBITDA}-{c}{R_DEP}"
R_OP=r; PL.append(("営業利益",r)); r+=2
sec(ws4,r,"■ 対売上比","E"); r+=1
for c,t in zip(COLS,["満年度","1年目","2年目","3年目以降"]):
    S(ws4,f"{c}{r}",bold=True,border=True,fill=TOT,align="center"); ws4[f"{c}{r}"]=t
S(ws4,f"A{r}",bold=True,border=True,fill=TOT); ws4[f"A{r}"]="科　目"; r+=1
for lab,src in [("売上高",R_REV)]+PL:
    hl=lab in ("EBITDA","営業利益","部門利益","GOP（店舗営業総利益）")
    S(ws4,f"A{r}",border=True,bold=hl); ws4[f"A{r}"]=lab
    for c in COLS:
        S(ws4,f"{c}{r}",fmt=PCT,border=True,align="center",bold=hl,
          fill=KEY if lab in ("EBITDA","営業利益") else (TOT if hl else None))
        ws4[f"{c}{r}"]=f"=IF({c}${R_REV}=0,0,{c}{src}/{c}${R_REV})"
    r+=1

# ══════════ 5. 店舗群別PL ══════════
ws5=wb.create_sheet("店舗群別PL"); ws5.sheet_view.showGridLines=False
ws5.column_dimensions["A"].width=30; ws5.column_dimensions["B"].width=11
for c in YC: ws5.column_dimensions[c].width=13
band(ws5,1,"店舗群（出店年度コホート）別PL（百万円）","G")
S(ws5,"A2",size=9,color="7F7F7F")
ws5["A2"]="各群の売上は「売上前提」の標準店年商に立上り率・店舗数・稼働月数を掛けたもの。経費率は年次に応じて自動参照。"
S(ws5,"A3",bold=True,border=True,fill=TOT); ws5["A3"]="店舗群"
S(ws5,"B3",bold=True,border=True,fill=TOT,align="center"); ws5["B3"]="店舗数"
for i,y in enumerate(YEARS):
    S(ws5,f"{YC[i]}3",bold=True,border=True,fill=TOT,align="center"); ws5[f"{YC[i]}3"]=y
COHORTS=[("① 東京（六本木ミッドタウン）",1,f"$B${R_MON_TKY}","B",[1,2,3,3,3]),
         ("② 札幌・大阪",2,f"$B${R_MON_NEW}","C",[0,1,2,3,3]),
         ("③ 首都圏・他政令指定都市",2,f"$B${R_MON_NEW}","D",[0,0,1,2,3]),
         ("④ FY2030出店（リゾート含む）",2,f"$B${R_MON_NEW}","D",[0,0,0,1,2]),
         ("⑤ FY2031出店（リゾート含む）",2,f"$B${R_MON_NEW}","D",[0,0,0,0,1])]
blocks=[]; r=5
for name,n,mref,invcol,yidx in COHORTS:
    sec(ws5,r,f"　{name}","G")
    r_n,r_m,r_y,r_rev=r+1,r+2,r+3,r+4
    S(ws5,f"A{r_n}",size=9); ws5[f"A{r_n}"]="店舗数（店）"
    S(ws5,f"B{r_n}",color=BLUE,fmt=MM0,border=True,align="center",fill=INP); ws5[f"B{r_n}"]=n
    S(ws5,f"A{r_m}",size=9); ws5[f"A{r_m}"]="開業年度の稼働月数"
    S(ws5,f"B{r_m}",color=GREEN,fmt=NUM,border=True,align="center"); ws5[f"B{r_m}"]=f"=売上前提!{mref}"
    S(ws5,f"A{r_y}",size=9); ws5[f"A{r_y}"]="年次（0=未開業/1/2/3）"
    for i,cl in enumerate(YC):
        S(ws5,f"{cl}{r_y}",color=BLUE,fmt=MM0,border=True,align="center",fill=INP); ws5[f"{cl}{r_y}"]=yidx[i]
    S(ws5,f"A{r_rev}",bold=True,border=True); ws5[f"A{r_rev}"]="売上高"
    for cl in YC:
        S(ws5,f"{cl}{r_rev}",bold=True,fmt=MM,border=True)
        ws5[f"{cl}{r_rev}"]=(f"=IF({cl}{r_y}=0,0,売上前提!$K${R_GT}"
                             f"*INDEX(売上前提!$B${R_RP1}:$B${R_RP3},{cl}{r_y})"
                             f"*$B${r_n}*IF({cl}{r_y}=1,$B${r_m}/12,1))")
    rr=r_rev+1; rowmap={}
    for lab,kind,idx in PL_ORDER:
        S(ws5,f"A{rr}",size=9,bold=kind.startswith("sub"),border=True); ws5[f"A{rr}"]=lab
        for cl in YC:
            if kind=="var":
                S(ws5,f"{cl}{rr}",fmt=MM,border=True)
                ws5[f"{cl}{rr}"]=(f"=IF({cl}${r_y}=0,0,{cl}${r_rev}*"
                                  f"INDEX(経費前提!$B${RATEROWS[idx]}:$D${RATEROWS[idx]},1,{cl}${r_y}))")
            elif kind=="sub_dept":
                S(ws5,f"{cl}{rr}",bold=True,fmt=MM,border=True,fill=TOT)
                ws5[f"{cl}{rr}"]=f"={cl}{r_rev}-SUM({cl}{r_rev+1}:{cl}{rr-1})"
            elif kind=="sub_gop":
                S(ws5,f"{cl}{rr}",bold=True,fmt=MM,border=True,fill=TOT)
                dr=rowmap["部門利益"]; ws5[f"{cl}{rr}"]=f"={cl}{dr}-SUM({cl}{dr+1}:{cl}{rr-1})"
            elif kind=="rent_fix":
                S(ws5,f"{cl}{rr}",fmt=MM,border=True)
                ws5[f"{cl}{rr}"]=(f"=IF({cl}{r_y}=0,0,経費前提!$B${R_RENT_Y}*$B${r_n}"
                                  f"*IF({cl}{r_y}=1,$B${r_m}/12,1))")
            elif kind=="rent_pct":
                S(ws5,f"{cl}{rr}",fmt=MM,border=True)
                ws5[f"{cl}{rr}"]=(f"=IF({cl}{r_y}=0,0,MAX(0,{cl}{r_rev}/$B${r_n}-経費前提!$B${R_THR}"
                                  f"*IF({cl}{r_y}=1,$B${r_m}/12,1))*経費前提!$B${R_PCTR}*$B${r_n})")
        rowmap[lab]=rr; rr+=1
    r_eb,r_dep,r_op=rr,rr+1,rr+2
    S(ws5,f"A{r_eb}",bold=True,border=True); ws5[f"A{r_eb}"]="店舗EBITDA"
    S(ws5,f"A{r_dep}",border=True); ws5[f"A{r_dep}"]="減価償却費"
    S(ws5,f"A{r_op}",bold=True,border=True); ws5[f"A{r_op}"]="店舗営業利益"
    gr=rowmap["GOP（店舗営業総利益）"]
    for cl in YC:
        S(ws5,f"{cl}{r_eb}",bold=True,fmt=MM,border=True,fill=TOT)
        ws5[f"{cl}{r_eb}"]=f"={cl}{gr}-SUM({cl}{gr+1}:{cl}{r_eb-1})"
        S(ws5,f"{cl}{r_dep}",fmt=MM,border=True)
        ws5[f"{cl}{r_dep}"]=(f"=IF({cl}{r_y}=0,0,経費前提!${invcol}${R_DEPRE}*$B${r_n}"
                             f"*IF({cl}{r_y}=1,$B${r_m}/12,1))")
        S(ws5,f"{cl}{r_op}",bold=True,fmt=MM,border=True,fill=TOT)
        ws5[f"{cl}{r_op}"]=f"={cl}{r_eb}-{cl}{r_dep}"
    blocks.append(dict(rev=r_rev,rows=rowmap,eb=r_eb,dep=r_dep,op=r_op,n=r_n,m=r_m,y=r_y,invcol=invcol))
    r=r_op+2
print("単店PL/店舗群別PL OK")

# ══════════ 6. 連結PL ══════════
ws6=wb.create_sheet("連結PL"); ws6.sheet_view.showGridLines=False
ws6.column_dimensions["A"].width=30; ws6.column_dimensions["B"].width=10
for c in YC: ws6.column_dimensions[c].width=14
ws6.column_dimensions["H"].width=14
band(ws6,1,"5か年事業計画　連結損益計画（FY2027-FY2031・百万円）","H")
S(ws6,"A2",size=9,color="7F7F7F")
ws6["A2"]="2027年4月 六本木開業 → 2028年 札幌・大阪 → 2029年 首都圏・政令指定都市 → 2030年以降 リゾート含め年2店舗"
S(ws6,"A3",bold=True,border=True,fill=TOT); ws6["A3"]="科　目"
S(ws6,"B3",bold=True,border=True,fill=TOT,align="center",wrap=True); ws6["B3"]="対売上比\n(FY2031)"
for i,y in enumerate(YEARS):
    S(ws6,f"{YC[i]}3",bold=True,border=True,fill=TOT,align="center"); ws6[f"{YC[i]}3"]=y
S(ws6,"H3",bold=True,border=True,fill=TOT,align="center"); ws6["H3"]="5年累計"
def agg(k,cl): return "+".join(f"店舗群別PL!{cl}{b[k]}" for b in blocks)
def aggrow(lab,cl): return "+".join(f"店舗群別PL!{cl}{b['rows'][lab]}" for b in blocks)
R4={}; row=4
S(ws6,f"A{row}",bold=True,border=True); ws6[f"A{row}"]="期末店舗数（店）"
for i,cl in enumerate(YC):
    S(ws6,f"{cl}{row}",bold=True,color=GREEN,fmt=MM0,border=True,align="center")
    ws6[f"{cl}{row}"]=f"=経費前提!{PC[i]}{R_STORES}"
R4["stores"]=row; row+=1
S(ws6,f"A{row}",bold=True,border=True); ws6[f"A{row}"]="売上高"
for cl in YC:
    S(ws6,f"{cl}{row}",bold=True,color=GREEN,fmt=MM,border=True); ws6[f"{cl}{row}"]="="+agg("rev",cl)
R4["rev"]=row; row+=1
for lab,kind,idx in PL_ORDER:
    S(ws6,f"A{row}",size=9,bold=kind.startswith("sub"),border=True); ws6[f"A{row}"]=lab
    for cl in YC:
        S(ws6,f"{cl}{row}",bold=kind.startswith("sub"),color=GREEN,fmt=MM,border=True,
          fill=TOT if kind.startswith("sub") else None)
        ws6[f"{cl}{row}"]="="+aggrow(lab,cl)
    R4[lab]=row; row+=1
S(ws6,f"A{row}",bold=True,border=True); ws6[f"A{row}"]="店舗EBITDA　計"
for cl in YC:
    S(ws6,f"{cl}{row}",bold=True,color=GREEN,fmt=MM,border=True,fill=TOT); ws6[f"{cl}{row}"]="="+agg("eb",cl)
R4["st_eb"]=row; row+=1
S(ws6,f"A{row}",border=True); ws6[f"A{row}"]="本部費"
for i,cl in enumerate(YC):
    S(ws6,f"{cl}{row}",color=GREEN,fmt=MM,border=True); ws6[f"{cl}{row}"]=f"=経費前提!{PC[i]}{R_HQ}"
R4["hq"]=row; row+=1
S(ws6,f"A{row}",bold=True,border=True); ws6[f"A{row}"]="連結EBITDA"
for cl in YC:
    S(ws6,f"{cl}{row}",bold=True,fmt=MM,border=True,fill=KEY)
    ws6[f"{cl}{row}"]=f"={cl}{R4['st_eb']}-{cl}{R4['hq']}"
R4["eb"]=row; row+=1
S(ws6,f"A{row}",bold=True,border=True); ws6[f"A{row}"]="　EBITDA率"
for cl in YC:
    S(ws6,f"{cl}{row}",bold=True,fmt=PCT,border=True,align="center",fill=KEY)
    ws6[f"{cl}{row}"]=f"=IF({cl}{R4['rev']}=0,0,{cl}{R4['eb']}/{cl}{R4['rev']})"
R4["ebm"]=row; row+=1
S(ws6,f"A{row}",border=True); ws6[f"A{row}"]="減価償却費"
for cl in YC:
    S(ws6,f"{cl}{row}",color=GREEN,fmt=MM,border=True); ws6[f"{cl}{row}"]="="+agg("dep",cl)
R4["dep"]=row; row+=1
S(ws6,f"A{row}",bold=True,border=True); ws6[f"A{row}"]="営業利益"
for cl in YC:
    S(ws6,f"{cl}{row}",bold=True,fmt=MM,border=True,fill=KEY)
    ws6[f"{cl}{row}"]=f"={cl}{R4['eb']}-{cl}{R4['dep']}"
R4["op"]=row; row+=1
S(ws6,f"A{row}",bold=True,border=True); ws6[f"A{row}"]="　営業利益率"
for cl in YC:
    S(ws6,f"{cl}{row}",bold=True,fmt=PCT,border=True,align="center",fill=KEY)
    ws6[f"{cl}{row}"]=f"=IF({cl}{R4['rev']}=0,0,{cl}{R4['op']}/{cl}{R4['rev']})"
R4["opm"]=row; row+=1
plrows=[R4["rev"]]+[R4[l] for l,_,_ in PL_ORDER]+[R4["st_eb"],R4["hq"],R4["eb"],R4["dep"],R4["op"]]
for rr in plrows:
    S(ws6,f"H{rr}",bold=True,fmt=MM,border=True,fill=TOT); ws6[f"H{rr}"]=f"=SUM(C{rr}:G{rr})"
    S(ws6,f"B{rr}",size=9,fmt=PCT,border=True,align="center")
    ws6[f"B{rr}"]=f"=IF($G${R4['rev']}=0,0,G{rr}/$G${R4['rev']})"
for k in ["ebm","opm"]:
    rr=R4[k]
    S(ws6,f"H{rr}",bold=True,fmt=PCT,border=True,align="center",fill=TOT)
    ws6[f"H{rr}"]=f"=H{R4['eb' if k=='ebm' else 'op']}/H{R4['rev']}"
S(ws6,f"A{row+1}",size=9,color="7F7F7F")
ws6[f"A{row+1}"]="※ FY＝4月〜3月。新店初年度は売上90%・開業7.5か月で計上。金利・税金は含まない。"

# ══════════ 7. 投資・CF ══════════
ws7=wb.create_sheet("投資・CF"); ws7.sheet_view.showGridLines=False
ws7.column_dimensions["A"].width=34
for c in "BCDEFG": ws7.column_dimensions[c].width=13
ws7.column_dimensions["H"].width=14; ws7.column_dimensions["I"].width=3
ws7.column_dimensions["J"].width=46
band(ws7,1,"投資・キャッシュフロー（百万円）","H")
S(ws7,"A2",size=9,color="7F7F7F")
ws7["A2"]="上段＝1号店の投資回収、下段＝連結5か年のキャッシュフローと必要資金。"
sec(ws7,4,"■ 1号店 キャッシュフロー推移","H")
SC=["B","C","D","E","F","G"]
for c,t in zip(SC,["開業時","1年目","2年目","3年目","4年目","5年目"]):
    S(ws7,f"{c}5",bold=True,border=True,fill=TOT,align="center"); ws7[f"{c}5"]=t
S(ws7,"A5",bold=True,border=True,fill=TOT); ws7["A5"]="項　目"
EBCOL=[None,"C","D","E","E","E"]
S(ws7,"A6",border=True); ws7["A6"]="EBITDA"
for c,src in zip(SC,EBCOL):
    S(ws7,f"{c}6",color=GREEN if src else BLACK,fmt=MM,border=True,align="center")
    ws7[f"{c}6"]=(f"=単店PL!{src}{R_EBITDA}" if src else 0)
S(ws7,"A7",border=True); ws7["A7"]="開業時 現金支出"
for i,c in enumerate(SC):
    S(ws7,f"{c}7",fmt=MM,border=True,align="center")
    ws7[f"{c}7"]=(f"=-経費前提!$B${R_CASH}" if i==0 else 0)
S(ws7,"A8",border=True); ws7["A8"]="分割弁済"
for i,c in enumerate(SC):
    S(ws7,f"{c}8",fmt=MM,border=True,align="center")
    ws7[f"{c}8"]=(0 if i==0 else f"=-経費前提!$B${R_PAY}")
S(ws7,"A9",bold=True,border=True); ws7["A9"]="単店FCF"
for c in SC:
    S(ws7,f"{c}9",bold=True,fmt=MM,border=True,align="center",fill=TOT); ws7[f"{c}9"]=f"={c}6+{c}7+{c}8"
S(ws7,"A10",bold=True,border=True); ws7["A10"]="累計FCF"
for i,c in enumerate(SC):
    S(ws7,f"{c}10",bold=True,fmt=MM,border=True,align="center",fill=KEY)
    ws7[f"{c}10"]=f"={c}9" if i==0 else f"={SC[i-1]}10+{c}9"
S(ws7,"A12",border=True); ws7["A12"]="累計FCFが赤字の年数"
S(ws7,"B12",fmt=MM0,border=True,align="center"); ws7["B12"]='=COUNTIF(C10:G10,"<0")'
S(ws7,"A13",bold=True,border=True); ws7["A13"]="累計FCFの黒字転換"
S(ws7,"B13",bold=True,border=True,align="center",fill=KEY)
ws7["B13"]='=IF(B12>=5,"5年内に未回収",B12+1&"年目")'
S(ws7,"A14",bold=True,border=True); ws7["A14"]="投資回収年数（年）"
S(ws7,"B14",bold=True,fmt=NUM2,border=True,align="center",fill=KEY)
ws7["B14"]=('=IF(B12>=5,"",IF(INDEX(B9:G9,B12+2)=0,"",B12+(-INDEX(B10:G10,B12+1))/INDEX(B9:G9,B12+2)))')
S(ws7,"A16",bold=True,size=10,color=RED)
ws7["A16"]='="→ 1号店は開業"&B13&"に累計FCFが黒字転換（投資回収 約"&TEXT(B14,"0.0")&"年）。"'

sec(ws7,18,"■ 連結キャッシュフロー計画","H")
S(ws7,"A19",bold=True,border=True,fill=TOT); ws7["A19"]="項　目"
for i,y in enumerate(YEARS):
    S(ws7,f"{YC[i]}19",bold=True,border=True,fill=TOT,align="center"); ws7[f"{YC[i]}19"]=y
S(ws7,"H19",bold=True,border=True,fill=TOT,align="center"); ws7["H19"]="5年累計"
INV_MAP=["B","C","D","D","D"]
r=20
S(ws7,f"A{r}",border=True); ws7[f"A{r}"]="新規出店数（店）"
for i,cl in enumerate(YC):
    S(ws7,f"{cl}{r}",color=GREEN,fmt=MM0,border=True,align="center"); ws7[f"{cl}{r}"]=f"=経費前提!{PC[i]}{R_OPEN}"
ROW_OPEN=r; r+=1
S(ws7,f"A{r}",border=True); ws7[f"A{r}"]="期末店舗数（店）"
for i,cl in enumerate(YC):
    S(ws7,f"{cl}{r}",color=GREEN,fmt=MM0,border=True,align="center"); ws7[f"{cl}{r}"]=f"=経費前提!{PC[i]}{R_STORES}"
r+=1
S(ws7,f"A{r}",border=True); ws7[f"A{r}"]="初期投資（総額）"
for i,cl in enumerate(YC):
    S(ws7,f"{cl}{r}",fmt=MM,border=True); ws7[f"{cl}{r}"]=f"=-{cl}{ROW_OPEN}*経費前提!${INV_MAP[i]}${R_INVTOT}"
ROW_INV=r; r+=1
S(ws7,f"A{r}",border=True); ws7[f"A{r}"]="うち 分割対象額"
for cl in YC:
    S(ws7,f"{cl}{r}",fmt=MM,border=True); ws7[f"{cl}{r}"]=f"={cl}{ROW_OPEN}*経費前提!$B${R_DEF}"
ROW_DEF=r; r+=1
S(ws7,f"A{r}",bold=True,border=True); ws7[f"A{r}"]="開業時 現金支出"
for cl in YC:
    S(ws7,f"{cl}{r}",bold=True,fmt=MM,border=True,fill=TOT); ws7[f"{cl}{r}"]=f"={cl}{ROW_INV}+{cl}{ROW_DEF}"
ROW_CASH=r; r+=1
S(ws7,f"A{r}",border=True); ws7[f"A{r}"]="分割弁済額"
for cl in YC:
    S(ws7,f"{cl}{r}",fmt=MM,border=True)
    terms=[(f"IF(店舗群別PL!{cl}{b['y']}=0,0,経費前提!$B${R_PAY}*店舗群別PL!$B${b['n']}"
            f"*IF(店舗群別PL!{cl}{b['y']}=1,店舗群別PL!$B${b['m']}/12,1))") for b in blocks]
    ws7[f"{cl}{r}"]="=-("+"+".join(terms)+")"
ROW_PAY=r; r+=1
S(ws7,f"A{r}",bold=True,border=True); ws7[f"A{r}"]="投資キャッシュアウト　計"
for cl in YC:
    S(ws7,f"{cl}{r}",bold=True,fmt=MM,border=True,fill=TOT); ws7[f"{cl}{r}"]=f"={cl}{ROW_CASH}+{cl}{ROW_PAY}"
ROW_ICF=r; r+=1
S(ws7,f"A{r}",bold=True,border=True); ws7[f"A{r}"]="連結EBITDA"
for cl in YC:
    S(ws7,f"{cl}{r}",bold=True,color=GREEN,fmt=MM,border=True); ws7[f"{cl}{r}"]=f"=連結PL!{cl}{R4['eb']}"
ROW_EB=r; r+=1
ROW_TAXRATE=r+5
S(ws7,f"A{r}",border=True); ws7[f"A{r}"]="法人税等（推計）"
for cl in YC:
    S(ws7,f"{cl}{r}",fmt=MM,border=True)
    ws7[f"{cl}{r}"]=f"=-MAX(0,連結PL!{cl}{R4['op']})*$B${ROW_TAXRATE}"
ROW_TAX=r; r+=1
S(ws7,f"A{r}",bold=True,border=True); ws7[f"A{r}"]="簡易フリーキャッシュフロー"
for cl in YC:
    S(ws7,f"{cl}{r}",bold=True,fmt=MM,border=True,fill=KEY)
    ws7[f"{cl}{r}"]=f"={cl}{ROW_EB}+{cl}{ROW_TAX}+{cl}{ROW_ICF}"
ROW_FCF=r; r+=1
S(ws7,f"A{r}",bold=True,border=True); ws7[f"A{r}"]="累計フリーキャッシュフロー"
for i,cl in enumerate(YC):
    S(ws7,f"{cl}{r}",bold=True,fmt=MM,border=True,fill=KEY)
    ws7[f"{cl}{r}"]=f"={cl}{ROW_FCF}" if i==0 else f"={YC[i-1]}{r}+{cl}{ROW_FCF}"
ROW_CUM=r
for rr in [ROW_OPEN,ROW_INV,ROW_DEF,ROW_CASH,ROW_PAY,ROW_ICF,ROW_EB,ROW_TAX,ROW_FCF]:
    S(ws7,f"H{rr}",bold=True,fmt=MM0 if rr==ROW_OPEN else MM,border=True,fill=TOT)
    ws7[f"H{rr}"]=f"=SUM(C{rr}:G{rr})"
r=ROW_CUM+2
sec(ws7,r,"■ 必要資金と調達","H"); r+=1
assert r==ROW_TAXRATE, f"tax rate row mismatch {r} vs {ROW_TAXRATE}"
for i,(lab,f_,fmt,note) in enumerate([
    ("法人税 実効税率",0.30,PCT,"簡易。繰越欠損金の控除は考慮せず"),
    ("最大資金需要（累計FCFの最小値）",f"=-MIN(C{ROW_CUM}:G{ROW_CUM})",MM,"税引後ベース"),
    ("既存調達額（エクイティ＋デット）",200.0,MM,"事業収支資料p11「2億調達済み」"),
    ("第三者割当①（2026年10月）",150.0,MM,"Pre3億・Post4.5億"),
    ("調達額　計",None,MM,""),
    ("過不足（マイナスが不足）",None,MM,"デット元本返済・分割払い金利は未考慮")]):
    rr=r+i
    S(ws7,f"A{rr}",border=True,bold=(i>=4)); ws7[f"A{rr}"]=lab
    isin=(f_ is not None) and not isinstance(f_,str)
    S(ws7,f"B{rr}",color=BLUE if isin else BLACK,bold=(i>=4),fmt=fmt,border=True,align="center",
      fill=INP if isin else (KEY if i==5 else (TOT if i==4 else None)))
    if i==4: ws7[f"B{rr}"]=f"=B{r+2}+B{r+3}"
    elif i==5: ws7[f"B{rr}"]=f"=B{r+4}-B{r+1}"
    else: ws7[f"B{rr}"]=f_
    S(ws7,f"J{rr}",size=9,color="7F7F7F"); ws7[f"J{rr}"]=note
R_NEED,R_RAISE,R_GAP=r+1,r+4,r+5
S(ws7,f"A{r+7}",size=9,color="7F7F7F")
ws7[f"A{r+7}"]="※ 簡易FCF＝連結EBITDA－法人税等－投資キャッシュアウト。運転資本増減・借入返済は含まない。"
print("連結PL/投資CF OK")

# ══════════ 8. 感応度分析 ══════════
ws8=wb.create_sheet("感応度分析"); ws8.sheet_view.showGridLines=False
ws8.column_dimensions["A"].width=26
for c in "BCDEFG": ws8.column_dimensions[c].width=13
ws8.column_dimensions["H"].width=3; ws8.column_dimensions["I"].width=50
band(ws8,1,"感応度分析　店内席効率 × ディナー店内回転率（1号店・満年度）","G")
S(ws8,"A2",size=9,color="7F7F7F")
ws8["A2"]="原価は「経費前提」で選択中の方法に連動。軸の値（水色）は自由に変更できます。"
EFFS=[0.65,0.70,0.75,0.80,0.85]; TURNS=[0.70,0.80,0.90,1.00,1.10]
L_IN,C_IN,D_IN=R_G0,R_G0+2,R_G0+4
L_TE,C_TE,D_TE=R_G0+1,R_G0+3,R_G0+5
IN_FIX=(f"売上前提!$C${L_IN}*{{x}}*売上前提!$E${L_IN}*売上前提!$H${L_IN}"
        f"+売上前提!$C${C_IN}*{{x}}*売上前提!$E${C_IN}*売上前提!$H${C_IN}"
        f"+売上前提!$C${D_IN}*{{x}}*{{y}}*売上前提!$H${D_IN}")
TE_FIX=f"(売上前提!$J${L_TE}+売上前提!$J${C_TE}+売上前提!$J${D_TE})"
VSUM=f"経費前提!$B${R_VSUM}"
def sx(xc,yc): return f"(({IN_FIX.format(x=xc,y=yc)})+{TE_FIX})*売上前提!$B${R_DAYS}/1000000"
def grid(top,title,metric):
    sec(ws8,top,title,"G"); hdr=top+1
    S(ws8,f"A{hdr}",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True)
    ws8[f"A{hdr}"]="店内席効率 ＼ ディナー回転率"; ws8.row_dimensions[hdr].height=28
    for j,tv in enumerate(TURNS):
        c=chr(ord("B")+j)
        S(ws8,f"{c}{hdr}",bold=True,border=True,fill=INP,align="center",fmt=NUM2,color=BLUE)
        ws8[f"{c}{hdr}"]=tv
    for i,ev in enumerate(EFFS):
        rr=hdr+1+i
        S(ws8,f"A{rr}",bold=True,border=True,fill=INP,align="center",fmt=PCT,color=BLUE); ws8[f"A{rr}"]=ev
        for j in range(len(TURNS)):
            c=chr(ord("B")+j); s=sx(f"$A{rr}",f"{c}${hdr}")
            prof=(f"(({s})*(1-{VSUM})-経費前提!$B${R_RENT_Y}"
                  f"-MAX(0,({s})-経費前提!$B${R_THR})*経費前提!$B${R_PCTR}-経費前提!$B${R_DEPRE})")
            if metric=="sales": f_=f"={s}"; fmt=MM
            elif metric=="op": f_=f"={prof}"; fmt=MM
            else: f_=f"=IF(({s})=0,0,{prof}/({s}))"; fmt=PCT
            S(ws8,f"{c}{rr}",fmt=fmt,border=True,align="center"); ws8[f"{c}{rr}"]=f_
    ws8.conditional_formatting.add(f"B{hdr+1}:F{hdr+len(EFFS)}", FormulaRule(
        formula=[f'AND(ROUND($A{hdr+1},4)=ROUND(売上前提!$D${L_IN},4),'
                 f'ROUND(B${hdr},4)=ROUND(売上前提!$E${D_IN},4))'], fill=HIT, stopIfTrue=False))
    return hdr+len(EFFS)+1
nxt=grid(4,"■ 年間売上高（百万円）","sales")
nxt=grid(nxt+1,"■ 営業利益（百万円）","op")
nxt=grid(nxt+1,"■ 営業利益率","opm")
S(ws8,f"A{nxt+1}",size=9,color=RED)
ws8[f"A{nxt+1}"]="薄黄色＝現在の前提に一致するセル（席効率・回転率を変えると自動で移動します）。"

# ══════════ 1. サマリー ══════════
ws1=wb.create_sheet("サマリー",0); ws1.sheet_view.showGridLines=False
ws1.column_dimensions["A"].width=32; ws1.column_dimensions["B"].width=16
for c in "CDEFG": ws1.column_dimensions[c].width=14
ws1.column_dimensions["H"].width=14
band(ws1,1,"BRASSERIE Thierry Marx　事業収支 統合モデル（2026年10月）","H")
S(ws1,"A2",size=9,color="7F7F7F")
ws1["A2"]="単位：百万円。「売上前提」の席数・席効率・回転率・客単価を変えると、単店PLから5か年連結・資金計画まで一気に再計算されます。"
sec(ws1,4,"■ よく変える前提（編集は「売上前提」シートで）","H")
for i,(lab,f_,fmt) in enumerate([
    ("総席数（店内＋テラス）",f"=売上前提!$B${R_ALL_SEATS}",MM0),
    ("　店内",f"=売上前提!$B${R_IN_SEATS}",MM0),("　テラス",f"=売上前提!$B${R_TE_SEATS}",MM0),
    ("店内 席効率",f"=売上前提!$D${R_G0}",PCT),
    ("ディナー 回転率（店内）",f"=売上前提!$E${R_G0+4}",NUM2),
    ("客単価 ランチ",f"=売上前提!$H${R_G0}",YEN),
    ("客単価 カフェ",f"=売上前提!$H${R_G0+2}",YEN),
    ("客単価 ディナー（店内）",f"=売上前提!$H${R_G0+4}",YEN),
    ("客単価 ディナー（テラス）",f"=売上前提!$H${R_G0+5}",YEN),
    ("営業日数（日／年）",f"=売上前提!$B${R_DAYS}",NUM),
    ("標準店 満年度 年商",f"=売上前提!$K${R_GT}",MM)]):
    r=5+i; sub=lab.startswith("　")
    S(ws1,f"A{r}",border=True,size=9 if sub else 10,color="7F7F7F" if sub else BLACK,
      bold=("標準店" in lab)); ws1[f"A{r}"]=lab
    S(ws1,f"B{r}",color=GREEN,fmt=fmt,border=True,align="center",bold=("標準店" in lab),
      fill=KEY if "標準店" in lab else None)
    ws1[f"B{r}"]=f_
sec(ws1,17,"■ 単店PL（1号店）","H")
for c,t in zip("BCDE",["満年度","1年目","2年目","3年目以降"]):
    S(ws1,f"{c}18",bold=True,border=True,fill=TOT,align="center"); ws1[f"{c}18"]=t
S(ws1,"A18",bold=True,border=True,fill=TOT); ws1["A18"]="項　目"
r=19
for lab,src,hl in [("売上高",R_REV,False),("GOP（店舗営業総利益）",R_GOP,False),
                   ("EBITDA",R_EBITDA,True),("減価償却費",R_DEP,False),("営業利益",R_OP,True)]:
    S(ws1,f"A{r}",bold=hl,border=True); ws1[f"A{r}"]=lab
    for c in "BCDE":
        S(ws1,f"{c}{r}",bold=hl,color=GREEN,fmt=MM,border=True,align="center",fill=KEY if hl else None)
        ws1[f"{c}{r}"]=f"=単店PL!{c}{src}"
    r+=1
for lab,src in [("　EBITDA率",R_EBITDA),("　営業利益率",R_OP)]:
    S(ws1,f"A{r}",bold=True,border=True); ws1[f"A{r}"]=lab
    for c in "BCDE":
        S(ws1,f"{c}{r}",bold=True,fmt=PCT,border=True,align="center",fill=KEY)
        ws1[f"{c}{r}"]=f"=IF(単店PL!{c}{R_REV}=0,0,単店PL!{c}{src}/単店PL!{c}{R_REV})"
    r+=1
r+=1
sec(ws1,r,"■ 5か年連結","H"); r+=1
S(ws1,f"A{r}",bold=True,border=True,fill=TOT); ws1[f"A{r}"]="項　目"
for i,y in enumerate(YEARS):
    S(ws1,f"{YC[i]}{r}",bold=True,border=True,fill=TOT,align="center"); ws1[f"{YC[i]}{r}"]=y
S(ws1,f"H{r}",bold=True,border=True,fill=TOT,align="center"); ws1[f"H{r}"]="5年累計"
r+=1
for lab,key,fmt,hl in [("期末店舗数（店）","stores",MM0,False),("売上高","rev",MM,False),
                       ("連結EBITDA","eb",MM,True),("　EBITDA率","ebm",PCT,True),
                       ("営業利益","op",MM,True),("　営業利益率","opm",PCT,True)]:
    S(ws1,f"A{r}",bold=hl,border=True); ws1[f"A{r}"]=lab
    for cl in YC:
        S(ws1,f"{cl}{r}",bold=hl,color=GREEN,fmt=fmt,border=True,
          align="center" if fmt in (PCT,MM0) else None,fill=KEY if hl else None)
        ws1[f"{cl}{r}"]=f"=連結PL!{cl}{R4[key]}"
    if key!="stores":
        S(ws1,f"H{r}",bold=True,color=GREEN,fmt=fmt,border=True,fill=TOT,
          align="center" if fmt==PCT else None)
        ws1[f"H{r}"]=f"=連結PL!H{R4[key]}"
    r+=1
r+=1
sec(ws1,r,"■ 目標到達時期（自動判定）","H"); r+=1
for lab,key,thr,old in [("連結EBITDA率 10%超","ebm",0.1,"旧計画：FY2028"),
                        ("営業利益率 5%超","opm",0.05,"旧計画：FY2029")]:
    S(ws1,f"A{r}",border=True); ws1[f"A{r}"]=lab
    S(ws1,f"B{r}",bold=True,border=True,align="center",fill=KEY)
    ws1[f"B{r}"]=(f'=IF(COUNTIF(連結PL!C{R4[key]}:G{R4[key]},">={thr}")=0,"5年内に未達",'
                  f'INDEX({{"FY2027","FY2028","FY2029","FY2030","FY2031"}},'
                  f'6-COUNTIF(連結PL!C{R4[key]}:G{R4[key]},">={thr}")))')
    S(ws1,f"D{r}",size=9,color="7F7F7F"); ws1[f"D{r}"]=old
    r+=1
r+=1
sec(ws1,r,"■ 投資回収・資金","H"); r+=1
for lab,ref,fmt in [("1号店 投資回収","=投資・CF!B13",None),
                    ("1号店 投資回収年数（年）","=投資・CF!B14",NUM2),
                    ("初期投資（5年累計）",f"=投資・CF!H{ROW_INV}",MM),
                    ("開業時 現金支出（5年累計）",f"=投資・CF!H{ROW_CASH}",MM),
                    ("最大資金需要（税引後）",f"=投資・CF!B{R_NEED}",MM),
                    ("調達額 計（既存2億＋今回1.5億）",f"=投資・CF!B{R_RAISE}",MM),
                    ("過不足（マイナスが不足）",f"=投資・CF!B{R_GAP}",MM)]:
    S(ws1,f"A{r}",border=True,bold=("過不足" in lab)); ws1[f"A{r}"]=lab
    S(ws1,f"B{r}",color=GREEN,bold=("過不足" in lab),fmt=fmt,border=True,align="center",
      fill=KEY if ("過不足" in lab or "投資回収" == lab[-4:]) else None)
    ws1[f"B{r}"]=ref
    r+=1

if "Sheet" in wb.sheetnames: del wb["Sheet"]
SHEETS=["サマリー","売上前提","経費前提","単店PL","店舗群別PL","連結PL","投資・CF","感応度分析"]
for sh in wb.worksheets:
    for row_ in sh.iter_rows():
        for c in row_:
            if isinstance(c.value,str) and c.value.startswith("="):
                f=c.value
                for nm in SHEETS: f=f.replace(f"'{nm}'!",f"{nm}!")
                for nm in SHEETS: f=f.replace(f"{nm}!",f"'{nm}'!")
                c.value=f
wb.calculation.fullCalcOnLoad=True
wb.save(OUT)
print("saved:",OUT)
