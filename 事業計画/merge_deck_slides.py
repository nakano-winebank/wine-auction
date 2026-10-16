# -*- coding: utf-8 -*-
"""投資家向けデック（THIERRY MARX × WineBank）の事業収支スライドを、
統合モデルの最新値で元資料のトンマナ（Meiryo・銅色バー・ベージュ表）のまま差し替える。

使い方:
  python3 merge_deck_slides.py <統合モデル.xlsx> <展開済みデックdir>
  展開済みデックの slide12 / slide13 を書き換え、slide24 / slide25（slide13を複製済み）を新規作成。
"""
import sys, re
from decimal import Decimal, ROUND_HALF_UP
from xml.sax.saxutils import escape
import openpyxl

MODEL, DECK = sys.argv[1], sys.argv[2]
wv = openpyxl.load_workbook(MODEL, data_only=True)

EMU = 914400
COPPER, TAN, BEIGE, WHITE, BLACK, GRAY = "C68A65", "D5AA8F", "E3C5B2", "FFFFFF", "000000", "888888"
FONT = '<a:latin typeface="Meiryo"/><a:ea typeface="Meiryo"/><a:cs typeface="Meiryo"/><a:sym typeface="Meiryo"/>'
END = ('<a:endParaRPr sz="1400" b="0" i="0" u="none" strike="noStrike" cap="none"><a:solidFill>'
       '<a:srgbClr val="000000"/></a:solidFill><a:latin typeface="Arial"/><a:ea typeface="Arial"/>'
       '<a:cs typeface="Arial"/><a:sym typeface="Arial"/></a:endParaRPr>')


# ───────── 数値整形 ─────────
def hu(v, nd):
    q = Decimal(1).scaleb(-nd)
    return Decimal(str(round(v, 9))).quantize(q, rounding=ROUND_HALF_UP)

def mm(v, nd=0, plus=False):
    """百万円。マイナスは▲。"""
    d = hu(v, nd)
    if d == 0: d = abs(d)
    s = f"{abs(d):,.{nd}f}"
    return ("▲" + s) if d < 0 else (("＋" + s) if plus and d > 0 else s)

def man(v, plus=False):
    """百万円 → 万円（整数）。"""
    return mm(v * 100, 0, plus)

def pct(v, nd=1):
    return f"{hu(v * 100, nd):.{nd}f}%"


# ───────── モデルの値 ─────────
def rows_by_label(sh, c0=1, c1=8):
    d = {}
    for row in wv[sh].iter_rows():
        lab = row[0].value
        if isinstance(lab, str):
            d.setdefault(lab.strip().replace("　", " "), [c.value for c in row[c0:c1]])
    return d

U = rows_by_label("単店PL", 1, 5)       # 満年度,1年目,2年目,3年目以降
C = rows_by_label("連結PL", 2, 8)       # FY2027..FY2031, 5年累計
I = rows_by_label("投資・CF", 1, 8)
K = wv["経費前提"]; S = wv["売上前提"]; SEN = wv["感応度分析"]

def u(lab): return U[lab]
def c(lab): return C[lab][:5]

seats = {k: S[f"B{r}"].value for k, r in
         [("cnt", 13), ("din", 14), ("pri", 15), ("in", 16), ("ter", 17), ("all", 18)]}
day_sales = S["B36"].value; mon_sales = S["B37"].value; yr_sales = S["B38"].value
turn_all = S["B41"].value
eff_in, eff_te = S["D22"].value, S["D23"].value
lunch, cafe, din_in, din_te = S["H22"].value, S["H24"].value, S["H26"].value, S["H27"].value
din_food, din_bev = S["F26"].value, S["G26"].value

# 初期投資（1店あたり） 経費前提 B/C/D 列
inv_rows = {}
for row in K.iter_rows(min_row=30, max_row=50):
    lab = row[0].value
    if isinstance(lab, str):
        inv_rows[lab.strip().replace("　", " ")] = [x.value for x in row[1:4]]

def inv(lab_prefix):
    for k, v in inv_rows.items():
        if k.startswith(lab_prefix): return v
    raise KeyError(lab_prefix)


# ───────── XML 部品 ─────────
_id = [5000]
def nid():
    _id[0] += 1; return _id[0]

def E(x): return int(round(x * EMU))

def run(text, sz, color, b=False):
    return (f'<a:r><a:rPr lang="ja-JP" sz="{sz}" b="{1 if b else 0}" i="0" u="none" strike="noStrike" cap="none">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>{FONT}</a:rPr><a:t>{escape(text)}</a:t></a:r>')

def para(runs, algn="l", ln=115000, bu_sz=700):
    return (f'<a:p><a:pPr marL="0" marR="0" lvl="0" indent="0" algn="{algn}" rtl="0"><a:lnSpc><a:spcPct val="{ln}"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft>'
            f'<a:buClr><a:srgbClr val="000000"/></a:buClr><a:buSzPts val="{bu_sz}"/><a:buFont typeface="Arial"/><a:buNone/></a:pPr>'
            f'{runs}{END}</a:p>')

def textbox(x, y, w, h, paras, autofit=True, anchor="t"):
    i = nid()
    fit = "<a:spAutoFit/>" if autofit else "<a:noAutofit/>"
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="TextBox {i}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(h)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>'
            f'<p:txBody><a:bodyPr spcFirstLastPara="1" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}" anchorCtr="0">{fit}</a:bodyPr>'
            f'<a:lstStyle/>{"".join(paras)}</p:txBody></p:sp>')

def title(t):   return textbox(0.26, 0.23, 8.0, 0.31, [para(run(t, 1400, BLACK, True), ln=115000, bu_sz=1400)])
def subtitle(t): return textbox(0.28, 0.56, 9.2, 0.18, [para(run(t, 750, GRAY), bu_sz=750)])
def footnote(t, y=5.10): return textbox(0.48, y, 9.0, 0.16, [para(run(t, 600, GRAY), bu_sz=600)])
def confidential(): return textbox(0.06, 5.33, 0.9, 0.14, [para(run("CONFIDENTIAL", 600, GRAY), bu_sz=600)])

def slidenum():
    i = nid()
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="Slide Number {i}"/><p:cNvSpPr txBox="1"><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph type="sldNum" idx="12"/></p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="8898878" y="4847099"/>'
            f'<a:ext cx="264900" cy="183600"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>'
            f'<p:txBody><a:bodyPr spcFirstLastPara="1" wrap="square" lIns="68575" tIns="34275" rIns="68575" bIns="34275" anchor="ctr" anchorCtr="0"><a:noAutofit/></a:bodyPr>'
            f'<a:lstStyle/><a:p><a:pPr marL="0" lvl="0" indent="0" algn="r" rtl="0"><a:lnSpc><a:spcPct val="100000"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft><a:buClr><a:srgbClr val="000000"/></a:buClr>'
            f'<a:buSzPts val="900"/><a:buFont typeface="Arial"/><a:buNone/></a:pPr><a:fld id="{{00000000-1234-1234-1234-123412341234}}" type="slidenum">'
            f'<a:rPr lang="en-US" altLang="ja-JP" sz="700"/><a:t>‹#›</a:t></a:fld><a:endParaRPr sz="700"/></a:p></p:txBody></p:sp>')

def bar(x, y, w, t, h=0.24):
    i = nid()
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="Header {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(h)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{COPPER}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>'
            f'<p:txBody><a:bodyPr spcFirstLastPara="1" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr" anchorCtr="0"><a:noAutofit/></a:bodyPr>'
            f'<a:lstStyle/>{para(run(t, 800, WHITE, True), algn="ctr", ln=100000, bu_sz=800)}</p:txBody></p:sp>')

def bullets(x, y, w, h, lines, sz=750, ln=145000):
    ps = []
    for ln_ in lines:
        if isinstance(ln_, tuple):           # (text, color, bold)
            t, col, b = ln_
        else:
            t, col, b = ln_, BLACK, False
        ps.append(para(run(t, sz, col, b), ln=ln, bu_sz=sz))
    return textbox(x, y, w, h, ps)

# kind: h=見出し行, b=太字行(ベージュ), n=通常行(ラベル茶・値白), k=強調行(銅色・白字)
def cell(text, kind, first, algn, sz=700, color=None):
    if kind == "h":   fill, fc, b = COPPER, WHITE, True
    elif kind == "b": fill, fc, b = BEIGE, BLACK, True
    elif kind == "k": fill, fc, b = COPPER, WHITE, True
    elif kind == "s": fill, fc, b = (TAN, WHITE, False) if first else (WHITE, BLACK, False)  # 小見出し的な通常行
    else:             fill, fc, b = (TAN, WHITE, False) if first else (WHITE, BLACK, False)
    if color: fc = color
    return (f'<a:tc><a:txBody><a:bodyPr/><a:lstStyle/>{para(run(text, sz, fc, b), algn=algn, ln=100000, bu_sz=sz)}</a:txBody>'
            f'<a:tcPr marL="45725" marR="45725" marT="0" marB="0" anchor="ctr"><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill></a:tcPr></a:tc>')

def table(x, y, colw, rows, rh=0.213, sz=700, first_algn="l"):
    """rows: [(kind, [cells...]) or (kind, [cells...], {col: color})]"""
    i = nid()
    grid = "".join(
        f'<a:gridCol w="{E(cw)}"><a:extLst><a:ext uri="{{9D8B030D-6E8A-4147-A177-3AD203B41FA5}}">'
        f'<a16:colId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" val="{20000 + j}"/></a:ext></a:extLst></a:gridCol>'
        for j, cw in enumerate(colw))
    trs = []
    for ri, r in enumerate(rows):
        kind, vals = r[0], r[1]
        colors = r[2] if len(r) > 2 else {}
        tcs = "".join(cell(str(v), kind, j == 0,
                           (first_algn if j == 0 else ("ctr" if kind == "h" else "r")), sz, colors.get(j))
                      for j, v in enumerate(vals))
        trs.append(f'<a:tr h="{E(rh)}">{tcs}<a:extLst><a:ext uri="{{0D108BD9-81ED-4DB2-BD59-A6C34878D82A}}">'
                   f'<a16:rowId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" val="{10000 + ri}"/></a:ext></a:extLst></a:tr>')
    W, H = sum(colw), rh * len(rows)
    return (f'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="{i}" name="Table {i}"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
            f'<p:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(W)}" cy="{E(H)}"/></p:xfrm>'
            f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table"><a:tbl>'
            f'<a:tblPr><a:noFill/><a:tableStyleId>{{23FF2375-F9B8-41A6-9E25-2D63B5D407DD}}</a:tableStyleId></a:tblPr>'
            f'<a:tblGrid>{grid}</a:tblGrid>{"".join(trs)}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')

def write_slide(n, shapes):
    p = f"{DECK}/ppt/slides/slide{n}.xml"
    s = open(p, encoding="utf8").read()
    a = s.index("</p:grpSpPr>") + len("</p:grpSpPr>")
    b = s.rindex("</p:spTree>")
    s = s[:a] + "".join(shapes) + s[b:]
    open(p, "w", encoding="utf8").write(s)
    print(f"slide{n}.xml written ({len(shapes)} shapes)")


# ═════════ slide12：六本木ミッドタウン店 単体事業収支 ═════════
sales = u("売上高")[0]
def pl_row(lab, key=None, kind="n"):
    v = u(key or lab)[0]
    return (kind, [lab, mm(v), pct(v / sales)])

left = [("h", ["項目", "金額", "対売上比"]),
        pl_row("売上高", kind="b"), pl_row("売上原価"), pl_row("人件費"), pl_row("業務委託費・消耗品他"),
        pl_row("部門利益", kind="b"), pl_row("広告宣伝費・販売促進費"), pl_row("その他経費"), pl_row("水道光熱費"),
        pl_row("GOP（店舗営業総利益）", kind="b"), pl_row("ロイヤリティ"), pl_row("保険・租税公課"),
        pl_row("地代家賃（固定）"), pl_row("歩合家賃"), pl_row("修繕維持費（FFE）"),
        pl_row("EBITDA", kind="b"), pl_row("減価償却費"), pl_row("営業利益", kind="b")]
yr = lambda lab: u(lab)[1:4]
eb, op, sl = yr("EBITDA"), yr("営業利益"), yr("売上高")
right = [("h", ["項目", "1年目", "2年目", "3年目以降"]),
         ("n", ["売上立上り率"] + [pct(x, 0) for x in u("立上り率")[1:4]]),
         ("n", ["売上高"] + [mm(x) for x in sl]),
         ("b", ["EBITDA"] + [mm(x) for x in eb]),
         ("b", ["EBITDA率"] + [pct(a / b) for a, b in zip(eb, sl)]),
         ("b", ["営業利益"] + [mm(x) for x in op]),
         ("b", ["営業利益率"] + [pct(a / b) for a, b in zip(op, sl)])]
lab_ratio = u("人件費")
pts12 = [
    f"● 席数{seats['all']}席（カウンター{seats['cnt']}・ダイニング{seats['din']}・個室{seats['pri']}・テラス{seats['ter']}）",
    f"● 客単価　ランチ {lunch:,.0f}円／カフェ {cafe:,.0f}円",
    f"　　　　　ディナー{din_in:,.0f}円（料理{din_food:,.0f}＋飲料{din_bev:,.0f}）・テラス{din_te:,.0f}円",
    f"● 席効率 店内{pct(eff_in, 0)}・テラス{pct(eff_te, 0)}／延べ席回転率 {turn_all:.2f}回/日",
    f"● 日商 約{day_sales / 10000:,.0f}万円／月商 約{mon_sales * 100:,.0f}万円／年商 {yr_sales / 100:.2f}億円",
    f"● 人員26名。売上原価{pct(u('売上原価')[0] / sales)}・人件費{pct(lab_ratio[0] / sales)}",
    "● ロイヤリティ4.0%・歩合家賃7%（閾値超過分）・地代家賃 月222万円",
    f"● 人件費は2年目{pct(lab_ratio[2] / u('売上高')[2])}→3年目{pct(lab_ratio[3] / u('売上高')[3])}へ段階改善（本部集中・多能工化）",
    ("● 内装1億円は家主立替・10年分割返済（資産計上・減価償却は自社）", "C00000", False),
]
write_slide(12, [
    title("六本木ミッドタウン店　単体事業収支"),
    subtitle(f"2026年10月 事業収支統合モデル準拠　｜　183.59㎡(55.5坪)・{seats['all']}席・年間363日営業"),
    bar(0.48, 0.86, 4.35, "満年度ベース 損益計算（百万円）"),
    table(0.48, 1.16, [2.15, 1.1, 1.1], left, rh=0.213),
    bar(5.27, 0.86, 4.25, "立上りを織り込んだ年次推移（百万円）"),
    table(5.27, 1.16, [1.45, 0.93, 0.93, 0.94], right, rh=0.24),
    bar(5.27, 3.00, 4.25, "主要前提"),
    bullets(5.35, 3.32, 4.15, 1.70, pts12, sz=700, ln=135000),
    footnote(f"※ 満年度ベースは店内席効率{pct(eff_in, 0)}・ディナー店内回転率{S['E26'].value:.2f}回の標準値。新店初年度は保守的に売上90%で計画。金利・税金は含まない。"),
    confidential(), slidenum()])


# ═════════ slide13：5か年事業計画 事業収支（連結） ═════════
def crow(lab, keys, kind="n"):
    vals = [sum(C[k][i] for k in keys) for i in range(5)]
    return (kind, [lab] + [mm(v) for v in vals])
sales5 = c("売上高")
ebr, opr = c("連結EBITDA"), c("営業利益")
cons = [("h", ["項目", "FY2027", "FY2028", "FY2029", "FY2030", "FY2031"]),
        ("b", ["期末店舗数"] + [str(x) for x in c("期末店舗数（店）")]),
        crow("売上高", ["売上高"], "b"), crow("売上原価", ["売上原価"]), crow("人件費", ["人件費"]),
        crow("業務委託費・消耗品他", ["業務委託費・消耗品他"]), crow("部門利益", ["部門利益"], "b"),
        crow("販促費・水道光熱費等", ["広告宣伝費・販売促進費", "その他経費", "水道光熱費"]),
        crow("GOP（店舗営業総利益）", ["GOP（店舗営業総利益）"], "b"),
        crow("ロイヤリティ", ["ロイヤリティ"]),
        crow("地代家賃・歩合家賃", ["地代家賃（固定）", "歩合家賃"]),
        crow("保険・租税公課・修繕費", ["保険・租税公課", "修繕維持費（FFE）"]),
        crow("店舗EBITDA計", ["店舗EBITDA 計"]), crow("本部費", ["本部費"]),
        crow("連結EBITDA", ["連結EBITDA"], "b"),
        ("b", ["EBITDA率"] + [pct(a / b) for a, b in zip(ebr, sales5)]),
        crow("減価償却費", ["減価償却費"]), crow("営業利益", ["営業利益"], "b"),
        ("b", ["営業利益率"] + [pct(a / b) for a, b in zip(opr, sales5)])]
FY = ["FY2027", "FY2028", "FY2029", "FY2030", "FY2031"]
def first_fy(series, th):
    for fy, a, b in zip(FY, series, sales5):
        if a / b > th: return fy, a / b
fy_e, r_e = first_fy(ebr, 0.10); fy_o, r_o = first_fy(opr, 0.05)
m_eb = u("EBITDA")[3] / u("売上高")[3]; m_op = u("営業利益")[3] / u("売上高")[3]
util = [u("水道光熱費")[i] / u("売上高")[i] for i in (0, 3)]
outs = [u("業務委託費・消耗品他")[i] / u("売上高")[i] for i in (0, 3)]
labr = [u("人件費")[i] / u("売上高")[i] for i in (0, 3)]
pts13 = [f"● 1号店はFY2027（開業初年度）から", "　 売上立上り90%で保守的に計上", "",
         f"● {fy_e}に連結EBITDA率10%超（{pct(r_e)}）", "",
         f"● {fy_o}に営業利益率5%超（{pct(r_o)}）", "",
         f"● FY2031はEBITDA率{pct(ebr[4] / sales5[4])}", f"　 営業利益率{pct(opr[4] / sales5[4])}", "",
         f"● 成熟単店はEBITDA率{pct(m_eb)}", f"　 営業利益率{pct(m_op)}"]
lev13 = [f"● 人件費 {pct(labr[0])} → {pct(labr[1])}", "　 本部集中化・多能工化",
         f"● 水道光熱 {pct(util[0])} → {pct(util[1])}", "　 設備運用の標準化",
         f"● 業務委託・消耗品 {pct(outs[0])} → {pct(outs[1])}"]
write_slide(13, [
    title("5か年事業計画　事業収支（連結）"),
    subtitle("2027年4月 六本木開業 → 2028年 札幌・大阪 → 2029年 首都圏・政令指定都市 → 2030年以降 リゾート含め年2店舗"),
    bar(0.48, 0.90, 6.15, "連結損益計画（百万円）"),
    table(0.48, 1.20, [1.95, 0.84, 0.84, 0.84, 0.84, 0.84], cons, rh=0.198),
    bar(6.90, 0.90, 2.62, "計画のポイント"),
    bullets(6.98, 1.22, 2.50, 2.40, pts13, sz=750, ln=120000),
    bar(6.90, 3.75, 2.62, "経費改善レバー"),
    bullets(6.98, 4.07, 2.50, 1.00, lev13, sz=750, ln=120000),
    footnote("※ FY＝4月〜3月。新店初年度は売上90%・開業7.5か月で保守的に計上。ロイヤリティ4%・歩合家賃7%・地代家賃は契約条件どおり全店適用。金利・税金は含まない。"),
    confidential(), slidenum()])


# ═════════ slide24：投資回収と資金計画 ═════════
contract, biz, prep, dep, wine = inv("契約金"), inv("事業費"), inv("開業準備金"), inv("保証金"), inv("ワイン")
total = [a + b + c_ + d + e for a, b, c_, d, e in zip(contract, biz, prep, dep, wine)]
def_in, def_ho = inv("分割対象：内装"), inv("分割対象：保証金")
split = [a + b for a, b in zip(def_in, def_ho)]
cash = [a - b for a, b in zip(total, split)]
f1 = lambda xs: [mm(x, 1) for x in xs]
invt = [("h", ["項目", "1号店", "2〜3号店", "4号店〜"]),
        ("n", ["契約金"] + f1(contract)),
        ("n", ["事業費（設計・内装施工・OSE/FF&E）"] + f1(biz)),
        ("n", ["　うち 内装 家主立替（10年分割返済）"] + f1([-x for x in def_in])),
        ("n", ["開業準備金（広告・採用等）"] + f1(prep)),
        ("n", ["保証金（非償却・返還対象）"] + f1(dep)),
        ("n", ["ワイン・美術品在庫（非償却）"] + f1(wine)),
        ("b", ["初期投資 合計"] + f1(total)),
        ("n", [f"分割対象額（内装{mm(def_in[0])}＋保証金{mm(def_ho[0], 1)}）"] + f1(split)),
        ("k", ["開業時 現金支出"] + f1(cash))]
cf = {k: I[k][:6] for k in ["EBITDA", "開業時 現金支出", "分割弁済", "単店FCF", "累計FCF"]}
dash = lambda v: "—" if abs(v) < 1e-9 else mm(v, 1)
cft = [("h", ["項目", "開業時", "1年目", "2年目", "3年目", "4年目", "5年目"]),
       ("n", ["EBITDA"] + [dash(x) for x in cf["EBITDA"]]),
       ("n", ["開業時 現金支出"] + [dash(x) for x in cf["開業時 現金支出"]]),
       ("n", ["分割弁済（内装＋保証金）"] + [dash(x) for x in cf["分割弁済"]]),
       ("b", ["単店FCF"] + [mm(x, 1) for x in cf["単店FCF"]]),
       ("b", ["累計FCF"] + [mm(x, 1) for x in cf["累計FCF"]])]
payback = I["投資回収年数（年）"][0]; turn_yr = I["累計FCFの黒字転換"][0]
ccf_keys = [("投資キャッシュアウト", "投資キャッシュアウト 計"), ("連結EBITDA", "連結EBITDA"),
            ("法人税等（推計）", "法人税等（推計）"), ("簡易フリーCF", "簡易フリーキャッシュフロー")]
ccf = [("h", ["項目"] + FY)]
for lab, key in ccf_keys:
    ccf.append(("n", [lab] + [mm(x, 1) for x in I[key][1:6]]))
ccf.append(("b", ["累計フリーCF"] + [mm(x, 1) for x in I["累計フリーキャッシュフロー"][1:6]]))
need_open = I["開業時に必要な現金 小計"][0]; need_fy = I["FY2027末までに必要な資金 累計"][0]
need_5 = I["最大資金需要（累計FCFの最小値）"][0]
closed = I["調達済（アピシウス70.0＋エニーカラー20.0）"][0]; pending = I["募集中（今回ラウンド残額）"][0]
raise_t = I["調達額 計（ラウンド完了時）"][0]
gap_now = I["過不足（開業時点・調達済のみ）"][0]; gap_open = I["過不足（開業時点・ラウンド完了後）"][0]
gap_fy = I["過不足（FY2027末・ラウンド完了後）"][0]; gap_5 = I["過不足（5年・9店舗展開ベース）"][0]
assert abs(closed - need_open - gap_now) < 1e-6 and abs(raise_t - need_open - gap_open) < 1e-6
RED = "C00000"
fund = [("h", ["項目", "金額（万円）"]),
        ("n", ["① 1号店 開業時に必要な現金", man(need_open)]),
        ("n", ["② FY2027末までに必要な資金", man(need_fy)]),
        ("n", ["③ 5年・9店舗の最大資金需要", man(need_5)]),
        ("n", ["調達済（アピシウス投資家7,000＋エニーカラー創業者2,000）", man(closed)]),
        ("n", ["募集中（今回ラウンド残額）", man(pending)]),
        ("b", ["調達額 計（ラウンド1.8億円 完了時）", man(raise_t)]),
        ("k", ["現時点の不足（① － 調達済）", man(gap_now)]),
        ("n", ["ラウンド完了後　①開業時", man(gap_open, plus=True)]),
        ("n", ["ラウンド完了後　②FY2027末", man(gap_fy, plus=True)]),
        ("b", ["ラウンド完了後　③5年・9店舗", man(gap_5)], {1: RED})]
write_slide(24, [
    title("投資回収と資金計画"),
    subtitle(f"開業時に必要な現金{man(need_open)}万円に対し調達済{man(closed)}万円 ― 不足{man(-gap_now)}万円は募集中の{man(pending)}万円で充当"),
    bar(0.48, 0.86, 4.40, "初期投資（1店あたり・百万円）"),
    table(0.48, 1.16, [2.05, 0.78, 0.78, 0.79], invt, rh=0.19, sz=650),
    bar(0.48, 3.20, 4.40, "1号店 キャッシュフロー（百万円）"),
    table(0.48, 3.50, [1.36, 0.50, 0.50, 0.51, 0.51, 0.51, 0.51], cft, rh=0.19, sz=650),
    bullets(0.48, 4.70, 4.40, 0.20, [(f"● 開業{turn_yr}に累計FCFが黒字転換　／　投資回収 約{payback:.1f}年　／　定常期FCF {mm(cf['単店FCF'][5], 1)}百万円", BLACK, True)], sz=700, ln=115000),
    bar(5.12, 0.86, 4.40, "連結キャッシュフロー（百万円）"),
    table(5.12, 1.16, [1.40, 0.60, 0.60, 0.60, 0.60, 0.60], ccf, rh=0.19, sz=650),
    bar(5.12, 2.45, 4.40, "必要資金と調達"),
    table(5.12, 2.75, [3.25, 1.15], fund, rh=0.19, sz=650),
    footnote("※ 内装100百万円は家主立替（月833千円×120回）、保証金のうち38.9百万円は月463千円×84回の分割（賃貸借契約）。資産計上・減価償却は自社で継続。"
             "既存2億円は調達済9,000万円と家主分割分で構成されるため調達額に含めない。運転資金は含まない。", y=4.98),
    confidential(), slidenum()])


# ═════════ slide25：前提の振れ幅に対する感応度 ═════════
turns = [SEN.cell(5, j).value for j in range(2, 7)]
effs = [SEN.cell(r, 1).value for r in range(6, 11)]
base_e, base_t = eff_in, S["E26"].value
def grid(r0, fmt):
    rows = [("h", ["席効率 ＼ 回転率"] + [f"{t:.2f}" for t in turns])]
    for i, e in enumerate(effs):
        vals = [SEN.cell(r0 + i, j).value for j in range(2, 7)]
        rows.append(("n", [pct(e, 0)] + [fmt(v) for v in vals]))
    return rows
g_sales = grid(6, lambda v: mm(v, 1))
g_opm = grid(22, lambda v: pct(v))
s_lo, s_hi = SEN.cell(6, 2).value, SEN.cell(10, 6).value
o_lo, o_hi = SEN.cell(22, 2).value, SEN.cell(26, 6).value

# 計画前提セルを銅色で強調するため、該当行を個別に組む
def grid_rows(rows):
    out = []
    bi = [i for i, e in enumerate(effs) if abs(e - base_e) < 1e-9]
    bj = [j for j, t in enumerate(turns) if abs(t - base_t) < 1e-9]
    for k, r in enumerate(rows):
        if k > 0 and bi and k - 1 == bi[0] and bj:
            out.append(("n", r[1], {bj[0] + 1: "HILITE"}))
        else:
            out.append(r)
    return out

def table_hl(x, y, colw, rows, rh, sz):
    xml = table(x, y, colw, [(r[0], r[1]) for r in rows], rh=rh, sz=sz, first_algn="ctr")
    # 強調セル：該当行・列のセルを銅色・白太字に置換
    for k, r in enumerate(rows):
        if len(r) > 2:
            col = list(r[2].keys())[0]
            trs = re.findall(r"<a:tr .*?</a:tr>", xml, flags=re.S)
            tr = trs[k]
            tcs = re.findall(r"<a:tc>.*?</a:tc>", tr, flags=re.S)
            tc = tcs[col]
            new = (tc.replace(f'<a:srgbClr val="{WHITE}"/></a:solidFill></a:tcPr>', f'<a:srgbClr val="{COPPER}"/></a:solidFill></a:tcPr>')
                     .replace(f'b="0" i="0" u="none" strike="noStrike" cap="none"><a:solidFill><a:srgbClr val="{BLACK}"/>',
                              f'b="1" i="0" u="none" strike="noStrike" cap="none"><a:solidFill><a:srgbClr val="{WHITE}"/>'))
            tr2 = tr.replace(tc, new, 1)
            xml = xml.replace(tr, tr2, 1)
    return xml

pts25 = [("● 下振れ耐性", COPPER, True),
         f"　 最も保守的な前提（席効率{pct(effs[0], 0)}・回転{turns[0]:.2f}）でも年商{s_lo / 100:.2f}億円・営業利益率{pct(o_lo)}で黒字を維持",
         ("● 上振れ余地", COPPER, True),
         f"　 席効率{pct(effs[-1], 0)}・回転{turns[-1]:.2f}では年商{s_hi / 100:.2f}億円・営業利益率{pct(o_hi)}。テラスの通年稼働が進めばさらに上振れ",
         ("● 固定費の効き", COPPER, True),
         "　 地代家賃と減価償却は売上に連動しないため、売上が伸びるほど利益率は加速的に改善"]
write_slide(25, [
    title("前提の振れ幅に対する感応度"),
    subtitle("1号店・満年度ベース　｜　店内席効率 × ディナー店内回転率（銅色＝計画前提）"),
    bar(0.48, 0.86, 4.35, "年間売上高（百万円）"),
    table_hl(0.48, 1.16, [1.10, 0.65, 0.65, 0.65, 0.65, 0.65], grid_rows(g_sales), 0.30, 750),
    bar(5.17, 0.86, 4.35, "営業利益率"),
    table_hl(5.17, 1.16, [1.10, 0.65, 0.65, 0.65, 0.65, 0.65], grid_rows(g_opm), 0.30, 750),
    bar(0.48, 3.25, 9.04, "読み取りポイント"),
    bullets(0.56, 3.60, 8.90, 1.40, pts25, sz=800, ln=150000),
    footnote(f"※ 計画前提は店内席効率{pct(base_e, 0)}・ディナー店内回転率{base_t:.2f}（銅色セル）。ランチ・カフェ・テラスの前提は固定。"),
    confidential(), slidenum()])
