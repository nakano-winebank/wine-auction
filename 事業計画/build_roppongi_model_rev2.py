# -*- coding: utf-8 -*-
"""六本木ミッドタウン店 事業収支（可変モデル）rev2
   2026年10月 ユーザー変更（席数・席効率・回転率・客単価・原価方法・ワイン/美術品）を反映し、
   ①償却対象の是正 ②投資内訳の整合 ③投資回収の自動計算 ④感応度の整合 を織り込む。"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule

OUT = "/home/user/wine-auction/事業計画/六本木ミッドタウン店_事業収支_可変モデル_202610_rev2.xlsx"
JP = "Meiryo"
BLUE, BLACK, GREEN, RED = "0000FF", "000000", "008000", "C00000"
HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E2F3")
TOT = PatternFill("solid", fgColor="F2F2F2")
KEY = PatternFill("solid", fgColor="FFFF00")
INP = PatternFill("solid", fgColor="EAF1FB")
HIT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

MM='#,##0.0;(#,##0.0);-'; YEN='#,##0;(#,##0);-'; NUM1='#,##0.0;(#,##0.0);-'
NUM2='0.00;(0.00);-'; PCT='0.0%;(0.0%);-'; PCT2='0.00%;(0.00%);-'; INT='#,##0;(#,##0);-'

wb = openpyxl.Workbook()
def S(ws, cell, *, bold=False, color=BLACK, fmt=None, fill=None, align=None,
      size=10, border=False, wrap=False):
    c = ws[cell]
    c.font = Font(name=JP, bold=bold, color=color, size=size)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align or wrap: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border: c.border = BOX
    return c
def band(ws, row, text, last, fill=HDR, color="FFFFFF", size=11):
    ws.cell(row=row, column=1, value=text)
    for col in range(1, openpyxl.utils.column_index_from_string(last)+1):
        c = ws.cell(row=row, column=col); c.fill = fill
        c.font = Font(name=JP, bold=True, color=color, size=size)
def sec(ws, row, text, last): band(ws, row, text, last, fill=SUB, color="000000", size=10)

# ═══════════════ 売上前提 ═══════════════
ws2 = wb.create_sheet("売上前提"); ws2.sheet_view.showGridLines = False
ws2.column_dimensions["A"].width = 22; ws2.column_dimensions["B"].width = 10
for c in "CDEFGHIJKL": ws2.column_dimensions[c].width = 12
ws2.column_dimensions["M"].width = 3; ws2.column_dimensions["N"].width = 48
band(ws2, 1, "売上前提　／　席数 × 席効率 × 回転率 × 客単価 からの積上げ", "L")
S(ws2,"A2",size=9,color="7F7F7F")
ws2["A2"]="水色セル＝入力値。ここを変えると全シートが自動再計算されます。単位：売上は百万円、客単価は円。"

sec(ws2, 4, "■ 基本条件", "L")
for i,(lab,v,f,note) in enumerate([
    ("営業日数（日／年）",363.0,NUM1,"年363日営業（提出P&L準拠）"),
    ("立上り率　1年目",0.90,PCT,"新店初年度は保守的に90%"),
    ("立上り率　2年目",0.98,PCT,""),
    ("立上り率　3年目以降（成熟）",1.00,PCT,"満年度＝この水準")]):
    r=5+i
    S(ws2,f"A{r}"); ws2[f"A{r}"]=lab
    S(ws2,f"B{r}",color=BLUE,fmt=f,border=True,align="center",fill=INP); ws2[f"B{r}"]=v
    S(ws2,f"N{r}",size=9,color="7F7F7F"); ws2[f"N{r}"]=note

sec(ws2, 10, "■ 席数（席）", "L")
for i,(lab,v) in enumerate([("カウンター",14),("ダイニング",24),("個室",25)]):
    r=11+i
    S(ws2,f"A{r}"); ws2[f"A{r}"]=lab
    S(ws2,f"B{r}",color=BLUE,fmt=INT,border=True,align="center",fill=INP); ws2[f"B{r}"]=v
S(ws2,"A14",bold=True); ws2["A14"]="店内　小計"
S(ws2,"B14",bold=True,fmt=INT,border=True,align="center",fill=TOT); ws2["B14"]="=SUM(B11:B13)"
S(ws2,"A15"); ws2["A15"]="テラス"
S(ws2,"B15",color=BLUE,fmt=INT,border=True,align="center",fill=INP); ws2["B15"]=52
S(ws2,"A16",bold=True); ws2["A16"]="合　計"
S(ws2,"B16",bold=True,fmt=INT,border=True,align="center",fill=KEY); ws2["B16"]="=B14+B15"
S(ws2,"N11",size=9,color="7F7F7F")
ws2["N11"]="2026年10月改定：ダイニング30→24、個室27→25（店内71→63席／合計123→115席）。面積は55.5坪で変更なし"

sec(ws2, 18, "■ 時間帯 × エリア 別の売上前提", "L")
for c,t in [("A","時間帯"),("B","エリア"),("C","席数"),("D","席効率"),("E","回転率\n(回/日)"),
            ("F","客単価\n料理(円)"),("G","客単価\n飲料(円)"),("H","客単価\n計(円)"),
            ("I","1日\n客数(人)"),("J","日商(円)"),("K","年間売上\n(百万円)"),("L","構成比")]:
    S(ws2,f"{c}19",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True); ws2[f"{c}19"]=t
ws2.row_dimensions[19].height=34
GRID=[("ランチ","店内","$B$14",0.75,1.15,3000,1000),
      ("ランチ","テラス","$B$15",0.50,1.00,3000,1000),
      ("カフェ（アイドルタイム）","店内","$B$14",0.75,0.50,1500,1000),
      ("カフェ（アイドルタイム）","テラス","$B$15",0.50,0.50,1500,1000),
      ("ディナー","店内","$B$14",0.75,0.90,9000,6000),
      ("ディナー","テラス","$B$15",0.50,0.50,5000,5000)]
R_G0=20
for i,(jk,ar,sref,eff,tn,ry,ir) in enumerate(GRID):
    r=R_G0+i
    S(ws2,f"A{r}",border=True,size=9); ws2[f"A{r}"]=jk
    S(ws2,f"B{r}",border=True,size=9,align="center"); ws2[f"B{r}"]=ar
    S(ws2,f"C{r}",color=GREEN,fmt=INT,border=True,align="center"); ws2[f"C{r}"]=f"={sref}"
    S(ws2,f"D{r}",color=BLUE,fmt=PCT,border=True,align="center",fill=INP); ws2[f"D{r}"]=eff
    S(ws2,f"E{r}",color=BLUE,fmt=NUM2,border=True,align="center",fill=INP); ws2[f"E{r}"]=tn
    S(ws2,f"F{r}",color=BLUE,fmt=YEN,border=True,align="center",fill=INP); ws2[f"F{r}"]=ry
    S(ws2,f"G{r}",color=BLUE,fmt=YEN,border=True,align="center",fill=INP); ws2[f"G{r}"]=ir
    S(ws2,f"H{r}",fmt=YEN,border=True,align="center"); ws2[f"H{r}"]=f"=F{r}+G{r}"
    S(ws2,f"I{r}",fmt=NUM1,border=True,align="center"); ws2[f"I{r}"]=f"=C{r}*D{r}*E{r}"
    S(ws2,f"J{r}",fmt=YEN,border=True); ws2[f"J{r}"]=f"=I{r}*H{r}"
    S(ws2,f"K{r}",fmt=MM,border=True); ws2[f"K{r}"]=f"=J{r}*$B$5/1000000"
    S(ws2,f"L{r}",fmt=PCT,border=True,align="center"); ws2[f"L{r}"]=f"=IF($K$26=0,0,K{r}/$K$26)"
R_G1, R_GT = R_G0+5, R_G0+6
S(ws2,f"A{R_GT}",bold=True,border=True); ws2[f"A{R_GT}"]="合　計"
for c in "BCDEFGH": S(ws2,f"{c}{R_GT}",border=True,fill=TOT)
for c,f_,fm,fl in [("I",f"=SUM(I{R_G0}:I{R_G1})",NUM1,TOT),("J",f"=SUM(J{R_G0}:J{R_G1})",YEN,TOT),
                   ("K",f"=SUM(K{R_G0}:K{R_G1})",MM,KEY),("L",f"=SUM(L{R_G0}:L{R_G1})",PCT,TOT)]:
    S(ws2,f"{c}{R_GT}",bold=True,fmt=fm,border=True,align="center" if c!="J" else None,fill=fl)
    ws2[f"{c}{R_GT}"]=f_
S(ws2,f"N{R_G0}",size=9,color="7F7F7F")
ws2[f"N{R_G0}"]="店内席効率は80%→75%に変更。ディナー店内回転率0.80→0.90"
S(ws2,f"N{R_G0+2}",size=9,color="7F7F7F")
ws2[f"N{R_G0+2}"]="カフェ客単価2,500円（料理1,500＋飲料1,000）に変更"
S(ws2,f"N{R_G0+5}",size=9,color="7F7F7F")
ws2[f"N{R_G0+5}"]="テラスのディナーは客単価10,000円（店内15,000円）・回転率0.50に変更"

sec(ws2, 28, "■ 料理・飲料別の年間売上（原価計算用）", "L")
S(ws2,"A29"); ws2["A29"]="料理売上"
S(ws2,"B29",fmt=MM,border=True,align="center")
ws2["B29"]=f"=SUMPRODUCT(I{R_G0}:I{R_G1},F{R_G0}:F{R_G1})*$B$5/1000000"
S(ws2,"A30"); ws2["A30"]="飲料売上"
S(ws2,"B30",fmt=MM,border=True,align="center")
ws2["B30"]=f"=SUMPRODUCT(I{R_G0}:I{R_G1},G{R_G0}:G{R_G1})*$B$5/1000000"
S(ws2,"A31",bold=True); ws2["A31"]="合　計"
S(ws2,"B31",bold=True,fmt=MM,border=True,align="center",fill=TOT); ws2["B31"]="=B29+B30"
S(ws2,"A32"); ws2["A32"]="料理 構成比"
S(ws2,"B32",fmt=PCT,border=True,align="center"); ws2["B32"]="=IF($B$31=0,0,B29/$B$31)"
S(ws2,"A33"); ws2["A33"]="飲料 構成比"
S(ws2,"B33",fmt=PCT,border=True,align="center"); ws2["B33"]="=IF($B$31=0,0,B30/$B$31)"

sec(ws2, 35, "■ 参考指標", "L")
for i,(lab,f_,fm,note) in enumerate([
    ("日商（円）",f"=J{R_GT}",YEN,""),("月商（百万円）",f"=K{R_GT}/12",MM,""),
    ("年商（百万円・満年度）",f"=K{R_GT}",MM,"提出P&L：4.60億円"),
    ("1日あたり総客数（人）",f"=I{R_GT}",NUM1,""),
    ("平均客単価（円）",f"=IF(I{R_GT}=0,0,J{R_GT}/I{R_GT})",YEN,""),
    ("延べ席回転率（回/日）",f"=IF($B$16=0,0,I{R_GT}/$B$16)",NUM2,"総客数 ÷ 総席数"),
    ("坪数（坪）",55.5,NUM1,"183.59㎡"),
    ("坪売上（百万円/坪・年）",f"=IF(B42=0,0,K{R_GT}/B42)",MM,""),
    ("1席あたり年間売上（百万円）",f"=IF($B$16=0,0,K{R_GT}/$B$16)",MM,"席数変更の効果を見る指標")]):
    r=36+i
    S(ws2,f"A{r}"); ws2[f"A{r}"]=lab
    isin = not isinstance(f_,str)
    S(ws2,f"B{r}",color=BLUE if isin else BLACK,fmt=fm,border=True,align="center",fill=INP if isin else None)
    ws2[f"B{r}"]=f_
    S(ws2,f"N{r}",size=9,color="7F7F7F"); ws2[f"N{r}"]=note

sec(ws2, 46, "■ 立上りを織り込んだ年間売上（百万円）", "L")
for c,t in zip("BCDE",["満年度","1年目","2年目","3年目以降"]):
    S(ws2,f"{c}47",bold=True,border=True,fill=TOT,align="center"); ws2[f"{c}47"]=t
S(ws2,"A47",bold=True,border=True,fill=TOT); ws2["A47"]="項　目"
S(ws2,"A48"); ws2["A48"]="立上り率"
S(ws2,"A49",bold=True); ws2["A49"]="年間売上高"
for c,ramp in zip("BCDE",["$B$8","$B$6","$B$7","$B$8"]):
    S(ws2,f"{c}48",fmt=PCT,border=True,align="center"); ws2[f"{c}48"]=f"={ramp}"
    S(ws2,f"{c}49",bold=True,fmt=MM,border=True,align="center",fill=KEY)
    ws2[f"{c}49"]=f"=$K${R_GT}*{c}48"
R_RAMPROW, R_SALES = 48, 49

# ═══════════════ 経費前提 ═══════════════
ws3 = wb.create_sheet("経費前提"); ws3.sheet_view.showGridLines=False
ws3.column_dimensions["A"].width=34
for c in "BCD": ws3.column_dimensions[c].width=15
ws3.column_dimensions["E"].width=3; ws3.column_dimensions["F"].width=58
band(ws3,1,"経費前提（水色セル＝入力値）","F")
S(ws3,"A2",size=9,color="7F7F7F")
ws3["A2"]="経費率は2026年8月3日提出P&L準拠。2026年10月改定で売上原価は料理・飲料別（各30.0%）に変更。"

sec(ws3,4,"■ 売上原価の計算方法","F")
S(ws3,"A5"); ws3["A5"]="計算方法（1＝売上比一括／2＝料理・飲料別）"
S(ws3,"B5",color=BLUE,fmt=INT,border=True,align="center",fill=INP); ws3["B5"]=2
S(ws3,"F5",size=9,color="7F7F7F")
ws3["F5"]="2026年10月改定で方法2を選択。下の変動費計・感応度分析もこの選択に連動します。"
R_METHOD=5

sec(ws3,7,"■ 経費率（対売上高）","F")
for c,t in zip("ABCD",["費　目","提出P&L水準\n（満年度・1年目）","2年目","3年目以降"]):
    S(ws3,f"{c}8",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True); ws3[f"{c}8"]=t
ws3.row_dimensions[8].height=32
RATES=[("売上原価（方法1：売上比一括）",0.295,0.290,0.285,"方法1を選んだ場合のみ使用"),
       ("　料理原価率（方法2）",0.300,0.300,0.300,"2026年10月改定：全年度30.0%で固定"),
       ("　飲料原価率（方法2）",0.300,0.300,0.300,"同上。直輸入ワインによる低減は織り込んでいない"),
       ("人件費",0.345,0.335,0.325,"本部集中化・多能工化。人員26名"),
       ("業務委託費・消耗品他",0.024,0.022,0.020,"本部集中購買・規格統一"),
       ("広告宣伝費・販売促進費",0.014,0.014,0.014,"据置"),
       ("その他経費",0.005,0.005,0.005,"据置"),
       ("水道光熱費",0.040,0.038,0.036,"高効率厨房機器・営業時間最適化"),
       ("ロイヤリティ",0.040,0.040,0.040,"契約条件（ブラッスリー4.0%）"),
       ("保険・租税公課",0.007,0.007,0.007,"据置"),
       ("修繕維持費（FFE）",0.025,0.025,0.025,"据置")]
R_R0=9
for i,(lab,a,b,c_,note) in enumerate(RATES):
    r=R_R0+i; sub=lab.startswith("　")
    S(ws3,f"A{r}",border=True,size=9 if sub else 10,color="7F7F7F" if sub else BLACK); ws3[f"A{r}"]=lab
    for col,v in zip("BCD",[a,b,c_]):
        S(ws3,f"{col}{r}",color=BLUE,fmt=PCT2 if sub else PCT,border=True,align="center",fill=INP)
        ws3[f"{col}{r}"]=v
    S(ws3,f"F{r}",size=9,color="7F7F7F"); ws3[f"F{r}"]=note
R_COST1,R_RY,R_IN = R_R0,R_R0+1,R_R0+2
R_JIN,R_ITAKU,R_KOK,R_SON,R_SUI,R_ROY,R_HOK,R_SHU = [R_R0+i for i in range(3,11)]
R_VSUM = R_R0+11   # 20
S(ws3,f"A{R_VSUM}",bold=True,border=True); ws3[f"A{R_VSUM}"]="変動費計（選択中の原価方法ベース）"
for col in "BCD":
    S(ws3,f"{col}{R_VSUM}",bold=True,fmt=PCT,border=True,align="center",fill=TOT)
    ws3[f"{col}{R_VSUM}"]=(f"=IF($B${R_METHOD}=1,{col}{R_COST1},"
                           f"売上前提!$B$32*{col}{R_RY}+売上前提!$B$33*{col}{R_IN})"
                           f"+SUM({col}{R_JIN}:{col}{R_SHU})")
S(ws3,f"F{R_VSUM}",size=9,color=RED)
ws3[f"F{R_VSUM}"]="原価の計算方法（B5）に連動。感応度分析シートもこの値を参照します。"

sec(ws3,22,"■ 固定費・家賃条件","F")
for i,(lab,v,f_,note) in enumerate([
    ("地代家賃（固定）月額",2.22,MM,"月222万円（提出P&L）。席数を変えても面積は同じため不変"),
    ("地代家賃（固定）年額",None,MM,"＝月額×12"),
    ("歩合家賃率",0.07,PCT,"閾値超過分に対して7%"),
    ("歩合家賃 閾値（年商）",331.4,MM,"満年度の歩合家賃9百万円から逆算（提出P&L時点）"),
    ("人員数（名・参考）",26,INT,"提出P&L。人件費は売上比で計算")]):
    r=23+i
    S(ws3,f"A{r}"); ws3[f"A{r}"]=lab
    isin = v is not None
    S(ws3,f"B{r}",color=BLUE if isin else BLACK,fmt=f_,border=True,align="center",fill=INP if isin else None)
    ws3[f"B{r}"]= v if isin else "=B23*12"
    S(ws3,f"F{r}",size=9,color="7F7F7F"); ws3[f"F{r}"]=note
R_RENT_M,R_RENT_Y,R_PCTRENT,R_THRESH,R_STAFF = 23,24,25,26,27

sec(ws3,29,"■ 初期投資・減価償却（百万円）","F")
for i,(lab,v,note) in enumerate([
    ("契約金",10.0,"契約金3,000万円÷3店舗"),
    ("事業費（設計・内装施工・OSE/FF&E）","=55.5*3.5","55.5坪×350万"),
    ("開業準備金（広告・採用等）",4.0,""),
    ("保証金",39.0,"非償却・退去時返還対象"),
    ("ワイン＋美術品",50.0,"美術品は数百万前半。非償却（下記参照）")]):
    r=30+i
    S(ws3,f"A{r}"); ws3[f"A{r}"]=lab
    S(ws3,f"B{r}",color=BLUE,fmt=MM,border=True,align="center",fill=INP); ws3[f"B{r}"]=v
    S(ws3,f"F{r}",size=9,color="7F7F7F"); ws3[f"F{r}"]=note
R_WINE=34
S(ws3,"A35",bold=True); ws3["A35"]="初期投資　合計"
S(ws3,"B35",bold=True,fmt=MM,border=True,align="center",fill=KEY); ws3["B35"]="=SUM(B30:B34)"
S(ws3,"A36",bold=True); ws3["A36"]="償却対象資産（＝合計－保証金－ワイン/美術品）"
S(ws3,"B36",bold=True,fmt=MM,border=True,align="center",fill=TOT); ws3["B36"]="=B35-B33-B34"
S(ws3,"F36",size=9,color=RED)
ws3["F36"]="ワイン在庫は棚卸資産、美術品は取得価額100万円以上なら原則非償却のため除外（2026年10月修正）"
S(ws3,"A37"); ws3["A37"]="償却年数（年）"
S(ws3,"B37",color=BLUE,fmt=NUM1,border=True,align="center",fill=INP); ws3["B37"]=10.0
S(ws3,"A38",bold=True); ws3["A38"]="年間減価償却費"
S(ws3,"B38",bold=True,fmt=MM,border=True,align="center",fill=TOT); ws3["B38"]="=B36/B37"
S(ws3,"F38",size=9,color="7F7F7F")
ws3["F38"]="美術品のうち100万円未満で償却したい分があれば、その金額をB34から事業費（B31）へ移してください"
R_INVTOT,R_DEPBASE,R_DEPRE = 35,36,38

sec(ws3,40,"■ 分割払いスキーム（百万円）","F")
for i,(lab,v,note) in enumerate([
    ("分割対象：内装工事費",100.0,"月833千円 × 120回（10年）"),
    ("分割対象：保証金",38.9,"月463千円 × 84回（7年）"),
    ("分割対象額　計",None,""),
    ("開業時 現金支出",None,"＝初期投資合計 － 分割対象額"),
    ("年間分割弁済額",15.6,"月額合計1,296千円 × 12か月")]):
    r=41+i
    S(ws3,f"A{r}",bold=(v is None)); ws3[f"A{r}"]=lab
    isin = v is not None
    S(ws3,f"B{r}",color=BLUE if isin else BLACK,bold=(v is None),fmt=MM,border=True,
      align="center",fill=INP if isin else TOT)
    if lab=="分割対象額　計": ws3[f"B{r}"]="=B41+B42"
    elif lab=="開業時 現金支出": ws3[f"B{r}"]="=B35-B43"
    else: ws3[f"B{r}"]=v
    S(ws3,f"F{r}",size=9,color="7F7F7F"); ws3[f"F{r}"]=note
R_DEF,R_CASH,R_PAY = 43,44,45
print("経費前提 OK")

# ═══════════════ 単店PL ═══════════════
ws4 = wb.create_sheet("単店PL"); ws4.sheet_view.showGridLines=False
ws4.column_dimensions["A"].width=30
for c in "BCDE": ws4.column_dimensions[c].width=14
ws4.column_dimensions["F"].width=3; ws4.column_dimensions["G"].width=44
band(ws4,1,"六本木ミッドタウン店　単店PL（百万円）","E")
S(ws4,"A2",size=9,color="7F7F7F")
ws4["A2"]="科目は2026年8月3日提出P&Lのまま。売上は「売上前提」、経費率は「経費前提」に連動。金利・税金は含まない。"
COLS=["B","C","D","E"]; RATECOL=["B","B","C","D"]
for c,t in zip(COLS,["満年度\n（提出P&L水準）","1年目","2年目","3年目以降"]):
    S(ws4,f"{c}4",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True); ws4[f"{c}4"]=t
S(ws4,"A4",bold=True,border=True,fill=TOT); ws4["A4"]="科　目"
ws4.row_dimensions[4].height=30
r=5
S(ws4,f"A{r}"); ws4[f"A{r}"]="立上り率"
for c in COLS:
    S(ws4,f"{c}{r}",color=GREEN,fmt=PCT,border=True,align="center")
    ws4[f"{c}{r}"]=f"=売上前提!{c}{R_RAMPROW}"
r+=1
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="売上高"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,color=GREEN,fmt=MM,border=True)
    ws4[f"{c}{r}"]=f"=売上前提!{c}${R_SALES}"
R_REV=r; r+=1
for lab,ref in [("　うち 料理売上","$B$32"),("　うち 飲料売上","$B$33")]:
    S(ws4,f"A{r}",size=9,color="7F7F7F",border=True); ws4[f"A{r}"]=lab
    for c in COLS:
        S(ws4,f"{c}{r}",size=9,color="7F7F7F",fmt=MM,border=True)
        ws4[f"{c}{r}"]=f"={c}${R_REV}*売上前提!{ref}"
    r+=1
R_RYS,R_INS = R_REV+1,R_REV+2
PL=[]
S(ws4,f"A{r}",border=True); ws4[f"A{r}"]="売上原価"
for c,rc in zip(COLS,RATECOL):
    S(ws4,f"{c}{r}",fmt=MM,border=True)
    ws4[f"{c}{r}"]=(f"=IF(経費前提!$B${R_METHOD}=1,{c}${R_REV}*経費前提!${rc}${R_COST1},"
                    f"{c}${R_RYS}*経費前提!${rc}${R_RY}+{c}${R_INS}*経費前提!${rc}${R_IN})")
R_CGS=r; PL.append(("売上原価",r)); r+=1
def var_row(label, rate_row):
    global r
    S(ws4,f"A{r}",border=True); ws4[f"A{r}"]=label
    for c,rc in zip(COLS,RATECOL):
        S(ws4,f"{c}{r}",fmt=MM,border=True)
        ws4[f"{c}{r}"]=f"={c}${R_REV}*経費前提!${rc}${rate_row}"
    PL.append((label,r)); r+=1; return r-1
R_JINK=var_row("人件費",R_JIN); R_ITK=var_row("業務委託費・消耗品他",R_ITAKU)
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="部門利益"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=TOT)
    ws4[f"{c}{r}"]=f"={c}{R_REV}-{c}{R_CGS}-{c}{R_JINK}-{c}{R_ITK}"
R_BUMON=r; PL.append(("部門利益",r)); r+=1
R_KK=var_row("広告宣伝費・販売促進費",R_KOK); R_SN=var_row("その他経費",R_SON); R_SU=var_row("水道光熱費",R_SUI)
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="GOP（店舗営業総利益）"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=TOT)
    ws4[f"{c}{r}"]=f"={c}{R_BUMON}-{c}{R_KK}-{c}{R_SN}-{c}{R_SU}"
R_GOP=r; PL.append(("GOP（店舗営業総利益）",r)); r+=1
R_RO=var_row("ロイヤリティ",R_ROY); R_HK=var_row("保険・租税公課",R_HOK)
S(ws4,f"A{r}",border=True); ws4[f"A{r}"]="地代家賃（固定）"
for c in COLS:
    S(ws4,f"{c}{r}",color=GREEN,fmt=MM,border=True); ws4[f"{c}{r}"]=f"=経費前提!$B${R_RENT_Y}"
R_RENT=r; PL.append(("地代家賃（固定）",r)); r+=1
S(ws4,f"A{r}",border=True); ws4[f"A{r}"]="歩合家賃"
for c in COLS:
    S(ws4,f"{c}{r}",fmt=MM,border=True)
    ws4[f"{c}{r}"]=f"=MAX(0,{c}{R_REV}-経費前提!$B${R_THRESH})*経費前提!$B${R_PCTRENT}"
R_PRENT=r; PL.append(("歩合家賃",r)); r+=1
R_SH=var_row("修繕維持費（FFE）",R_SHU)
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="EBITDA"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=KEY)
    ws4[f"{c}{r}"]=f"={c}{R_GOP}-{c}{R_RO}-{c}{R_HK}-{c}{R_RENT}-{c}{R_PRENT}-{c}{R_SH}"
R_EBITDA=r; PL.append(("EBITDA",r)); r+=1
S(ws4,f"A{r}",border=True); ws4[f"A{r}"]="減価償却費"
for c in COLS:
    S(ws4,f"{c}{r}",color=GREEN,fmt=MM,border=True); ws4[f"{c}{r}"]=f"=経費前提!$B${R_DEPRE}"
R_DEP=r; PL.append(("減価償却費",r)); r+=1
S(ws4,f"A{r}",bold=True,border=True); ws4[f"A{r}"]="営業利益"
for c in COLS:
    S(ws4,f"{c}{r}",bold=True,fmt=MM,border=True,fill=KEY)
    ws4[f"{c}{r}"]=f"={c}{R_EBITDA}-{c}{R_DEP}"
R_OP=r; PL.append(("営業利益",r)); r+=2
sec(ws4,r,"■ 対売上比","E"); r+=1
for c,t in zip(COLS,["満年度","1年目","2年目","3年目以降"]):
    S(ws4,f"{c}{r}",bold=True,border=True,fill=TOT,align="center"); ws4[f"{c}{r}"]=t
S(ws4,f"A{r}",bold=True,border=True,fill=TOT); ws4[f"A{r}"]="科　目"; r+=1
for lab,src in [("売上高",R_REV)]+PL:
    hl = lab in ("EBITDA","営業利益","部門利益","GOP（店舗営業総利益）")
    S(ws4,f"A{r}",border=True,bold=hl); ws4[f"A{r}"]=lab
    for c in COLS:
        S(ws4,f"{c}{r}",fmt=PCT,border=True,align="center",bold=hl,
          fill=KEY if lab in ("EBITDA","営業利益") else (TOT if hl else None))
        ws4[f"{c}{r}"]=f"=IF({c}${R_REV}=0,0,{c}{src}/{c}${R_REV})"
    r+=1
S(ws4,f"G{R_EBITDA}",size=9,color=RED,bold=True); ws4[f"G{R_EBITDA}"]="提出P&L（旧前提）：満年度EBITDA 59百万円"
S(ws4,f"G{R_OP}",size=9,color=RED,bold=True); ws4[f"G{R_OP}"]="提出P&L（旧前提）：満年度 営業利益 38百万円"

# ═══════════════ 投資・CF ═══════════════
ws5 = wb.create_sheet("投資・CF"); ws5.sheet_view.showGridLines=False
ws5.column_dimensions["A"].width=34
for c in "BCDEFG": ws5.column_dimensions[c].width=13
ws5.column_dimensions["H"].width=3; ws5.column_dimensions["I"].width=52
band(ws5,1,"初期投資・キャッシュフロー（百万円）","G")
S(ws5,"A2",size=9,color="7F7F7F")
ws5["A2"]="分割払いスキーム活用後の開業時現金支出と、単店の投資回収を自動計算します。"
sec(ws5,4,"■ 初期投資","G")
r=5
for lab,ref in [("契約金",30),("事業費（設計・内装施工・OSE/FF&E）",31),
                ("開業準備金（広告・採用等）",32),("保証金",33),("ワイン＋美術品",R_WINE)]:
    S(ws5,f"A{r}",border=True); ws5[f"A{r}"]=lab
    S(ws5,f"B{r}",color=GREEN,fmt=MM,border=True,align="center"); ws5[f"B{r}"]=f"=経費前提!$B${ref}"
    r+=1
S(ws5,f"I{9}",size=9,color=RED); ws5["I9"]="2026年10月追加。減価償却の対象外（経費前提F36参照）"
for lab,ref,hl in [("初期投資　合計",R_INVTOT,True),("うち 償却対象資産",R_DEPBASE,False),
                   ("分割対象額",R_DEF,False),("開業時 現金支出",R_CASH,True),
                   ("年間分割弁済額",R_PAY,False)]:
    S(ws5,f"A{r}",bold=hl,border=True); ws5[f"A{r}"]=lab
    S(ws5,f"B{r}",bold=hl,color=GREEN,fmt=MM,border=True,align="center",fill=KEY if hl else None)
    ws5[f"B{r}"]=f"=経費前提!$B${ref}"
    r+=1
R_CHECK=r
S(ws5,f"A{r}",border=True,size=9); ws5[f"A{r}"]="（検算）内訳合計と一致するか"
S(ws5,f"B{r}",fmt=MM,border=True,align="center",size=9); ws5[f"B{r}"]="=SUM(B5:B9)-B10"
S(ws5,f"I{r}",size=9,color="7F7F7F"); ws5[f"I{r}"]="0なら内訳と合計が一致"

sec(ws5,17,"■ 単店キャッシュフロー推移","G")
YC=["B","C","D","E","F","G"]
for c,t in zip(YC,["開業時","1年目","2年目","3年目","4年目","5年目"]):
    S(ws5,f"{c}18",bold=True,border=True,fill=TOT,align="center"); ws5[f"{c}18"]=t
S(ws5,"A18",bold=True,border=True,fill=TOT); ws5["A18"]="項　目"
EBCOL=[None,"C","D","E","E","E"]
S(ws5,"A19",border=True); ws5["A19"]="EBITDA"
for c,src in zip(YC,EBCOL):
    S(ws5,f"{c}19",color=GREEN if src else BLACK,fmt=MM,border=True,align="center")
    ws5[f"{c}19"]=(f"=単店PL!{src}{R_EBITDA}" if src else 0)
S(ws5,"A20",border=True); ws5["A20"]="開業時 現金支出"
for i,c in enumerate(YC):
    S(ws5,f"{c}20",fmt=MM,border=True,align="center")
    ws5[f"{c}20"]=(f"=-経費前提!$B${R_CASH}" if i==0 else 0)
S(ws5,"A21",border=True); ws5["A21"]="分割弁済"
for i,c in enumerate(YC):
    S(ws5,f"{c}21",fmt=MM,border=True,align="center")
    ws5[f"{c}21"]=(0 if i==0 else f"=-経費前提!$B${R_PAY}")
S(ws5,"A22",bold=True,border=True); ws5["A22"]="単店FCF"
for c in YC:
    S(ws5,f"{c}22",bold=True,fmt=MM,border=True,align="center",fill=TOT)
    ws5[f"{c}22"]=f"={c}19+{c}20+{c}21"
S(ws5,"A23",bold=True,border=True); ws5["A23"]="累計FCF"
for i,c in enumerate(YC):
    S(ws5,f"{c}23",bold=True,fmt=MM,border=True,align="center",fill=KEY)
    ws5[f"{c}23"]=f"={c}22" if i==0 else f"={YC[i-1]}23+{c}23".replace(f"{c}23",f"{c}22",1) if False else (f"={YC[i-1]}23+{c}22")
R_FCF,R_CUM=22,23

sec(ws5,25,"■ 投資回収（自動計算）","G")
S(ws5,"A26",border=True); ws5["A26"]="累計FCFが赤字の年数"
S(ws5,"B26",fmt=INT,border=True,align="center"); ws5["B26"]=f'=COUNTIF(C{R_CUM}:G{R_CUM},"<0")'
S(ws5,"A27",bold=True,border=True); ws5["A27"]="累計FCFの黒字転換"
S(ws5,"B27",bold=True,border=True,align="center",fill=KEY)
ws5["B27"]=f'=IF(B26>=5,"5年内に未回収",B26+1&"年目")'
S(ws5,"A28",bold=True,border=True); ws5["A28"]="投資回収年数（年）"
S(ws5,"B28",bold=True,fmt=NUM2,border=True,align="center",fill=KEY)
ws5["B28"]=(f'=IF(B26>=5,"",IF(INDEX(B{R_FCF}:G{R_FCF},B26+2)=0,"",'
            f'B26+(-INDEX(B{R_CUM}:G{R_CUM},B26+1))/INDEX(B{R_FCF}:G{R_FCF},B26+2)))')
S(ws5,"A29",border=True); ws5["A29"]="定常期の年間FCF"
S(ws5,"B29",fmt=MM,border=True,align="center"); ws5["B29"]=f"=G{R_FCF}"
S(ws5,"I26",size=9,color="7F7F7F"); ws5["I26"]="固定文ではなく計算結果です。前提を変えると自動で更新されます。"
S(ws5,"A31",bold=True,size=10,color=RED)
ws5["A31"]=('=\"→ 開業\"&B27&\"に累計FCFが黒字転換（投資回収 約\"&TEXT(B28,\"0.0\")&\"年）。\"')
S(ws5,"A32",size=9,color="7F7F7F")
ws5["A32"]="※ 税金・運転資本の増減・借入返済・分割払いに係る金利相当は含まない簡易ベース。"
S(ws5,"A33",size=9,color="7F7F7F")
ws5["A33"]="※ 保証金39百万円は退去時返還対象。ワイン在庫も換価可能なため、実質的な回収はさらに短期。"
R_PBYEAR,R_PBNUM=27,28
print("投資CF OK")

# ═══════════════ 感応度分析 ═══════════════
ws6 = wb.create_sheet("感応度分析"); ws6.sheet_view.showGridLines=False
ws6.column_dimensions["A"].width=26
for c in "BCDEFG": ws6.column_dimensions[c].width=13
ws6.column_dimensions["H"].width=3; ws6.column_dimensions["I"].width=50
band(ws6,1,"感応度分析　店内席効率 × ディナー店内回転率","G")
S(ws6,"A2",size=9,color="7F7F7F")
ws6["A2"]="満年度・提出P&L水準。原価は「経費前提」で選択中の方法に連動。軸の値（水色）は自由に変更できます。"
EFFS=[0.65,0.70,0.75,0.80,0.85]; TURNS=[0.70,0.80,0.90,1.00,1.10]
IN_FIX=("売上前提!$C$20*{x}*売上前提!$E$20*売上前提!$H$20"
        "+売上前提!$C$22*{x}*売上前提!$E$22*売上前提!$H$22"
        "+売上前提!$C$24*{x}*{y}*売上前提!$H$24")
TE_FIX="(売上前提!$J$21+売上前提!$J$23+売上前提!$J$25)"
VSUM=f"経費前提!$B${R_VSUM}"
def sales_expr(xc,yc): return f"(({IN_FIX.format(x=xc,y=yc)})+{TE_FIX})*売上前提!$B$5/1000000"
def build_grid(top,title,metric):
    sec(ws6,top,title,"G")
    hdr=top+1
    S(ws6,f"A{hdr}",bold=True,border=True,fill=TOT,align="center",size=9,wrap=True)
    ws6[f"A{hdr}"]="店内席効率 ＼ ディナー回転率"
    ws6.row_dimensions[hdr].height=28
    for j,tv in enumerate(TURNS):
        c=chr(ord("B")+j)
        S(ws6,f"{c}{hdr}",bold=True,border=True,fill=INP,align="center",fmt=NUM2,color=BLUE)
        ws6[f"{c}{hdr}"]=tv
    for i,ev in enumerate(EFFS):
        rr=hdr+1+i
        S(ws6,f"A{rr}",bold=True,border=True,fill=INP,align="center",fmt=PCT,color=BLUE)
        ws6[f"A{rr}"]=ev
        for j in range(len(TURNS)):
            c=chr(ord("B")+j); xc=f"$A{rr}"; yc=f"{c}${hdr}"; s=sales_expr(xc,yc)
            prof=(f"(({s})*(1-{VSUM})-経費前提!$B${R_RENT_Y}"
                  f"-MAX(0,({s})-経費前提!$B${R_THRESH})*経費前提!$B${R_PCTRENT}"
                  f"-経費前提!$B${R_DEPRE})")
            if metric=="sales": f_=f"={s}"; fmt=MM
            elif metric=="op":  f_=f"={prof}"; fmt=MM
            else:               f_=f"=IF(({s})=0,0,{prof}/({s}))"; fmt=PCT
            S(ws6,f"{c}{rr}",fmt=fmt,border=True,align="center"); ws6[f"{c}{rr}"]=f_
    rng=f"B{hdr+1}:F{hdr+len(EFFS)}"
    ws6.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(ROUND($A{hdr+1},4)=ROUND(売上前提!$D$20,4),'
                 f'ROUND(B${hdr},4)=ROUND(売上前提!$E$24,4))'], fill=HIT, stopIfTrue=False))
    return hdr+len(EFFS)+1
nxt=build_grid(4,"■ 年間売上高（百万円）","sales")
nxt=build_grid(nxt+1,"■ 営業利益（百万円）","op")
nxt=build_grid(nxt+1,"■ 営業利益率","opm")
S(ws6,f"A{nxt+1}",size=9,color=RED)
ws6[f"A{nxt+1}"]="薄黄色＝現在の前提に一致するセル（席効率・回転率を変えると自動で移動します）。"
S(ws6,f"A{nxt+2}",size=9,color="7F7F7F")
ws6[f"A{nxt+2}"]="※ 料理・飲料の構成比は基準ケースの比率を用いた近似。地代家賃と減価償却は売上非連動のため、売上増で利益率が加速的に改善します。"

# ═══════════════ サマリー ═══════════════
ws1 = wb.create_sheet("サマリー",0); ws1.sheet_view.showGridLines=False
ws1.column_dimensions["A"].width=32
for c in "BCDE": ws1.column_dimensions[c].width=15
ws1.column_dimensions["F"].width=3; ws1.column_dimensions["G"].width=40
band(ws1,1,"BRASSERIE Thierry Marx　六本木ミッドタウン店　事業収支（2026年10月改定 rev2）","E")
S(ws1,"A2",size=9,color="7F7F7F")
ws1["A2"]="単位：百万円／183.59㎡(55.5坪)・年363日営業。水色セルを変えると全体が再計算されます。"
sec(ws1,4,"■ よく変える前提","E")
S(ws1,"A5",size=9,color="7F7F7F")
ws1["A5"]="詳細は「売上前提」。時間帯×エリアごとの席効率・回転率・客単価はそちらで個別に設定できます。"
for i,(lab,f_,fmt) in enumerate([
    ("総席数（店内＋テラス）","=売上前提!$B$16",INT),("　店内","=売上前提!$B$14",INT),
    ("　テラス","=売上前提!$B$15",INT),("店内 席効率","=売上前提!$D$20",PCT),
    ("ディナー 回転率（店内）","=売上前提!$E$24",NUM2),("客単価 ランチ","=売上前提!$H$20",YEN),
    ("客単価 カフェ","=売上前提!$H$22",YEN),("客単価 ディナー（店内）","=売上前提!$H$24",YEN),
    ("客単価 ディナー（テラス）","=売上前提!$H$25",YEN),("営業日数（日／年）","=売上前提!$B$5",NUM1)]):
    r=6+i; sub=lab.startswith("　")
    S(ws1,f"A{r}",border=True,size=9 if sub else 10,color="7F7F7F" if sub else BLACK); ws1[f"A{r}"]=lab
    S(ws1,f"B{r}",color=GREEN,fmt=fmt,border=True,align="center"); ws1[f"B{r}"]=f_
sec(ws1,17,"■ 単店PL サマリー","E")
for c,t in zip("BCDE",["満年度","1年目","2年目","3年目以降"]):
    S(ws1,f"{c}18",bold=True,border=True,fill=TOT,align="center"); ws1[f"{c}18"]=t
S(ws1,"A18",bold=True,border=True,fill=TOT); ws1["A18"]="項　目"
r=19
for lab,src,hl in [("売上高",R_REV,False),("部門利益",R_BUMON,False),
                   ("GOP（店舗営業総利益）",R_GOP,False),("EBITDA",R_EBITDA,True),
                   ("減価償却費",R_DEP,False),("営業利益",R_OP,True)]:
    S(ws1,f"A{r}",bold=hl,border=True); ws1[f"A{r}"]=lab
    for c in "BCDE":
        S(ws1,f"{c}{r}",bold=hl,color=GREEN,fmt=MM,border=True,align="center",fill=KEY if hl else None)
        ws1[f"{c}{r}"]=f"=単店PL!{c}{src}"
    r+=1
for lab,src in [("　EBITDA率",R_EBITDA),("　営業利益率",R_OP)]:
    S(ws1,f"A{r}",bold=True,border=True); ws1[f"A{r}"]=lab
    for c in "BCDE":
        S(ws1,f"{c}{r}",bold=True,fmt=PCT,border=True,align="center",fill=KEY)
        ws1[f"{c}{r}"]=f"=IF(単店PL!{c}{R_REV}=0,0,単店PL!{c}{src}/単店PL!{c}{R_REV})"
    r+=1
r+=1
sec(ws1,r,"■ 売上の内訳（満年度）","E"); r+=1
for c,t in zip("BCD",["年間売上","構成比","1日客数"]):
    S(ws1,f"{c}{r}",bold=True,border=True,fill=TOT,align="center"); ws1[f"{c}{r}"]=t
S(ws1,f"A{r}",bold=True,border=True,fill=TOT); ws1[f"A{r}"]="時間帯 × エリア"; r+=1
for i in range(6):
    src=R_G0+i
    S(ws1,f"A{r}",border=True,size=9); ws1[f"A{r}"]=f'=売上前提!A{src}&"　"&売上前提!B{src}'
    S(ws1,f"B{r}",color=GREEN,fmt=MM,border=True,align="center"); ws1[f"B{r}"]=f"=売上前提!K{src}"
    S(ws1,f"C{r}",fmt=PCT,border=True,align="center"); ws1[f"C{r}"]=f"=売上前提!L{src}"
    S(ws1,f"D{r}",fmt=NUM1,border=True,align="center"); ws1[f"D{r}"]=f"=売上前提!I{src}"
    r+=1
S(ws1,f"A{r}",bold=True,border=True); ws1[f"A{r}"]="合　計"
S(ws1,f"B{r}",bold=True,color=GREEN,fmt=MM,border=True,align="center",fill=KEY); ws1[f"B{r}"]=f"=売上前提!K{R_GT}"
S(ws1,f"C{r}",bold=True,fmt=PCT,border=True,align="center",fill=TOT); ws1[f"C{r}"]=f"=売上前提!L{R_GT}"
S(ws1,f"D{r}",bold=True,fmt=NUM1,border=True,align="center",fill=TOT); ws1[f"D{r}"]=f"=売上前提!I{R_GT}"
r+=2
sec(ws1,r,"■ 投資回収","E"); r+=1
for lab,ref,fmt in [("初期投資 合計",f"=経費前提!$B${R_INVTOT}",MM),
                    ("　うち 償却対象資産",f"=経費前提!$B${R_DEPBASE}",MM),
                    ("　うち ワイン＋美術品（非償却）",f"=経費前提!$B${R_WINE}",MM),
                    ("開業時 現金支出（分割払い活用後）",f"=経費前提!$B${R_CASH}",MM),
                    ("定常期の年間FCF",f"=投資・CF!B29",MM),
                    ("累計FCFの黒字転換",f"=投資・CF!B{R_PBYEAR}",None),
                    ("投資回収年数（年）",f"=投資・CF!B{R_PBNUM}",NUM2)]:
    sub=lab.startswith("　")
    S(ws1,f"A{r}",border=True,size=9 if sub else 10,color="7F7F7F" if sub else BLACK); ws1[f"A{r}"]=lab
    S(ws1,f"B{r}",color=GREEN,fmt=fmt,border=True,align="center",
      fill=KEY if "回収年数" in lab or "黒字転換" in lab else None)
    ws1[f"B{r}"]=ref
    r+=1
S(ws1,f"A{r+1}",bold=True,size=10,color=RED)
ws1[f"A{r+1}"]='="→ 開業"&投資・CF!B27&"に累計FCFが黒字転換（投資回収 約"&TEXT(投資・CF!B28,"0.0")&"年）。"'

if "Sheet" in wb.sheetnames: del wb["Sheet"]
SHEETS=["サマリー","売上前提","経費前提","単店PL","投資・CF","感応度分析"]
for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith("="):
                f=c.value
                for nm in SHEETS: f=f.replace(f"'{nm}'!",f"{nm}!")
                for nm in SHEETS: f=f.replace(f"{nm}!",f"'{nm}'!")
                c.value=f
    for rng in getattr(sh.conditional_formatting,"_cf_rules",{}):
        pass
wb.calculation.fullCalcOnLoad=True
wb.save(OUT)
print("saved:",OUT)
