# カンテサンス CF計画（2027〜2032年・三つ星維持）IM 42ページの計画（〜2029年）をベースに、2030年以降は同じ年平均成長率で延長
import sys,re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUT=sys.argv[1]
wb=Workbook()
F="Arial"
BLUE=Font(name=F,size=10,color="0000FF"); BLK=Font(name=F,size=10); GRN=Font(name=F,size=10,color="008000")
B=Font(name=F,size=10,bold=True); H=Font(name=F,size=10,bold=True,color="FFFFFF"); T=Font(name=F,size=14,bold=True)
NOTE=Font(name=F,size=9,color="6E6368")
HF=PatternFill("solid",fgColor="6D2E46"); YF=PatternFill("solid",fgColor="FFFF00"); TF=PatternFill("solid",fgColor="F7F3F4")
NUM='#,##0;(#,##0);"-"'; NUM1='#,##0.0;(#,##0.0);"-"'; PCT='0.0%;(0.0%);"-"'; X='0.00"x";(0.00"x");"-"'
YRS=["2027","2028","2029","2030","2031","2032"]; YC=["C","D","E","F","G","H"]
IMY=["C","D","E"]; EXT=["F","G","H"]

def header(ws,r,labels):
    for i,h in enumerate(labels):
        c=ws.cell(row=r,column=i+1,value=h); c.font=H; c.fill=HF; c.alignment=Alignment(horizontal="center")

# ================= 前提条件 =================
a=wb.active; a.title="前提条件"
a["A1"]="カンテサンス CF計画（2027〜2032年・三つ星維持ケース）前提条件"; a["A1"].font=T
a["A2"]="単位：百万円。青字＝入力値（変更可）、黒字＝計算式。黄色セルは主要な仮定です。"; a["A2"].font=NOTE
header(a,4,["項目","FY2026／固定値"]+YRS+["出所・根拠"])
def src(r,t): a.cell(row=r,column=9,value=t).font=NOTE
def lab(r,t,bold=False): a.cell(row=r,column=1,value=t).font=B if bold else BLK
def setc(cell,v,fmt=NUM,key=False):
    c=a[cell]; c.value=v; c.number_format=fmt
    c.font=BLK if isinstance(v,str) and v.startswith("=") else BLUE
    if key: c.fill=YF

lab(5,"■ IM 42ページ「監修後）売上横ばいケース」の計画（2029年まで）",True)
im={6:("売上高(店舗)",949,[996,1029,1062]),7:("売上高(ガトー)",66,[66,66,66]),8:("売上原価",217,[225,232,240]),
    9:("販売費及び一般管理費",269,[282,287,291]),10:("営業利益",528,[556,576,598])}
for r,(t,b,vals) in im.items():
    lab(r,t); setc(f"B{r}",b)
    for col,v in zip(IMY,vals): setc(f"{col}{r}",v)
    src(r,"IM（2026年9月）42ページ 計画値")
lab(11,"売上高 合計")
for col,v in zip(["B"]+IMY,[1015,1062,1095,1129]): setc(f"{col}11",v)
src(11,"IM（2026年9月）42ページ 計画値（店舗＋ガトーの内訳は端数処理のため合計と1百万円ずれる年あり）")

lab(13,"■ 2030年以降の成長率（2026→2029年の年平均成長率を延長）",True)
lab(14,"売上高の年平均成長率（2026→2029）"); setc("B14","=(E11/B11)^(1/3)-1",PCT); src(14,"計算：IM計画の2026→2029年の年平均")
lab(15,"営業利益の年平均成長率（2026→2029）"); setc("B15","=(E10/B10)^(1/3)-1",PCT); src(15,"計算：IM計画の2026→2029年の年平均")
lab(16,"売上高の成長率（2030〜2032年に適用）")
for col in EXT: setc(f"{col}16","=$B$14",PCT,key=True)
src(16,"既定値＝上記の年平均。数値を直接入力して上書きできます")
lab(17,"営業利益の成長率（2030〜2032年に適用）")
for col in EXT: setc(f"{col}17","=$B$15",PCT,key=True)
src(17,"既定値＝上記の年平均。数値を直接入力して上書きできます")

lab(19,"■ 当社取得後の調整（IMの営業利益は岸田氏の現行報酬控除後）",True)
lab(20,"岸田氏の現行報酬（IM計画に含まれる）"); setc("B20",102.6,NUM1); src(20,"月額60万円×12＋事前確定届出給与95.4百万円（IM計画は現行報酬を前提＝売主QA②）")
lab(21,"岸田氏の役員報酬（取得後）")
for col,v in zip(YC,[20,20,20,0,0,0]): setc(f"{col}21",v,key=True)
src(21,"条件合意：ロックアップ3年（B案）は年2,000万円、以降は出勤義務なし")
lab(22,"後継料理長の追加人件費")
for col,v in zip(YC,[0,15,15,15,15,15]): setc(f"{col}22",v)
src(22,"仮定：2028年から年15百万円")
lab(23,"AM管理報酬")
for col in YC: setc(f"{col}23",61)
src(23,"SPC出資総額30.6億円×2%（AM契約で確定）")
lab(24,"岸田氏法人へのコンサルティング費（アーンアウト）")
for col,v in zip(YC,[0,0,0,400,400,0]): setc(f"{col}24",v,key=True)
src(24,"ユーザー指定：4年目・5年目に三つ星維持で各4億円、対象会社の費用として支払う")

lab(26,"■ 税・投資・資金",True)
rows=[(27,"実効税率",0.344,PCT,"2025年実績（銀行提出資料と同じ）"),
      (28,"減価償却費（年）",0,NUM,"IM計画に合わせて0（設備投資の償却による節税は保守的に考慮しない）"),
      (29,"設備投資（年）",30,NUM,"仮定（IMは計画なし。銀行提出資料と同じ保守的な水準）"),
      (30,"決済時の留保現金",100,NUM,"銀行提出資料：運転資金として1.0億円を留保"),
      (31,"FY2026 未払法人税等（退職金控除前）",114,NUM,"銀行提出資料：2026年12月期の未払税金1.14億円"),
      (32,"決済時の役員退職金（対象会社が支給）",225,NUM,"条件合意：前払い2.25億円を退職金に振替（ユーザー指定）"),
      (34,"借入額（ケースA）",1500,NUM,"ユーザー指定：15億円・7年"),
      (35,"金利",0.04,PCT,"ユーザー指定：金利上昇を含め4%"),
      (37,"借入額（ケースB）",2000,NUM,"ユーザー指定：20億円・7年"),
      (36,"返済年数",7,'0',"仮定（みずほ銀行と協議中）")]
for r,t,v,fmt,s_ in rows:
    lab(r,t); setc(f"B{r}",v,fmt,key=(r in (32,34,35,37))); src(r,s_)
lab(33,"■ 買収借入（合併後の会社が7年元利均等で返済）",True)
lab(39,"■ 子会社 カンテサンスプラス株式会社（プティ・ボノム100%子会社）",True)
lab(40,"子会社の現預金"); setc("B40",42.9,NUM1,key=True); src(40,"第10期（2026年6月期）決算書：42,906,344円")
lab(41,"子会社の役員借入金（岸田様への返済）"); setc("B41",5.5,NUM1); src(41,"第10期（2026年6月期）決算書：5,543,167円。決済までに返済する前提で控除")
lab(42,"決済時に使える子会社の現預金"); setc("B42","=B40-B41",NUM1); src(42,"配当（完全子法人からの配当は益金不算入・源泉徴収なし）又は合併で親会社に移す。年間の利益（約4百万円）はCFに含めず保守的に試算")
a.column_dimensions["A"].width=50; a.column_dimensions["B"].width=16; a.column_dimensions["I"].width=66
for col in YC: a.column_dimensions[col].width=10
a.freeze_panes="B5"

def build_cf(NAME,SUF,LOAN,LOANABS):
    global s
    # ================= CF計画 =================
    s=wb.create_sheet(NAME)
    s["A1"]=f"カンテサンス CF計画（2027〜2032年・三つ星維持ケース）{SUF}"; s["A1"].font=T
    s["A2"]="単位：百万円。2029年まではIM 42ページの計画値、2030年以降はIMの年平均成長率で延長。緑字＝前提条件シートの参照。"; s["A2"].font=NOTE
    header(s,4,["項目","FY2026（参考）"]+YRS)
    P="前提条件!"
    def put(r,label,bform,form,fmt=NUM,bold=False,fill=False):
        s.cell(row=r,column=1,value=label).font=B if bold else BLK
        for col,prev in zip(["B"]+YC,["A"]+["B"]+YC[:-1]):
            f=bform if col=="B" else form(col,prev)
            if f is None: continue
            c=s[f"{col}{r}"]; c.value=f; c.number_format=fmt
            direct=isinstance(f,str) and re.fullmatch(r"=前提条件!\$?[A-Z]+\$?\d+",f) is not None
            c.font=Font(name=F,size=10,bold=bold,color="008000" if direct else "000000")
            if fill: c.fill=TF
        if fill: s[f"A{r}"].fill=TF

    ext=lambda c: c in EXT
    put(5,"売上高（IM計画・2030年以降は延長）",f"={P}B11",lambda c,p: f"={p}5*(1+{P}{c}16)" if ext(c) else f"={P}{c}11",bold=True,fill=True)
    put(6,"　売上高成長率",None,lambda c,p:f"={c}5/{p}5-1",PCT)
    put(7,"営業利益（IM計画・岸田氏の現行報酬控除後）",f"={P}B10",lambda c,p: f"={p}7*(1+{P}{c}17)" if ext(c) else f"={P}{c}10",bold=True)
    put(8,"　営業利益率",f"=B7/B5",lambda c,p:f"=IF({c}5=0,0,{c}7/{c}5)",PCT)
    s["A9"]="■ 当社取得後の調整"; s["A9"].font=B
    put(10,"＋ 岸田氏の現行報酬の戻し",None,lambda c,p:f"={P}$B$20",NUM1)
    put(11,"− 岸田氏の役員報酬（取得後）",None,lambda c,p:f"=-{P}{c}21")
    put(12,"− 後継料理長の追加人件費",None,lambda c,p:f"=-{P}{c}22")
    put(13,"− AM管理報酬",None,lambda c,p:f"=-{P}{c}23")
    put(14,"− 岸田氏法人へのコンサルティング費",None,lambda c,p:f"=-{P}{c}24")
    put(15,"当社想定のEBITDA",None,lambda c,p:f"=SUM({c}7,{c}10:{c}14)",NUM1,bold=True,fill=True)
    put(16,"減価償却費",None,lambda c,p:f"={P}$B$28")
    put(17,"営業利益（当社想定）",None,lambda c,p:f"={c}15-{c}16",NUM1,bold=True)
    put(18,"支払利息",None,lambda c,p:f"={c}35*{P}$B$35",NUM1)
    put(19,"税引前利益",None,lambda c,p:f"={c}17-{c}18",NUM1)
    put(20,"法人税等",None,lambda c,p:f"=MAX(0,{c}19)*{P}$B$27",NUM1)
    put(21,"当期純利益",None,lambda c,p:f"={c}19-{c}20",NUM1,bold=True,fill=True)

    s["A23"]="■ キャッシュフロー"; s["A23"].font=B
    put(24,"当期純利益",None,lambda c,p:f"={c}21",NUM1)
    put(25,"減価償却費",None,lambda c,p:f"={c}16")
    put(26,"設備投資",None,lambda c,p:f"=-{P}$B$29")
    put(27,"FY2026分 法人税の納付（退職金の損金算入後）",None,lambda c,p:(f"=-({P}B31-{P}B32*{P}B27)" if c=="C" else 0),NUM1)
    put(28,"営業CF（借入返済前）",None,lambda c,p:f"=SUM({c}24:{c}27)",NUM1,bold=True)
    put(29,"借入元本の返済",None,lambda c,p:f"=-{c}38",NUM1)
    put(30,"借入返済後のCF",None,lambda c,p:f"={c}28+{c}29",NUM1,bold=True,fill=True)
    put(31,"期末現金残高（子会社の現預金を含む）",None,lambda c,p:(f"={P}B30+{P}B42+{c}30" if c=="C" else f"={p}31+{c}30"),NUM1,bold=True)

    s["A34"]="■ 買収借入（7年元利均等）"; s["A34"].font=B
    put(35,"期首残高",None,lambda c,p:(f"={P}{LOAN}" if c=="C" else f"={p}39"),NUM1)
    put(36,"元利返済額",None,lambda c,p:f"=IF({c}35<=0.001,0,-PMT({P}$B$35,{P}$B$36,{P}{LOANABS}))",NUM1)
    put(37,"　うち利息",None,lambda c,p:f"={c}18",NUM1)
    put(38,"　うち元本",None,lambda c,p:f"=MIN({c}35,{c}36-{c}37)",NUM1)
    put(39,"期末残高",None,lambda c,p:f"={c}35-{c}38",NUM1,bold=True)
    put(40,"DSCR（返済余力）",None,lambda c,p:f"=IF({c}36=0,0,({c}15-{c}20-{P}$B$29)/{c}36)",X,bold=True,fill=True)
    s["A41"]="DSCR＝（EBITDA−法人税等−設備投資）÷元利返済額。1.0倍未満の年は、それまでに積み上がった現金で返済します。"; s["A41"].font=NOTE

    s["A43"]="■ 損金算入による節税効果（ご参考）"; s["A43"].font=B
    put(44,"決済時の役員退職金（2026年度に損金算入）",None,lambda c,p:(f"={P}B32*{P}B27" if c=="C" else 0),NUM1)
    put(45,"コンサルティング費（アーンアウト）",None,lambda c,p:f"=-{c}14*{P}$B$27",NUM1)
    put(46,"節税効果 合計",None,lambda c,p:f"={c}44+{c}45",NUM1,bold=True)
    put(47,"節税効果 累計",None,lambda c,p:(f"={c}46" if c=="C" else f"={p}47+{c}46"),NUM1,bold=True,fill=True)
    s["A48"]="株式対価で支払った場合はいずれも損金にならず、この節税効果は生じません。退職金の効果は2026年度の税額減少として、2027年の納付時に現金化します。"; s["A48"].font=NOTE
    s.column_dimensions["A"].width=48
    for col in ["B"]+YC: s.column_dimensions[col].width=12
    s.freeze_panes="B5"

    return s

build_cf("CF_借入15億","｜借入15億円・7年・金利4%","B34","$B$34")
build_cf("CF_借入20億","｜借入20億円・7年・金利4%","B37","$B$37")

# ================= サマリー =================
m=wb.create_sheet("サマリー",0)
m["A1"]="カンテサンス CF計画 サマリー（三つ星維持ケース・IM計画ベース）"; m["A1"].font=T
m["A2"]="単位：百万円。数値はすべて各CFシートからの参照です。金利4%・7年元利均等。"; m["A2"].font=NOTE
items=[("売上高",5,NUM,True),("営業利益（IM計画）",7,NUM,True),("当社想定のEBITDA",15,NUM,True),("当期純利益",21,NUM,True),
       ("借入返済後のCF",30,NUM,True),("期末現金残高",31,NUM,False),("借入金期末残高",39,NUM,False),("DSCR",40,X,False)]
row0=4
for sheet,title in (("CF_借入15億","ケースA：借入15億円・7年"),("CF_借入20億","ケースB：借入20億円・7年")):
    m.cell(row=row0,column=1,value=title).font=Font(name=F,size=11,bold=True,color="6D2E46")
    header(m,row0+1,["指標"]+YRS+["6年計"])
    for k,(t,r,fmt,tot) in enumerate(items):
        rr=row0+2+k; m.cell(row=rr,column=1,value=t).font=BLK
        for j,col in enumerate(YC):
            c=m.cell(row=rr,column=2+j,value=f"='{sheet}'!{col}{r}"); c.font=GRN; c.number_format=fmt
        if tot: c=m.cell(row=rr,column=8,value=f"=SUM(B{rr}:G{rr})"); c.font=B; c.number_format=fmt
    rc=row0+2+len(items)
    cash_r=row0+2+5; dscr_r=row0+2+7
    m.cell(row=rc,column=1,value="最低DSCR").font=B
    c=m.cell(row=rc,column=2,value=f"=MIN(B{dscr_r}:G{dscr_r})"); c.number_format=X; c.font=B
    m.cell(row=rc+1,column=1,value="期中の最低現金残高").font=B
    c=m.cell(row=rc+1,column=2,value=f"=MIN(B{cash_r}:G{cash_r})"); c.number_format=NUM; c.font=B
    m.cell(row=rc+2,column=1,value="資金不足を避けるために必要な追加の留保現金").font=B
    c=m.cell(row=rc+2,column=2,value=f"=MAX(0,-B{rc+1})"); c.number_format=NUM; c.font=B; c.fill=YF
    row0=rc+4
m.cell(row=row0,column=1,value="売上高の年平均成長率（2030年以降に適用）").font=BLK
c=m.cell(row=row0,column=2,value="=前提条件!B14"); c.number_format=PCT; c.font=GRN
m.cell(row=row0+1,column=1,value="営業利益の年平均成長率（2030年以降に適用）").font=BLK
c=m.cell(row=row0+1,column=2,value="=前提条件!B15"); c.number_format=PCT; c.font=GRN
m.cell(row=row0+2,column=1,value="節税効果の6年累計（退職金・コンサル費）").font=BLK
c=m.cell(row=row0+2,column=2,value="='CF_借入20億'!H47"); c.number_format=NUM; c.font=GRN
NOTE_ROW=row0+4
notes=["【売上・営業利益】2027〜2029年はIM（2026年9月）42ページ「監修後）売上横ばいケース」の計画値（値上げ・サービス料12%化を織り込み済み）。2030〜2032年は、2026→2029年の年平均成長率（売上・営業利益それぞれ）で延長。",
 "【岸田様へのお支払い】決済時に役員退職金2.25億円（対象会社が支給・2026年度に損金算入）。2030年・2031年に岸田様の法人へコンサルティング費として各4億円（対象会社の費用）。IMの営業利益は岸田氏の現行報酬（約1.03億円）控除後のため、取得後の報酬との差額を調整。",
 "【税務上の留意点】退職金は賞与を含む実質月額での算定となるため否認リスクあり。コンサルティング費は役務提供の実態と対価の相当性が必要で、実態が伴わない場合は寄附金・株式対価の後払いと認定され損金不算入となるリスクが高い。双方の税理士の確認が前提。",
 "【子会社】カンテサンスプラスの現預金（役員借入金返済後 約0.37億円）を決済時の手元資金に加算。ブリッジ返済原資21.4億円（プティ・ボノム単体の現預金17.4＋役員貸付金5.0−留保1.0）には含めていないため、SPCの必要資金を減らす方に使うこともできる。",
 "【モデルの簡略化】運転資本の増減は無視。法人税は発生年度に納付（FY2026分のみ2027年に納付）。欠損金の繰越は考慮せず。消費税は除く。買収借入（15億円または20億円・7年・金利4%）は合併後の会社が返済する前提（WineBankで借りる形の場合は、借入返済後のCFが配当原資）。"]
for i,t in enumerate(notes):
    c=m.cell(row=NOTE_ROW+i,column=1,value=t); c.font=Font(name=F,size=9); c.alignment=Alignment(wrap_text=True,vertical="top")
    m.merge_cells(start_row=NOTE_ROW+i,start_column=1,end_row=NOTE_ROW+i,end_column=8); m.row_dimensions[NOTE_ROW+i].height=42
m.column_dimensions["A"].width=40
for col in "BCDEFGH": m.column_dimensions[col].width=11
wb.save(OUT); print("saved")
