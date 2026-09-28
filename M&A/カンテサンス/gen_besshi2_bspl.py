# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
F='BIZ UDPゴシック'
def k(y): return f"{round(y/1000):,}" if y>=0 else f"△{round(-y/1000):,}"
d=Document()
for s in d.sections: s.top_margin=s.bottom_margin=Cm(1.6); s.left_margin=s.right_margin=Cm(2.0)
st=d.styles['Normal']; st.font.name=F; st.font.size=Pt(9.5); st.element.rPr.rFonts.set(qn('w:eastAsia'),F)
st.paragraph_format.space_after=Pt(2)
def P(t,sz=9.5,b=False,al=None,after=2,before=0):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    if al: p.alignment=al
    r=p.add_run(t); r.font.size=Pt(sz); r.bold=b; r.font.name=F; r._element.rPr.rFonts.set(qn('w:eastAsia'),F)
    return p
def shade(c,fill):
    tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
def put(c,t,b=False,right=False,sz=9):
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    if right: p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run(t); r.font.size=Pt(sz); r.bold=b; r.font.name=F; r._element.rPr.rFonts.set(qn('w:eastAsia'),F)
TOT=('資産合計','負債・純資産合計','負債合計','純資産合計')
def bs(title, A, L):
    P(title,sz=10.5,b=True,before=4,after=3)
    n=max(len(A),len(L))
    A=A[:-1]+[None]*(n-len(A))+[A[-1]]; L=L[:-1]+[None]*(n-len(L))+[L[-1]]
    t=d.add_table(rows=n+1,cols=4); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    W=[Cm(5.0),Cm(3.2),Cm(5.0),Cm(3.2)]
    for j,c in enumerate(t.columns): c.width=W[j]
    for j,h in enumerate(['資産の部','金額','負債及び純資産の部','金額']):
        c=t.cell(0,j); put(c,h,b=True,right=(j%2==1)); shade(c,'F2F2F2')
    for i in range(n):
        for side,data in ((0,A),(2,L)):
            if data[i] is not None:
                nm,v,b=data[i]
                put(t.cell(i+1,side),nm if b else '　'+nm,b=b); put(t.cell(i+1,side+1),k(v),b=b,right=True)
                if nm in TOT: shade(t.cell(i+1,side),'F2F2F2'); shade(t.cell(i+1,side+1),'F2F2F2')
    for row in t.rows:
        for j,c in enumerate(row.cells): c.width=W[j]
def pl(title, rows):
    P(title,sz=10.5,b=True,before=8,after=3)
    t=d.add_table(rows=len(rows)+1,cols=2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Cm(9.0); t.columns[1].width=Cm(4.0)
    for j,h in enumerate(['科目','金額']):
        c=t.cell(0,j); put(c,h,b=True,right=(j==1)); shade(c,'F2F2F2')
    for i,(nm,v,b) in enumerate(rows):
        put(t.cell(i+1,0),nm if b else '　'+nm,b=b); put(t.cell(i+1,1),k(v),b=b,right=True)
    for row in t.rows: row.cells[0].width=Cm(9.0); row.cells[1].width=Cm(4.0)

# ===== 1. 株式会社WineBank 第54期
P('別紙2',sz=10,b=True)
P('貸借対照表及び損益計算書（抜粋）',sz=14,b=True,after=4)
P('1．株式会社WineBank　第54期（自 2024年10月1日　至 2025年9月30日）　単体',sz=10.5,b=True,after=1)
P('（単位：千円）',sz=8.5,al=WD_ALIGN_PARAGRAPH.RIGHT,after=0)
bs('(1) 貸借対照表（2025年9月30日現在）',
 [('【流動資産】',611624360,True),('現金及び預金',24390278,False),('売掛金',66464082,False),('商品',504638847,False),
  ('その他',611624360-24390278-66464082-504638847,False),
  ('【固定資産】',271745230,True),('有形固定資産',98991178,False),('無形固定資産',129897908,False),('投資その他の資産',42856144,False),
  ('資産合計',883369590,True)],
 [('【流動負債】',247844701,True),('買掛金',37602169,False),('短期借入金',162440980,False),('その他',247844701-37602169-162440980,False),
  ('【固定負債】',423769200,True),('長期借入金',423623000,False),('役員借入金',146200,False),
  ('負債合計',671613901,True),
  ('【純資産】',211755689,True),('資本金',10000000,False),('資本剰余金',266519000,False),('利益剰余金',-64763311,False),
  ('純資産合計',211755689,True),('負債・純資産合計',883369590,True)])
pl('(2) 損益計算書（自 2024年10月1日　至 2025年9月30日）',
 [('売上高',752864901,True),('売上原価',550916268,False),('売上総利益',201948633,True),('販売費及び一般管理費',199438126,False),
  ('営業利益',2510507,True),('営業外収益',1680098,False),('営業外費用',10003704,False),('経常損失',-5813099,True),
  ('特別利益',1839943,False),('特別損失（※2）',123391056,False),('税引前当期純損失',-127364212,True),
  ('法人税、住民税及び事業税',463105,False),('当期純損失',-127827317,True)])
P('',after=2)
for x in ['※1　第54期決算報告書から主要科目を抜粋し、千円未満を四捨五入して表示しております。このため、合計が内訳の和と一致しない場合がございます。',
 '※2　特別損失123,391千円の内訳は、事業譲渡損59,842千円、固定資産除却損18,102千円、抱合せ株式消滅差損45,447千円であり、いずれも事業譲渡および合併に伴う一時的な損失です。',
 '※3　（参考）営業利益に減価償却費および繰延資産償却を加えた金額は30,954千円です。',
 '※4　資本金および資本剰余金は2025年9月30日時点の金額です。期末後に増資を行っており、現在の資本金（資本準備金を含む）は4億3,716万9千円です。']: P(x,sz=8.5,after=2)

# ===== 2. （参考）株式会社アピシウス 第2期
p=d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
P('2．（参考）株式会社アピシウス　第2期（自 2025年1月1日　至 2025年12月31日）',sz=10.5,b=True,after=1)
P('当社グループが2024年6月に株式の100%を取得し、東京・銀座でグランメゾン「アピシウス」を運営する会社です。本件の承継モデルの参考として掲載いたします。',sz=9,after=1)
P('（単位：千円）',sz=8.5,al=WD_ALIGN_PARAGRAPH.RIGHT,after=0)
bs('(1) 貸借対照表（2025年12月31日現在）',
 [('【流動資産】',478953815,True),('現金及び預金',57372421,False),('売掛金',38018698,False),('商品',375793137,False),
  ('その他',478953815-57372421-38018698-375793137,False),
  ('【固定資産】',178950903,True),('有形固定資産',11097948,False),('投資その他の資産',167852955,False),
  ('資産合計',657904718,True)],
 [('【流動負債】',92523770,True),('買掛金',27150613,False),('未払金',25913811,False),('未払法人税等',28045700,False),('その他',92523770-27150613-25913811-28045700,False),
  ('【固定負債】',47161940,True),('退職給付引当金',47161940,False),
  ('負債合計',139685710,True),
  ('【純資産】',518219008,True),('資本金',10000000,False),('資本剰余金',458138337,False),('利益剰余金',50080671,False),
  ('純資産合計',518219008,True),('負債・純資産合計',657904718,True)])
pl('(2) 損益計算書（自 2025年1月1日　至 2025年12月31日）',
 [('売上高',705063048,True),('売上原価',203987046,False),('売上総利益',501076002,True),('販売費及び一般管理費',418247975,False),
  ('営業利益',82828027,True),('営業外収益',1668100,False),('経常利益',84496127,True),
  ('特別利益（※6）',5904192,False),('特別損失（※6）',8072862,False),('税引前当期純利益',82327457,True),
  ('法人税等',28058815,False),('当期純利益',54268642,True)])
P('',after=2)
for x in ['※5　第2期決算報告書から主要科目を抜粋し、千円未満を四捨五入して表示しております。借入金はございません。',
 '※6　特別利益および特別損失は、いずれも前期損益修正によるものです。',
 '※7　（参考）営業利益に減価償却費を加えた金額は84,463千円です。']: P(x,sz=8.5,after=2)
d.save('別紙2_貸借対照表及び損益計算書_抜粋_WineBank_2025年9月期.docx'); print('ok')
