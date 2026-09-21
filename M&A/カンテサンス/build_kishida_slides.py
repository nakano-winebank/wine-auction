# -*- coding: utf-8 -*-
import copy
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY="30373F"; NAVY2="36394B"; COPPER="C68A65"; CARD="F6F6F6"; GHOST="D8D8D8"
BLACK="000000"; WHITE="FFFFFF"; GREY="595959"; SAND="EAD5C7"
F="Meiryo"

pr=Presentation("deck.pptx")
BLANK=pr.slides[15].slide_layout

def _ea(run):
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:ea","a:cs","a:sym"):
        ns="{http://schemas.openxmlformats.org/drawingml/2006/main}"+tag.split(":")[1]
        el=rPr.find(ns)
        if el is None:
            from lxml import etree
            el=etree.SubElement(rPr,ns)
        el.set("typeface",F)

def tb(sl,x,y,w,h,parts,size=10.5,bold=False,color=NAVY,align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP,space=None,wrap=True):
    """parts: str | list of (text,{opts}) | list of lines (each a list of runs)"""
    sp=sl.shapes.add_textbox(I(x),I(y),I(w),I(h))
    tf=sp.text_frame; tf.word_wrap=wrap; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    if isinstance(parts,str): parts=[[(parts,{})]]
    if parts and isinstance(parts[0],tuple): parts=[parts]
    first=True
    for line in parts:
        p=tf.paragraphs[0] if first else tf.add_paragraph()
        first=False
        p.alignment=align
        if space: p.line_spacing=space
        for t,o in line:
            r=p.add_run(); r.text=t
            r.font.name=F; _ea(r)
            r.font.size=Pt(o.get("sz",size))
            r.font.bold=o.get("b",bold)
            r.font.color.rgb=C.from_string(o.get("c",color))
    return sp

def rect(sl,x,y,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.RECTANGLE,adj=None):
    sp=sl.shapes.add_shape(shape,I(x),I(y),I(w),I(h))
    if fill: sp.fill.solid(); sp.fill.fore_color.rgb=C.from_string(fill)
    else: sp.fill.background()
    if line:
        sp.line.color.rgb=C.from_string(line); sp.line.width=Pt(lw)
    else: sp.line.fill.background()
    sp.shadow.inherit=False
    sp.text_frame.text=""
    return sp

def chrome(sl, page_title, band_text=None):
    rect(sl,0,5.482,10,0.143,fill=NAVY)
    src=pr.slides[15]
    for s in src.shapes:                      # CONFIDENTIAL badge
        if s.has_text_frame and s.text_frame.text.strip()=="CONFIDENTIAL":
            sl.shapes._spTree.append(copy.deepcopy(s._element)); break
    tb(sl,0.262,0.314,9.4,0.30,page_title,size=14,bold=True,color=BLACK)
    if band_text:
        b=rect(sl,0,0.737,10,0.359,fill=NAVY2)
        tf=b.text_frame; tf.word_wrap=True
        tf.margin_left=tf.margin_right=0; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=band_text
        r.font.name=F; _ea(r); r.font.size=Pt(13); r.font.bold=True
        r.font.color.rgb=C.from_string(WHITE)

def add_after(idx):
    """new blank slide placed right after 0-based idx"""
    sl=pr.slides.add_slide(BLANK)
    xml=pr.slides._sldIdLst
    ids=list(xml)
    xml.remove(ids[-1]); xml.insert(idx+1, ids[-1])
    return sl

def pic(sl,path,x,y,w,h):
    """cover-crop insert"""
    from PIL import Image
    iw,ih=Image.open(path).size
    tr,sr=w/h, iw/ih
    p=sl.shapes.add_picture(path,I(x),I(y),I(w),I(h))
    if sr>tr:
        c=(1-tr/sr)/2; p.crop_left=c; p.crop_right=c
    else:
        c=(1-sr/tr)/2; p.crop_top=c; p.crop_bottom=c
    return p

# ===================== p.11  WineBank × アピシウス =====================
s=pr.slides[10]
for sh in list(s.shapes):                      # wipe, keep nothing
    sh._element.getparent().remove(sh._element)
chrome(s,"WineBank × アピシウス",
       "銀座の老舗グランメゾンを100％承継。2年で、人を替えずに伸ばしました")

kpi=[("売上高","＋30%","取得前比（2年間）"),
     ("営業利益","＋250%超","取得前比（2年間）"),
     ("幹部メンバー継続率","100%","一人も欠けていません")]
CX=[0.587,3.677,6.855]; CW=2.559
for i,(lab,val,note) in enumerate(kpi):
    x=CX[i]; on=(i==2)
    rect(s,x,1.22,CW,0.95,fill=NAVY2 if on else CARD)
    tb(s,x+0.20,1.33,CW-0.40,0.22,lab,size=9.5,bold=True,
       color=SAND if on else COPPER)
    tb(s,x+0.20,1.53,CW-0.40,0.40,val,size=24,bold=True,
       color=WHITE if on else NAVY2)
    tb(s,x+0.20,1.93,CW-0.40,0.20,note,size=8,color="C9B6A8" if on else GREY)

tb(s,0.587,2.32,5.2,0.24,"取得前と、取得から2年後",size=11,bold=True,color=BLACK)
rows=[("","取得前（〜2024年5月）","取得後2年（2026年）"),
      ("業績","39年間赤字。40年目に黒字化","売上＋30%／営業利益＋250%超"),
      ("人員","―","幹部メンバーは全員継続。入替えなし"),
      ("労働環境","サービス残業が常態化","サービス残業を撤廃。残業ゼロの体制へ"),
      ("店舗","―","一部改装・修繕を実施。業態は不変")]
TX=0.587; TW=[0.72,1.92,2.54]; TY=2.60; RH=0.40
for r,(a,b,c) in enumerate(rows):
    y=TY+r*RH
    hdr=(r==0)
    rect(s,TX,y,sum(TW),RH,fill=NAVY2 if hdr else (CARD if r%2 else WHITE),
         line=None if hdr else "E3E3E3",lw=0.75)
    for j,(txt,w) in enumerate(zip((a,b,c),TW)):
        cx=TX+sum(TW[:j])
        col=WHITE if hdr else (NAVY2 if j==0 else NAVY)
        hi=(not hdr and j==2)
        tb(s,cx+0.10,y+0.05,w-0.16,RH-0.10,txt,size=8.5,
           bold=hdr or j==0 or hi,color=COPPER if hi else col,
           anchor=MSO_ANCHOR.MIDDLE,space=1.05)

pic(s,"cuts/apx_main.png",5.95,2.60,3.46,0.86)
pic(s,"cuts/apx_private.png",5.95,3.52,3.46,0.86)
tb(s,5.95,4.44,3.46,0.20,"アピシウス（東京・銀座）　1983年開業",size=7.5,color=GREY)

rect(s,0.587,4.68,8.826,0.54,fill=CARD)
rect(s,0.587,4.68,0.05,0.54,fill=COPPER)
tb(s,0.76,4.76,8.60,0.42,
   [[("グランメゾンは、引き継いでも壊れない。",{"sz":11,"b":True,"c":BLACK})],
    [("私たちが学んだのは収益改善の手法ではありません。人を入れ替えず、労働環境をむしろ良くしたうえで、店は伸ばせるということです。",
      {"sz":8.5,"c":NAVY})]],space=1.25)

# ===================== new p.18  国内オーベルジュ構想 ① =====================
a=add_after(16)
chrome(a,"国内オーベルジュ構想",
       "日本のワイン産地に、ファインダイニングを主役とした小さな滞在施設を")
cards=[("料理","ファインダイニング","cuts/apx_main.png",
        [("その土地の生産者と",{}),("同じテーブルにつける距離",{"b":True,"c":COPPER}),
         ("で、一日十数名だけをお迎えする。",{})]),
       ("ワイン","産地とセラー","x/ppt/media/image10.jpg",
        [("ワイン産地に置くことで、",{}),("造り手と客が直接つながる場",{"b":True,"c":COPPER}),
         ("をつくる。当社のセラー機能を併設する。",{})]),
       ("滞在","小規模宿泊","cuts/apx_salon.png",
        [("客室は10室前後。夜を跨いでいただくことで、",{}),
         ("ワインリストの幅が一気に広がる",{"b":True,"c":COPPER}),("。",{})])]
for i,(lab,ttl,img,body) in enumerate(cards):
    x=CX[i]
    rect(a,x,1.24,CW,3.28,fill=CARD)
    pic(a,img,x+0.12,1.36,CW-0.24,1.12)
    rect(a,x+0.12,2.60,0.74,0.24,fill=COPPER)
    tb(a,x+0.12,2.60,0.74,0.24,lab,size=9,bold=True,color=WHITE,
       align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    tb(a,x+0.12,2.94,CW-0.24,0.30,ttl,size=14,bold=True,color=NAVY2)
    tb(a,x+0.20,3.34,CW-0.36,1.05,body,size=10,color=NAVY,space=1.35)

rect(a,0.587,4.62,8.826,0.52,fill=NAVY2)
tb(a,0.76,4.68,8.50,0.42,
   [[("2027年4月 ブラッスリー六本木 ▶ 2028年 札幌・大阪 ▶ 2029年 首都圏 ▶ ",{"sz":9,"c":"C9CCD6"}),
     ("2030年〜 リゾート含め年2店舗",{"sz":9,"b":True,"c":SAND}),("。",{"sz":9,"c":"C9CCD6"})],
    [("この先に、オーベルジュがあります。",{"sz":9,"b":True,"c":SAND})]],
   align=PP_ALIGN.CENTER,space=1.2)
tb(a,0.587,5.22,4.00,0.18,
   "※ 写真は当社運営店舗。オーベルジュのイメージです。",size=7,color=GREY)

# ===================== new p.19  国内オーベルジュ構想 ② =====================
b=add_after(17)
chrome(b,"国内オーベルジュ構想　―　なぜ、当社が担えるのか",
       "グランメゾンは、引き継いでも壊れない ― アピシウスで実証しました")
pic(b,"cuts/apx_private.png",0.587,1.24,3.36,2.38)
tb(b,0.587,3.66,3.36,0.20,"アピシウス（東京・銀座）　2024年6月に100％承継",size=7.5,color=GREY)
rect(b,0.587,3.94,3.36,1.18,fill=CARD)
tb(b,0.75,4.06,3.04,0.94,
   [[("ゼロから新しい店をつくるより、",{"sz":9.5,"c":NAVY})],
    [("長い時間をかけて積み上げられたものを、",{"sz":9.5,"c":NAVY})],
    [("壊さずに次の世代へ渡す。",{"sz":10.5,"b":True,"c":COPPER})],
    [("私たちが得意なのは、そちらです。",{"sz":10.5,"b":True,"c":BLACK})]],space=1.3)

pts=[("人を替えずに、伸ばした",
      [[("幹部メンバーは全員継続、経営陣の入替えもなし。そのうえで",{})],
       [("売上＋30%／営業利益＋250%超",{"b":True,"c":COPPER}),("（2年間・取得前比）。",{})]]),
     ("働く環境は、むしろ良くした",
      [[("常態化していた",{}),("サービス残業を撤廃",{"b":True,"c":COPPER}),("。12月を除き、",{})],
       [("残業自体が発生しない体制に切り替えました。",{})]]),
     ("ワインで、店を強くした",
      [[("日本トップクラスのワイン在庫とソムリエに、",{})],
       [("インポーターである当社の仕入ルートが加わりました",{"b":True,"c":COPPER}),("。",{})]])]
for i,(ttl,body) in enumerate(pts):
    y=1.24+i*1.30
    rect(b,4.27,y,5.14,1.16,fill=CARD)
    tb(b,4.32,y-0.06,0.46,0.62,str(i+1),size=30,bold=True,color=GHOST)
    tb(b,4.92,y+0.14,4.36,0.26,ttl,size=12.5,bold=True,color=NAVY2)
    tb(b,4.92,y+0.48,4.36,0.58,body,size=9.5,color=NAVY,space=1.3)

pr.save("out.pptx"); print("SAVED", len(pr.slides.__iter__.__self__._sldIdLst))
