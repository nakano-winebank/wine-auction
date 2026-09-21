# -*- coding: utf-8 -*-
import copy
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

NAVY="30373F"; NAVY2="36394B"; COPPER="C68A65"; CARD="F6F6F6"
BLACK="000000"; WHITE="FFFFFF"; GREY="595959"; SAND="EAD5C7"; F="Meiryo"
A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
R="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

pr=Presentation("base.pptx")

def _ea(run):
    rPr=run._r.get_or_add_rPr()
    for tag in ("ea","cs","sym"):
        el=rPr.find(A+tag)
        if el is None: el=etree.SubElement(rPr,A+tag)
        el.set("typeface",F)

def tb(sl,x,y,w,h,parts,size=10.5,bold=False,color=NAVY,align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP,space=None):
    sp=sl.shapes.add_textbox(I(x),I(y),I(w),I(h))
    tf=sp.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    if isinstance(parts,str): parts=[[(parts,{})]]
    if parts and isinstance(parts[0],tuple): parts=[parts]
    first=True
    for line in parts:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.alignment=align
        if space: p.line_spacing=space
        for t,o in line:
            r=p.add_run(); r.text=t; r.font.name=F; _ea(r)
            r.font.size=Pt(o.get("sz",size)); r.font.bold=o.get("b",bold)
            r.font.color.rgb=C.from_string(o.get("c",color))
    return sp

def rect(sl,x,y,w,h,fill=None,line=None,lw=1.0):
    sp=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,I(x),I(y),I(w),I(h))
    if fill: sp.fill.solid(); sp.fill.fore_color.rgb=C.from_string(fill)
    else: sp.fill.background()
    if line: sp.line.color.rgb=C.from_string(line); sp.line.width=Pt(lw)
    else: sp.line.fill.background()
    sp.shadow.inherit=False
    return sp

# ---------- delete p.19 ----------
lst=pr.slides._sldIdLst
el=list(lst)[18]
pr.part.drop_rel(el.get(R+"id")); lst.remove(el)

# ---------- rebuild p.18 ----------
s=pr.slides[17]
for sh in list(s.shapes): sh._element.getparent().remove(sh._element)

rect(s,0,5.482,10,0.143,fill=NAVY)
for sh in pr.slides[15].shapes:                       # CONFIDENTIAL badge
    if sh.has_text_frame and sh.text_frame.text.strip()=="CONFIDENTIAL":
        s.shapes._spTree.append(copy.deepcopy(sh._element)); break
tb(s,0.262,0.314,9.4,0.30,"国内オーベルジュ構想",size=14,bold=True,color=BLACK)
b=rect(s,0,0.737,10,0.359,fill=NAVY2)
tf=b.text_frame; tf.word_wrap=True; tf.margin_left=tf.margin_right=0
tf.vertical_anchor=MSO_ANCHOR.MIDDLE
p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
r=p.add_run(); r.text="ワイナリーに隣接し、温泉を持つ。日本にまだ無いオーベルジュをつくります"
r.font.name=F; _ea(r); r.font.size=Pt(13); r.font.bold=True
r.font.color.rgb=C.from_string(WHITE)

CX=[0.587,3.677,6.855]; CW=2.559
cards=[("料理","一日十数名だけ","その土地の生産者と、同じテーブルに",False),
       ("ワイン","ワイナリー隣接","造り手の畑が、窓の外にある",False),
       ("滞在","温泉 ＋ 10室","食後に、移動しなくていい",True)]
for i,(lab,val,note,on) in enumerate(cards):
    x=CX[i]
    rect(s,x,1.22,CW,0.95,fill=NAVY2 if on else CARD)
    tb(s,x+0.20,1.33,CW-0.40,0.22,lab,size=9.5,bold=True,color=SAND if on else COPPER)
    tb(s,x+0.20,1.53,CW-0.40,0.34,val,size=18,bold=True,color=WHITE if on else NAVY2)
    tb(s,x+0.20,1.92,CW-0.40,0.20,note,size=8,color="C9B6A8" if on else GREY)

tb(s,0.587,2.30,6.0,0.24,"構想の骨子",size=11,bold=True,color=BLACK)
rows=[("","一般的なオーベルジュ","当社が考えるオーベルジュ"),
      ("立地","景勝地・観光地に単独で建つ","ワイナリーに隣接。造り手が徒歩圏にいる"),
      ("ワイン","地元ワインが中心","当社在庫32億円から、世界のファインワインを産地価格で"),
      ("滞在","宿泊機能のみ","温泉を併設。食後に移動しないから、もう一本開けられる"),
      ("運営","オーナーシェフの個人事業","グランメゾン運営会社が担う。料理人に経営を背負わせない")]
TX=0.587; TW=[1.10,2.90,4.826]; TY=2.58; RH=0.38
for i,(a_,b_,c_) in enumerate(rows):
    y=TY+i*RH; hdr=(i==0)
    rect(s,TX,y,sum(TW),RH,fill=NAVY2 if hdr else (CARD if i%2 else WHITE),
         line=None if hdr else "E3E3E3",lw=0.75)
    for j,(txt,w) in enumerate(zip((a_,b_,c_),TW)):
        cx=TX+sum(TW[:j])
        col=WHITE if hdr else (NAVY2 if j==0 else NAVY)
        hi=(not hdr and j==2)
        tb(s,cx+0.12,y+0.04,w-0.20,RH-0.08,txt,size=8.5,
           bold=hdr or j==0 or hi,color=COPPER if hi else col,
           anchor=MSO_ANCHOR.MIDDLE,space=1.05)

tb(s,0.587,4.58,8.826,0.20,
   [[("ブラッスリー出店計画（2027年4月 六本木 → 2030年〜 リゾート含め年2店舗）の先に着手。　",{"sz":8,"c":GREY}),
     ("候補地：北海道（余市・仁木）／長野（東御）／山梨（勝沼）",{"sz":8,"b":True,"c":NAVY2})]])

rect(s,0.587,4.82,8.826,0.46,fill=CARD)
rect(s,0.587,4.82,0.05,0.46,fill=COPPER)
tb(s,0.76,4.87,8.50,0.36,
   [[("料理人が、一日の終わりまでお客様と過ごせる場所を。",{"sz":11,"b":True,"c":BLACK})],
    [("ゼロから新しい店をつくるより、長い時間をかけて積み上げられたものを、壊さずに次の世代へ渡す。私たちが得意なのは、そちらです。",
      {"sz":8.5,"c":NAVY})]],space=1.2)

pr.save("out2.pptx"); print("SAVED", len(pr.slides._sldIdLst))
