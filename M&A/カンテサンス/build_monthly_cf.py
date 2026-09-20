# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb=openpyxl.Workbook(); F="Arial"
BLUE=Font(name=F,size=10,color="0000FF"); BLK=Font(name=F,size=10); BLD=Font(name=F,size=10,bold=True)
TTL=Font(name=F,size=13,bold=True,color="6D2E46"); SEC=Font(name=F,size=10,bold=True,color="FFFFFF")
GRY=Font(name=F,size=9,color="808080"); GRN=Font(name=F,size=10,color="008000")
RED=Font(name=F,size=10,bold=True,color="9B2C2C")
SECF=PatternFill("solid",fgColor="6D2E46"); YEL=PatternFill("solid",fgColor="FFFF00")
TINT=PatternFill("solid",fgColor="F7F3F4")
NUM='#,##0;(#,##0);-'; PCT='0.0%'

def sec(ws,row,text,span):
    c=ws.cell(row=row,column=1,value=text); c.font=SEC; c.fill=SECF
    for i in range(2,span+1): ws.cell(row=row,column=i).fill=SECF

# ===================== 前提 =====================
ws=wb.active; ws.title="前提"
ws["A1"]="カンテサンス 2027年 月次キャッシュフロー計画 ― 前提条件"; ws["A1"].font=TTL
ws["A2"]="単位：百万円。青字＝入力値、黒字＝計算式、黄色＝感応度を見るための主要レバー。"; ws["A2"].font=GRY

sec(ws,4,"1. 損益前提（2026年12月期の会社計画と同水準で横ばい）",4)
for i,(lab,val,fmt,note) in enumerate([
  ("売上高（年間）",1015,None,"IM p.40 の2026年12月期 会社計画"),
  ("営業利益（年間）",528,None,"同上。役員報酬・賞与を控除した後の金額"),
  ("変動費率（食材・ドリンク原価）",0.254,PCT,"ドリンク原価0.9万円/人×1.4万人と食材原価から逆算"),
  ("月次固定費",None,None,"＝（売上高×(1−変動費率)−営業利益）÷12。人件費・役員報酬・賃料等")]):
    r=5+i; ws.cell(row=r,column=1,value=lab).font=BLK
    c=ws.cell(row=r,column=2)
    if val is None: c.value="=(B5*(1-B7)-B6)/12"; c.font=BLK
    else: c.value=val; c.font=BLUE
    c.number_format=fmt or NUM
    ws.cell(row=r,column=4,value=note).font=GRY

sec(ws,10,"2. 月次の営業日数（年間252日）",14)
ws.cell(row=11,column=1,value="月").font=BLD; ws.cell(row=12,column=1,value="営業日数").font=BLD
for i,d in enumerate([18,20,22,21,21,22,22,17,21,22,21,25]):
    ws.cell(row=11,column=2+i,value=f"{i+1}月").font=BLD
    c=ws.cell(row=12,column=2+i,value=d); c.font=BLUE; c.number_format=NUM
ws.cell(row=11,column=14,value="合計").font=BLD
c=ws.cell(row=12,column=14,value="=SUM(B12:M12)"); c.font=BLD; c.number_format=NUM
ws.cell(row=13,column=1,value="※ 1月は年始休業、8月は夏季休業、12月は繁忙期のため営業日数を増減させている。月次売上は営業日数に比例させている。").font=GRY

sec(ws,15,"3. 税金",4)
for i,(lab,val,note) in enumerate([
  ("2026年12月期の未払法人税等・未払消費税（2027年2月納付）",114,"2026年12月期末BSの負債計上額。譲受価額の算定では控除していないため、取得後の会社が納付する"),
  ("法人税等 2027年度 中間納付（2027年8月）",92,"前事業年度の確定法人税額の2分の1。仮決算による中間申告を行えば圧縮できる"),
  ("（期中の消費税）",0,"売上高は税抜表示であり、消費税は預り金の性質を持つ。顧客からの預りと納付が年間で相殺されるため計上していない")]):
    r=16+i; ws.cell(row=r,column=1,value=lab).font=BLK
    c=ws.cell(row=r,column=2,value=val); c.font=BLUE; c.number_format=NUM
    ws.cell(row=r,column=4,value=note).font=GRY

sec(ws,20,"4. 買収ファイナンス（返済は毎月）",4)
for i,(lab,val,fmt,note) in enumerate([
  ("タームローンA（元金均等・7年）",1400,None,"毎月返済。月次元本＝1,400百万円÷84か月"),
  ("タームローンB（期限一括）",600,None,"7年後に一括返済。7年間の累積余剰キャッシュフローで返済する"),
  ("有利子負債 合計（2027年1月1日）",None,None,"＝タームローンA＋B。ブリッジローンは合併直後に全額返済済み"),
  ("金利",0.03,PCT,"みずほ銀行との協議前提。利息は前月末残高×金利÷12で毎月発生"),
  ("月次の元本返済額",None,None,"＝タームローンA÷84か月"),
  ("期首現金（対象会社に留保する運転資金）",100,None,"★ 本モデルの主要レバー。ここを増減させると買収時の中野氏の拠出額が同額だけ変動する")]):
    r=21+i; ws.cell(row=r,column=1,value=lab).font=BLK
    c=ws.cell(row=r,column=2)
    if lab.startswith("有利子負債"): c.value="=B21+B22"; c.font=BLK
    elif lab.startswith("月次の元本"): c.value="=B21/84"; c.font=BLK
    else: c.value=val; c.font=BLUE
    c.number_format=fmt or NUM
    if lab.startswith("期首現金"): c.fill=YEL; c.font=Font(name=F,size=10,color="0000FF",bold=True)
    ws.cell(row=r,column=4,value=note).font=GRY

sec(ws,28,"5. 資産調整勘定（税務上ののれん）",4)
for i,(lab,val,note) in enumerate([
  ("株式譲受代金",5000,"譲受価額50.0億円のケース"),
  ("対象会社の純資産（2026年12月期末・帳簿）",2315,"時価評価によりワイン在庫等の含み益が加算される場合、資産調整勘定はその分減少する"),
  ("資産調整勘定",None,"＝株式譲受代金−時価純資産。非適格合併により計上される税務上ののれん"),
  ("償却年数",5,"法人税法上、60か月にわたり月割で損金算入される"),
  ("年間の損金算入額",None,"＝資産調整勘定÷償却年数"),
  ("税務上の所得（年間）",None,"＝営業利益−年間の損金算入額。マイナスであれば法人税は発生しない")]):
    r=29+i; ws.cell(row=r,column=1,value=lab).font=BLK
    c=ws.cell(row=r,column=2)
    if lab=="資産調整勘定": c.value="=B29-B30"; c.font=BLK
    elif lab=="年間の損金算入額": c.value="=B31/B32"; c.font=BLK
    elif lab.startswith("税務上の所得"): c.value="=B6-B33"; c.font=RED
    else: c.value=val; c.font=BLUE
    c.number_format=NUM
    ws.cell(row=r,column=4,value=note).font=GRY
ws.cell(row=36,column=1,value="※ 資産調整勘定の計上可否および金額は、合併の適格・非適格の判定と時価純資産の算定によって変動する。デュー・ディリジェンスの税務パートで確認を要する。").font=GRY
ws.cell(row=37,column=1,value="※ 本シートを前提に、シート『月次CF（保守）』は損金算入を織り込まない場合、シート『月次CF（のれん考慮）』は織り込む場合を示している。").font=GRY

ws.column_dimensions["A"].width=46; ws.column_dimensions["B"].width=12
ws.column_dimensions["C"].width=2; ws.column_dimensions["D"].width=92
for i in range(3,15): ws.column_dimensions[get_column_letter(i)].width=8

# ===================== 月次CF（2パターン） =====================
def cfsheet(name, use_goodwill, lead):
    cf=wb.create_sheet(name)
    cf["A1"]=f"2027年 月次キャッシュフロー計画（百万円）― {name[5:]}"; cf["A1"].font=TTL
    cf["A2"]=lead; cf["A2"].font=GRY
    for r,t,b in [(4,"月番号",0),(5,"月",0),(6,"営業日数",0),(7,"売上高",0),(8,"変動費",0),(9,"固定費",0),(10,"営業利益",1),
                  (11,"（参考）資産調整勘定の償却　※非資金",0),(12,"（参考）税務上の所得",0),
                  (14,"営業キャッシュフロー",1),(15,"法人税等の納付",0),(16,"支払利息",0),(17,"借入元本の返済（毎月）",0),
                  (18,"当月収支",1),(20,"現金 期首残高",0),(21,"現金 期末残高",1),(23,"有利子負債（期末）",1)]:
        c=cf.cell(row=r,column=1,value=t); c.font=BLD if b else BLK
    sec(cf,3,"2027年",15)
    for i in range(12):
        col=2+i; L=get_column_letter(col); pl=get_column_letter(col-1)
        h=cf.cell(row=3,column=col,value=f"{i+1}月"); h.font=SEC; h.fill=SECF
        cf.cell(row=4,column=col,value=i+1).font=GRY
        cf.cell(row=5,column=col,value=f"{i+1}月").font=BLK
        cf[f"{L}6"]=f"=前提!{L}12"
        cf[f"{L}7"]=f"=前提!$B$5*{L}6/前提!$N$12"
        cf[f"{L}8"]=f"=-{L}7*前提!$B$7"
        cf[f"{L}9"]="=-前提!$B$8"
        cf[f"{L}10"]=f"=SUM({L}7:{L}9)"
        cf[f"{L}11"]="=-前提!$B$33/12"
        cf[f"{L}12"]=f"={L}10+{L}11"
        cf[f"{L}14"]=f"={L}10"
        if use_goodwill:
            cf[f"{L}15"]=f"=IF({L}4=2,-前提!$B$16,0)"
        else:
            cf[f"{L}15"]=f"=IF({L}4=2,-前提!$B$16,IF({L}4=8,-前提!$B$17,0))"
        base = "前提!$B$23" if i==0 else f"{pl}23"
        cf[f"{L}16"]=f"=-{base}*前提!$B$24/12"
        cf[f"{L}17"]="=-前提!$B$25"
        cf[f"{L}18"]=f"=SUM({L}14:{L}17)"
        cf[f"{L}20"]="=前提!$B$26" if i==0 else f"={pl}21"
        cf[f"{L}21"]=f"={L}20+{L}18"
        cf[f"{L}23"]=f"={base}+{L}17"
    h=cf.cell(row=3,column=14,value="合計"); h.font=SEC; h.fill=SECF
    for r in (6,7,8,9,10,11,12,14,15,16,17,18): cf[f"N{r}"]=f"=SUM(B{r}:M{r})"
    cf["N20"]="=B20"; cf["N21"]="=M21"; cf["N23"]="=M23"
    for r in [6,7,8,9,10,11,12,14,15,16,17,18,20,21,23]:
        for col in range(2,15):
            c=cf.cell(row=r,column=col); c.number_format=NUM
            if r in (10,18,21,23): c.font=BLD
            if r==21: c.fill=TINT
            if r in (11,12): c.font=GRY
    cf["A25"]="最低現金残高"; cf["A25"].font=BLD
    cf["B25"]="=MIN(B21:M21)"; cf["B25"].font=RED; cf["B25"].number_format=NUM
    cf["D25"]="＝ 月次固定費の"; cf["D25"].font=BLK
    cf["E25"]="=B25/前提!$B$8"; cf["E25"].font=BLD; cf["E25"].number_format='0.0"か月分"'
    cf["A26"]="年間の営業キャッシュフロー"; cf["B26"]="=N14"; cf["B26"].number_format=NUM
    cf["A27"]="年間の法人税等の納付"; cf["B27"]="=-N15"; cf["B27"].number_format=NUM
    cf["A28"]="年間の元利返済額"; cf["B28"]="=-N16-N17"; cf["B28"].number_format=NUM
    cf["A29"]="DSCR"; cf["A29"].font=BLD
    cf["B29"]="=(N14+N15)/B28"; cf["B29"].number_format='0.00"倍"'; cf["B29"].font=Font(name=F,size=10,bold=True,color="008000")
    for r in (26,27,28): cf.cell(row=r,column=1).font=BLK
    cf["A31"]="※ 借入の返済は毎月。元本は毎月16.7百万円（タームローンA 1,400百万円÷84か月）、利息は前月末残高×金利÷12。"; cf["A31"].font=GRY
    cf["A32"]="※ 減価償却費・設備投資・運転資本の増減はいずれも軽微なため、営業利益をそのまま営業キャッシュフローとみなしている。消費税は預り金の性質のため計上していない。"; cf["A32"].font=GRY
    ch=LineChart(); ch.title=f"2027年 現金残高の推移（百万円）― {name[5:]}"; ch.style=2; ch.height=8.4; ch.width=22
    data=Reference(cf,min_col=1,max_col=13,min_row=21,max_row=21)
    cats=Reference(cf,min_col=2,max_col=13,min_row=5,max_row=5)
    ch.add_data(data,titles_from_data=True,from_rows=True); ch.set_categories(cats)
    ch.series[0].graphicalProperties.line.solidFill="6D2E46"
    ch.series[0].graphicalProperties.line.width=28000; ch.series[0].smooth=False
    cf.add_chart(ch,"A34")
    cf.column_dimensions["A"].width=34
    for i in range(2,15): cf.column_dimensions[get_column_letter(i)].width=10
    return cf

cfsheet("月次CF（保守）",False,
  "資産調整勘定（税務上ののれん）の損金算入を織り込まない場合。2月に2026年12月期の未払税金、8月に2027年度の中間納付が発生する。")
cfsheet("月次CF（のれん考慮）",True,
  "資産調整勘定の損金算入を織り込む場合。税務上の所得がマイナスとなるため、仮決算による中間申告により8月の中間納付を回避できる前提。2月の納付は2026年12月期に対応するものであり回避できない。")

# ===================== 買収資金 =====================
fz=wb.create_sheet("買収資金")
fz["A1"]="買収資金の調達と使途（譲受価額50.0億円のケース）"; fz["A1"].font=TTL
fz["A2"]="『前提』シートの期首現金（留保する運転資金）を変更すると、中野氏の必要拠出額が自動で再計算される。"; fz["A2"].font=GRY
sec(fz,4,"1. 資金使途（クロージング時）",3)
fz["A5"]="株式譲受代金"; fz["B5"]="=前提!$B$29"; fz["B5"].font=GRN
fz["A6"]="取得関連費用"; fz["B6"]=150; fz["B6"].font=BLUE
fz["A7"]="合計"; fz["A7"].font=BLD; fz["B7"]="=SUM(B5:B6)"; fz["B7"].font=BLD
sec(fz,9,"2. 対象会社の現預金（2026年12月期末）",3)
fz["A10"]="現預金"; fz["B10"]=1743; fz["B10"].font=BLUE
fz["A11"]="役員貸付金の回収（クロージング時に岸田氏が現金返済）"; fz["B11"]=500; fz["B11"].font=BLUE
fz["A12"]="運転資金として留保する現金"; fz["B12"]="=-前提!$B$26"; fz["B12"].font=GRN
fz["A13"]="ブリッジローンの返済原資"; fz["A13"].font=BLD; fz["B13"]="=SUM(B10:B12)"; fz["B13"].font=BLD
sec(fz,15,"3. 調達（クロージング時）",3)
fz["A16"]="みずほ銀行 タームローン（A＋B）"; fz["B16"]="=前提!$B$23"; fz["B16"].font=GRN
fz["A17"]="みずほ銀行 ブリッジローン"; fz["B17"]="=B13"; fz["B17"].font=BLK
fz["A18"]="榊原氏 出資"; fz["B18"]=1000; fz["B18"].font=BLUE
fz["A19"]="中野氏 出資（資本金10万円）＋株主貸付"; fz["B19"]="=B7-B16-B17-B18"; fz["B19"].font=RED
fz["A20"]="合計"; fz["A20"].font=BLD; fz["B20"]="=SUM(B16:B19)"; fz["B20"].font=BLD
fz["A22"]="差引（ゼロであること）"; fz["B22"]="=B7-B20"; fz["B22"].font=BLD
fz["A24"]="合併後の有利子負債"; fz["A24"].font=BLD; fz["B24"]="=B16"; fz["B24"].font=BLD
fz["A25"]="合併後の現預金"; fz["A25"].font=BLD; fz["B25"]="=前提!$B$26"; fz["B25"].font=BLD
fz["A26"]="ネットデット"; fz["A26"].font=BLD; fz["B26"]="=B24-B25"; fz["B26"].font=BLD
fz["A27"]="ネットデット／EBITDA"; fz["B27"]="=B26/529"; fz["B27"].number_format='0.0"倍"'
for r in [5,6,7,10,11,12,13,16,17,18,19,20,22,24,25,26]: fz.cell(row=r,column=2).number_format=NUM
fz["A29"]="※ ブリッジローンはクロージング後、SPCと対象会社を合併したうえで対象会社の余剰現預金をもって全額返済する。"; fz["A29"].font=GRY
fz["A30"]="※ 留保する現金を1億円増やすと、中野氏の必要拠出額も1億円増え、2027年の最低現金残高も1億円増える（1対1の関係）。"; fz["A30"].font=GRY
fz["A31"]="※ EBITDAは2026年12月期計画の529百万円。"; fz["A31"].font=GRY
fz.column_dimensions["A"].width=48; fz.column_dimensions["B"].width=14

wb.save("カンテサンス_2027年月次CF計画.xlsx"); print("saved")
