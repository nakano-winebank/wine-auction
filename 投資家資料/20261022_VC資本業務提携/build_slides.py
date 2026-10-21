# -*- coding: utf-8 -*-
"""投資家様向け会社紹介（2026/10/21版）の差替スライドを、元資料のテンプレートで作る。

元資料（Meiryo・紺の帯・銅色アクセント・CONFIDENTIALバッジ・ロゴ入りマスター）と同じ
マスターを持つ 9/30 会社案内をベースに新規スライドを追加し、元スライドは削除する。

使い方:
  python3 build_slides.py <テンプレートpptx> <model.json> <出力pptx>
"""
import copy
import json
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as I, Pt

NAVY, NAVY2, COPPER, CARD, SAND = "30373F", "36394B", "C68A65", "F6F6F6", "EAD5C7"
BLACK, WHITE, GREY, LINE, BEIGE = "000000", "FFFFFF", "595959", "E3E3E3", "F3EDE6"
F = "Meiryo"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

TEMPLATE, MODEL, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
M = json.load(open(MODEL, encoding="utf-8"))
pr = Presentation(TEMPLATE)
N_ORIG = len(pr.slides)

# テンプレートから流用する部品（CONFIDENTIALバッジ・ページ番号）
BADGE = PAGENUM = None
for s in pr.slides:
    for sh in s.shapes:
        if BADGE is None and sh.has_text_frame and sh.text_frame.text.strip() == "CONFIDENTIAL":
            BADGE = copy.deepcopy(sh._element)
        if PAGENUM is None and sh.is_placeholder and sh.placeholder_format.type == 13:  # SLIDE_NUMBER
            PAGENUM = copy.deepcopy(sh._element)
# ロゴはマスター2にあるため、元資料の本文スライドと同じレイアウトを使う
BLANK = next(sl.slide_layout for sl in pr.slides
             if sl.slide_layout.name == "BLANK" and sl.slide_layout.slide_master is not pr.slide_masters[0])


# ───────── 描画ヘルパー ─────────
def _ea(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs", "sym"):
        el = rPr.find(A + tag)
        if el is None:
            el = etree.SubElement(rPr, A + tag)
        el.set("typeface", F)


def _fill_runs(tf, parts, size, bold, color, align, space):
    if isinstance(parts, str):
        parts = [[(parts, {})]]
    if parts and isinstance(parts[0], tuple):
        parts = [parts]
    split = []
    for line_ in parts:
        cur = []
        for t, o in line_:
            chunks = t.split("\n")
            for k, ch in enumerate(chunks):
                if k:
                    split.append(cur)
                    cur = []
                cur.append((ch, o))
        split.append(cur)
    parts = split
    first = True
    for line in parts:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if space:
            p.line_spacing = space
        for t, o in line:
            r = p.add_run()
            r.text = t
            r.font.name = F
            _ea(r)
            r.font.size = Pt(o.get("sz", size))
            r.font.bold = o.get("b", bold)
            r.font.color.rgb = RGBColor.from_string(o.get("c", color))
        # 空セルは段落末の既定サイズ（18pt）で行が高くなるため、段落末にも同じサイズを入れる
        p._p.get_or_add_endParaRPr().set("sz", str(int(size * 100)))


def tb(sl, x, y, w, h, parts, size=10, bold=False, color=NAVY, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, space=None):
    sp = sl.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _fill_runs(tf, parts, size, bold, color, align, space)
    return sp


def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    sp = sl.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    style = sp._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        sp._element.remove(style)
    return sp


def boxtext(sp, parts, size=9, bold=False, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            space=None, margin=0.04):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = I(margin)
    tf.margin_top = tf.margin_bottom = I(0.02)
    _fill_runs(tf, parts, size, bold, color, align, space)


def line(sl, x1, y1, x2, y2, color=GREY, w=1.0):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(w)
    return c


def new_slide(title, band, band_size=13):
    s = pr.slides.add_slide(BLANK)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    rect(s, 0, 5.482, 10, 0.143, fill=NAVY)
    if BADGE is not None:
        s.shapes._spTree.append(copy.deepcopy(BADGE))
    if PAGENUM is not None:
        s.shapes._spTree.append(copy.deepcopy(PAGENUM))
    tb(s, 0.262, 0.314, 9.4, 0.30, title, size=14, bold=True, color=BLACK)
    b = rect(s, 0, 0.737, 10, 0.359, fill=NAVY2)
    boxtext(b, band, size=band_size, bold=True, color=WHITE)
    return s


def table(sl, x, y, col_w, row_h, rows, size=8, header_rows=1, merges=(), styles=None):
    """rows: セル文字列の2次元配列。styles[(r,c)] = dict(fill,color,bold,align,size)。"""
    nr, nc = len(rows), len(col_w)
    gt = sl.shapes.add_table(nr, nc, I(x), I(y), I(sum(col_w)), I(sum(row_h)))
    t = gt.table
    tblPr = t._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style_id = tblPr.find(A + "tableStyleId")
    if style_id is not None:
        style_id.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
    for j, w in enumerate(col_w):
        t.columns[j].width = I(w)
    for i, h in enumerate(row_h):
        t.rows[i].height = I(h)
    for (r0, c0, r1, c1) in merges:
        t.cell(r0, c0).merge(t.cell(r1, c1))
    styles = styles or {}
    for i in range(nr):
        for j in range(nc):
            cell = t.cell(i, j)
            if cell.is_spanned:
                # 結合で隠れたセルも既定18ptの空段落が行高を押し上げるため、文字サイズを揃えておく
                cell.text_frame.paragraphs[0].text = ""
                _fill_runs(cell.text_frame, [[("", {})]], size, False, NAVY, PP_ALIGN.LEFT, None)
                continue
            st = dict(fill=WHITE, color=NAVY, bold=False, align=PP_ALIGN.RIGHT if j > 0 else PP_ALIGN.LEFT,
                      size=size)
            if i < header_rows:
                st.update(fill=NAVY2, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            st.update(styles.get((i, None), {}))
            st.update(styles.get((i, j), {}))
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(st["fill"])
            cell.margin_left = cell.margin_right = I(0.05)
            cell.margin_top = cell.margin_bottom = I(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            _fill_runs(tf, [[(rows[i][j], {})]], st["size"], st["bold"], st["color"], st["align"], None)
            _borders(cell)
    return gt


def _borders(cell, color=LINE, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.find(A + tag[2:])
        if ln is not None:
            tcPr.remove(ln)
        ln = etree.SubElement(tcPr, A + tag[2:], w=str(w), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, A + "solidFill")
        etree.SubElement(sf, A + "srgbClr", val=color)
        etree.SubElement(ln, A + "prstDash", val="solid")
    # lnL..lnB must precede the cell fill in tcPr
    fill = tcPr.find(A + "solidFill")
    if fill is not None:
        tcPr.remove(fill)
        tcPr.append(fill)


def oku(v, nd=1):
    s = f"{abs(v) / 100:,.{nd}f}"
    return ("▲" + s) if v < -0.05 else s


def mil(v):
    s = f"{abs(v):,.0f}"
    return ("▲" + s) if v < -0.5 else s


def pct(v):
    return f"{v * 100:.1f}%"


E = {e["key"]: e for e in M["entities"]}
SEG = [("ワイン・WineTech", ["wb_base", "wb_fund", "wb_mile"]),
       ("グランメゾン", ["apicius", "quint"]),
       ("ブラッスリー", ["tm"]),
       ("レストラン運営", ["vk", "others"])]
TOT = M["total"]
YL = ["2026年度\n（予算）", "2027年度", "2028年度", "2029年度"]


def seg_sum(keys, kind):
    return [sum(E[k][kind][i] for k in keys) for i in range(4)]


# ═══════════ p.34 ビジネスモデルに伴う組織再編 ═══════════
s = new_slide("ビジネスモデルに伴う組織再編",
              f"持分をWineBankへ集約し連結運営へ ― 連結売上 {oku(TOT['rev'][0], 0)}億円（2026年度予算）→ "
              f"{oku(TOT['rev'][3], 0)}億円（2029年度）", band_size=12)

# 左：再編後のグループ構成
tb(s, 0.30, 1.19, 4.40, 0.18,
   "株主：中野邦人／マネーフォワードVP／PrivateBANK／ベクトル／寺田倉庫 ほか",
   size=6.5, color=GREY)
wb = rect(s, 0.30, 1.40, 4.40, 0.56, fill=NAVY2)
boxtext(wb, [[("株式会社WineBank（親会社）", {"sz": 11, "b": True})],
             [("ワイン卸・EC ｜ WineTech（ファンド・ST・オークション）｜ ワインマイル ｜ WineBank CLUB",
               {"sz": 6.5, "c": SAND})]], space=1.1)
tb(s, 0.30, 2.01, 4.40, 0.18, "▼ 外部投資家以外の持分（中野個人保有分等）をWineBankへ譲渡・集約",
   size=7, bold=True, color=COPPER, align=PP_ALIGN.CENTER)
line(s, 2.50, 2.21, 2.50, 2.29, color=GREY)
CXS = [0.30, 1.42, 2.54, 3.66]
CW = 1.04
line(s, CXS[0] + CW / 2, 2.29, CXS[3] + CW / 2, 2.29, color=GREY)
cols = [
    ("グランメゾン", COPPER, [("アピシウス", "100%", "銀座・創業41年"),
                         ("カンテサンス", "50〜66.7%", "SPC経由（外部33.3〜50%）\n三つ星／取得協議中")]),
    ("ブラッスリー", COPPER, [("Brasserie\nThierry Marx", "66.7%", "外部投資家33.3%\n2027年4月 六本木開業"),
                         ("ティエリーマルクス×TBS", "ライセンス", "サブライセンスで\nWineBankが運営")]),
    ("レストラン運営", COPPER, [("バリューキッチン", "60%", "Prime42・oto ほか"),
                           ("バリューテーブル\n・エルネット等", "60〜75%", "飲食・不動産管理")]),
    ("持分法等（連結外）", "9E9E9E", [("あどばる", "37%", "ビジョングループ\n（東証プライム）"),
                               ("KINKA／バーガー\nレボリューション", "20%／49%", "持分法の対象")]),
]
for (head, hc, items), x in zip(cols, CXS):
    line(s, x + CW / 2, 2.29, x + CW / 2, 2.36, color=GREY)
    h = rect(s, x, 2.36, CW, 0.24, fill=hc)
    boxtext(h, head, size=7.5, bold=True)
    for k, (nm, pc_, note) in enumerate(items):
        yy = 2.66 + k * 0.70
        grey = hc != COPPER
        b = rect(s, x, yy, CW, 0.64, fill=CARD, line=LINE)
        boxtext(b, [[(nm, {"sz": 6.5, "b": True, "c": GREY if grey else NAVY2})],
                    [(pc_, {"sz": 9, "b": True, "c": GREY if grey else COPPER})],
                    [(note, {"sz": 5.5, "c": GREY})]], space=1.0, margin=0.03)
flow = rect(s, 0.30, 4.08, 4.40, 0.44, fill=BEIGE)
boxtext(flow, [[("WineBank → 各社　", {"b": True, "c": NAVY2}), ("ワイン卸・経営指導・AM（カンテサンスSPC）", {})],
               [("各社 → WineBank　", {"b": True, "c": NAVY2}),
                ("ワイン代金・経営指導料・AM報酬／ライセンスフィー（TM）", {})]],
        size=6.5, color=NAVY, align=PP_ALIGN.LEFT, space=1.15, margin=0.10)

# 右：連結業績計画（億円）
tb(s, 4.90, 1.19, 4.80, 0.18, "連結業績計画（億円）", size=8.5, bold=True, color=NAVY2)
rows = [["セグメント", "", *YL]]
styles = {}
merges = [(0, 0, 0, 1)]
for name, keys in SEG:
    r_ = seg_sum(keys, "rev")
    o_ = seg_sum(keys, "op")
    i0 = len(rows)
    rows.append([name, "売上高", *[oku(v) for v in r_]])
    rows.append(["", "営業利益", *[oku(v) for v in o_]])
    merges.append((i0, 0, i0 + 1, 0))
    styles[(i0, 0)] = dict(bold=True, color=NAVY2, align=PP_ALIGN.LEFT)
    styles[(i0 + 1, None)] = dict(fill=CARD)
    styles[(i0 + 1, 0)] = dict(fill=WHITE, bold=True, color=NAVY2, align=PP_ALIGN.LEFT)
    for j in range(2, 6):
        styles[(i0 + 1, j)] = dict(fill=CARD, color=COPPER, bold=True)
for lab, vals, fmt in (("連結売上高", TOT["rev"], oku), ("連結営業利益", TOT["op"], oku),
                       ("営業利益率", [o / r for o, r in zip(TOT["op"], TOT["rev"])], pct),
                       ("（参考）持分按分後 営業利益", TOT["attrib_op"], oku)):
    i0 = len(rows)
    rows.append([lab, "", *[fmt(v) for v in vals]])
    merges.append((i0, 0, i0, 1))
    big = lab in ("連結売上高", "連結営業利益")
    styles[(i0, None)] = dict(fill=NAVY2 if big else BEIGE, color=WHITE if big else NAVY, bold=big)
    styles[(i0, 0)] = dict(fill=NAVY2 if big else BEIGE, color=WHITE if big else NAVY, bold=True,
                           align=PP_ALIGN.LEFT, size=7.5 if not big else 8)
for j in range(1, 6):
    for i in range(1, 9):
        styles.setdefault((i, j), {})
for i in range(1, 9):
    styles[(i, 1)] = dict(styles.get((i, 1), {}), align=PP_ALIGN.CENTER, size=7, color=GREY)
table(s, 4.90, 1.40, [1.12, 0.64, 0.76, 0.76, 0.76, 0.76],
      [0.34] + [0.215] * 8 + [0.25, 0.25, 0.22, 0.22], rows, size=8, merges=merges, styles=styles)
tb(s, 0.30, 4.62, 9.40, 0.60,
   [[("年度＝WineBank決算期（10月〜翌9月）。2026年度＝2027年9月期の予算。", {"b": True})],
    [("決算期が異なる子会社は月割りで組替え。カンテサンスは期初取得を仮定したプロフォーマ（IM計画値。取得後の報酬見直し等の改善と"
      "役員退職金等の一時費用は含まず）。グループ内取引の消去は未反映（WineBankのグループ内コンサル料は消去済）。"
      "持分按分後営業利益はカンテサンス50%・TM66.7%・バリューキッチン60%等の下限持分で按分。事業別の内訳と出所は次頁。", {})]],
   size=6.3, color=GREY, space=1.1)
s.notes_slide.notes_text_frame.text = (
    "元のp.34（売上4⇒30億・スキーム図追加予定）の差替。外部投資家以外の持分をWineBankへ集約した後の連結PL。"
    "数値は 組織再編_連結PL_2026-2029.xlsx と一致。")

# ═══════════ p.35 連結損益計画（事業別） ═══════════
s = new_slide("組織再編後の連結損益計画（事業別）",
              "2026年度予算〜2029年度 ｜ 単位：百万円 ｜ 年度＝WineBank決算期（10月〜翌9月）", band_size=12)
SHORT = {
    "wb_base": "銀行提出FY2027プランD（グループ内コンサル料82.5を消去）。以降年5%成長・増分粗利30%",
    "wb_fund": "銀行提出プランEの上積み（私募＋オークション300・粗利率30%）を据置",
    "wb_mile": "ワインマイル3か年PLを最新レートカード（9/29）で再計算",
    "apicius": "2026年12月期 着地予想（グループ管理表）。以降売上年3%・利益率据置",
    "quint": "IM（2026/9）監修後・売上横ばいケース。2030年は年4%で延長",
    "tm": "事業収支統合モデル v4（2026/10/16・4月〜3月期）。2027年4月開業",
    "vk": "着地予想700/100を2027年6月期に充当（2026年6月期 実績645）。以降年3%",
    "others": "バリューテーブル600/20・エルネット100/5・摩耶100/5 の合算を据置",
}
NAMES = {"wb_base": "WineBank｜既存ワイン事業", "wb_fund": "WineBank｜私募・オークション",
         "wb_mile": "WineBank｜ワインマイル", "apicius": "アピシウス", "quint": "カンテサンス（プロフォーマ）",
         "tm": "Brasserie Thierry Marx", "vk": "バリューキッチン", "others": "バリューテーブル・エルネット等"}
OWN = {"wb_base": "親会社", "wb_fund": "親会社", "wb_mile": "親会社", "apicius": "100%", "quint": "50〜66.7%",
       "tm": "66.7%", "vk": "60%", "others": "60〜75%"}
hdr1 = ["事業", "持分"]
hdr2 = ["", ""]
for y in ["2026年度（予算）", "2027年度", "2028年度", "2029年度"]:
    hdr1 += [y, ""]
    hdr2 += ["売上", "営業利益"]
hdr1.append("前提・出所")
hdr2.append("")
rows = [hdr1, hdr2]
merges = [(0, 0, 1, 0), (0, 1, 1, 1), (0, 10, 1, 10)] + [(0, 2 + 2 * k, 0, 3 + 2 * k) for k in range(4)]
styles = {(1, None): dict(size=7)}
order = ["wb_base", "wb_fund", "wb_mile", "apicius", "quint", "tm", "vk", "others"]
for k in order:
    e = E[k]
    i0 = len(rows)
    row = [NAMES[k], OWN[k]]
    for y in range(4):
        row += [mil(e["rev"][y]), mil(e["op"][y])]
    row.append(SHORT[k])
    rows.append(row)
    styles[(i0, 0)] = dict(bold=True, color=NAVY2, size=7.3)
    styles[(i0, 1)] = dict(align=PP_ALIGN.CENTER, size=7, color=GREY)
    for y in range(4):
        styles[(i0, 3 + 2 * y)] = dict(color=COPPER, bold=True, fill=CARD)
    styles[(i0, 10)] = dict(align=PP_ALIGN.LEFT, size=5.8, color=GREY)
i0 = len(rows)
row = ["連結合計", ""]
for y in range(4):
    row += [mil(TOT["rev"][y]), mil(TOT["op"][y])]
row.append("グループ内取引の消去は未反映")
rows.append(row)
merges.append((i0, 0, i0, 1))
styles[(i0, None)] = dict(fill=NAVY2, color=WHITE, bold=True)
styles[(i0, 10)] = dict(fill=NAVY2, color=SAND, size=5.8, align=PP_ALIGN.LEFT, bold=False)
i1 = len(rows)
row = ["営業利益率", ""]
for y in range(4):
    row += [pct(TOT["op"][y] / TOT["rev"][y]), ""]
row.append("連結営業利益／連結売上高")
rows.append(row)
i2 = len(rows)
row = ["（参考）持分按分後 営業利益", ""]
for y in range(4):
    row += [mil(TOT["attrib_op"][y]), ""]
row.append("下限持分で按分（親会社帰属の目安）")
rows.append(row)
for ii in (i1, i2):
    merges.append((ii, 0, ii, 1))
    merges.extend((ii, 2 + 2 * y, ii, 3 + 2 * y) for y in range(4))
    styles[(ii, None)] = dict(fill=BEIGE, bold=True, align=PP_ALIGN.CENTER)
    styles[(ii, 0)] = dict(fill=BEIGE, bold=True, align=PP_ALIGN.LEFT)
    styles[(ii, 10)] = dict(fill=BEIGE, size=5.8, align=PP_ALIGN.LEFT, bold=False, color=GREY)
table(s, 0.30, 1.22, [1.78, 0.50] + [0.555] * 8 + [2.68],
      [0.20, 0.18] + [0.275] * 8 + [0.25, 0.21, 0.21], rows, size=7.5, header_rows=2, merges=merges,
      styles=styles)
tb(s, 0.30, 4.66, 9.40, 0.55,
   [[("連結対象外：", {"b": True}), ("あどばる（37%・ビジョングループ）、KINKA（20%）、バーガーレボリューション（49%）は持分法等。"
                                  "ワイン・ラ・ターブルは店舗をバリューキッチンへ移管済のため除外。", {})],
    [("仮置きの前提：", {"b": True}),
     ("WineBank既存事業の年5%成長、アピシウス・バリューキッチンの年3%成長は計画未策定のための仮置き。"
      "カンテサンスの取得後改善（役員報酬見直し等 年約0.8億円）は上振れ要因として含めていない。", {})]],
   size=6.3, color=GREY, space=1.1)

# ═══════════ ワインマイル経済圏 ═══════════
s = new_slide("ワインマイル経済圏 ―「ワイン貯金」",
              "持つ・貯める・開ける・贈る・譲る。ワインにまつわる体験を、ひとつの通貨でつなぐ")
CX = [0.587, 3.677, 6.855]
CWD = 2.559
cards = [("貯まる", "管理料の、その先へ", "管理料2.75%に対し、マイル還元4〜6%", False),
         ("使える", "飲む・贈る・集う", "系列店・グランメゾン・オークション等", False),
         ("育つ", "値上がりの70%", "会員のもの。当社は成功報酬30%のみ", True)]
for i, (lab, val, note, on) in enumerate(cards):
    x = CX[i]
    rect(s, x, 1.22, CWD, 0.95, fill=NAVY2 if on else CARD)
    tb(s, x + 0.20, 1.33, CWD - 0.40, 0.22, lab, size=9.5, bold=True, color=SAND if on else COPPER)
    tb(s, x + 0.20, 1.53, CWD - 0.40, 0.34, val, size=16, bold=True, color=WHITE if on else NAVY2)
    tb(s, x + 0.20, 1.92, CWD - 0.40, 0.20, note, size=7.5, color="C9B6A8" if on else GREY)

tb(s, 0.587, 2.30, 4.4, 0.22, "新・会員制度（2026年9月改定）", size=10, bold=True, color=BLACK)
rows = [["ランク", "ご購入額", "管理手数料", "マイル還元", "実質"],
        ["PRESTIGE", "100万円〜", "2.75%", "4%", "+1.25%"],
        ["GOLD", "500万円〜", "2.75%", "5%", "+2.25%"],
        ["SIGNATURE", "1,000万円〜", "2.75%", "6%", "+3.25%"]]
styles = {(i, 0): dict(bold=True, color=NAVY2) for i in range(1, 4)}
for i in range(1, 4):
    styles[(i, 4)] = dict(bold=True, color=COPPER)
    for j in range(1, 4):
        styles[(i, j)] = dict(align=PP_ALIGN.CENTER)
    styles[(i, 4)].update(align=PP_ALIGN.CENTER)
table(s, 0.587, 2.55, [1.00, 0.95, 0.85, 0.85, 0.75], [0.24, 0.22, 0.22, 0.22], rows, size=8, styles=styles)
tb(s, 0.587, 3.52, 4.40, 0.85,
   [[("・", {}), ("所有権は会員のまま。", {"b": True}), ("当社は保管・保険・真正性管理を担い、賃貸借は使わない", {})],
    [("・", {}), ("マイルは管理料に対する無償付与。", {"b": True}), ("前払式支払手段に当たらないことを法律事務所に確認（2026年9月）", {})],
    [("・", {}), ("交流イベントとランク制で継続率を高める。", {"b": True}), ("使うほど貯まり、ランクが上がる", {})]],
   size=7.3, color=NAVY, space=1.25)

tb(s, 5.25, 2.30, 4.2, 0.22, "3か年計画（ワインマイル事業）", size=10, bold=True, color=BLACK)
mile = E["wb_mile"]
mrows = [["", "2027年度", "2028年度", "2029年度"],
         ["期末会員数", "37名", "102名", "192名"],
         ["期末預かり資産", "0.75億円", "2.07億円", "4.02億円"],
         ["売上高", *[f"{v / 100:.2f}億円" for v in [77.4, 138.7, 208.1]]],
         ["営業利益", *[f"{v / 100:.2f}億円" for v in [19.0, 34.6, 52.0]]]]
styles = {(i, 0): dict(bold=True, color=NAVY2) for i in range(1, 5)}
for j in range(1, 4):
    styles[(4, j)] = dict(bold=True, color=COPPER)
table(s, 5.25, 2.55, [1.30, 0.94, 0.94, 0.94], [0.24, 0.22, 0.22, 0.22, 0.22], mrows, size=8, styles=styles)
tb(s, 5.25, 3.78, 4.15, 0.50,
   "最新レートカードで3か年PLを再計算（ワイン販売粗利30%・保管1.05%・マイル費用は期待原価で発行時計上）。"
   "人件費・システム費は既存体制で吸収する前提。",
   size=6.3, color=GREY, space=1.1)

rect(s, 0.587, 4.46, 8.826, 0.66, fill=CARD)
rect(s, 0.587, 4.46, 0.05, 0.66, fill=COPPER)
tb(s, 0.76, 4.52, 8.55, 0.56,
   [[("同じキャッシュで、支配下のワインは2.5倍。", {"sz": 11, "b": True, "c": BLACK})],
    [("1億円の在庫を自社で持ち続ける代わりに、会員へ販売して預かり、1.2億円を新規に買い付ける。"
      "純CFはほぼ同じ（年▲405万円→▲536万円）のまま、在庫と預かり資産の合計は1億円から2.5億円へ。", {"sz": 7.8, "c": NAVY})]],
   space=1.2)
s.notes_slide.notes_text_frame.text = (
    "出所：WineBank_ワインマイル経済圏構想_20260929（レートカード・法務整理）、ワインマイル3か年PL（9/2版）を"
    "最新レートカードで再計算した値（参考_ワインマイル3か年PL_最新レートカード再計算.xlsx）。")

# ═══════════ 国内オーベルジュ構想 ═══════════
s = new_slide("国内オーベルジュ構想", "ワイナリーに隣接し、温泉を持つ。日本にまだ無いオーベルジュをつくります")
cards = [("料理", "一日十数名だけ", "その土地の生産者と、同じテーブルに", False),
         ("ワイン", "ワイナリー隣接", "造り手の畑が、窓の外にある", False),
         ("滞在", "温泉 ＋ 10室", "食後に、移動しなくていい", True)]
for i, (lab, val, note, on) in enumerate(cards):
    x = CX[i]
    rect(s, x, 1.22, CWD, 0.95, fill=NAVY2 if on else CARD)
    tb(s, x + 0.20, 1.33, CWD - 0.40, 0.22, lab, size=9.5, bold=True, color=SAND if on else COPPER)
    tb(s, x + 0.20, 1.53, CWD - 0.40, 0.34, val, size=18, bold=True, color=WHITE if on else NAVY2)
    tb(s, x + 0.20, 1.92, CWD - 0.40, 0.20, note, size=8, color="C9B6A8" if on else GREY)
tb(s, 0.587, 2.27, 6.0, 0.22, "構想の骨子", size=11, bold=True, color=BLACK)
arows = [("", "一般的なオーベルジュ", "当社が考えるオーベルジュ"),
         ("立地", "景勝地・観光地に単独で建つ", "ワイナリーに隣接。造り手が徒歩圏にいる"),
         ("ワイン", "地元ワインが中心", "グループ在庫（管理含む32億円）から、世界のファインワインを産地価格で"),
         ("滞在", "宿泊機能のみ", "温泉を併設。地下に大型カーヴとワインバーも備える"),
         ("運営", "オーナーシェフの個人事業", "アピシウス・カンテサンス※を擁するグランメゾン運営グループが担う"),
         ("集客", "開業後に一から集める", "ワインマイルの交換先・会員イベントの舞台として、開業前から送客できる")]
TX, TW, TY, RH = 0.587, [1.10, 2.90, 4.826], 2.50, 0.31
for i, (a_, b_, c_) in enumerate(arows):
    y = TY + i * RH
    hdr = i == 0
    rect(s, TX, y, sum(TW), RH, fill=NAVY2 if hdr else (CARD if i % 2 else WHITE), line=None if hdr else LINE)
    for j, (txt, w) in enumerate(zip((a_, b_, c_), TW)):
        cx = TX + sum(TW[:j])
        col = WHITE if hdr else (NAVY2 if j == 0 else NAVY)
        hi = (not hdr and j == 2)
        tb(s, cx + 0.12, y + 0.03, w - 0.20, RH - 0.06, txt, size=8.3, bold=hdr or j == 0 or hi,
           color=COPPER if hi else col, anchor=MSO_ANCHOR.MIDDLE, space=1.05)
tb(s, 0.587, 4.43, 8.826, 0.20,
   [[("ブラッスリー出店計画（2027年4月 六本木 → 2028年 札幌・大阪 → 2030年〜 リゾート含め年2店舗）の先に着手。　",
      {"sz": 7.5, "c": GREY}),
     ("候補地：北海道", {"sz": 7.5, "b": True, "c": NAVY2}),
     ("　※カンテサンスは取得に向けた基本合意の協議中", {"sz": 6.5, "c": GREY})]])
rect(s, 0.587, 4.66, 8.826, 0.48, fill=CARD)
rect(s, 0.587, 4.66, 0.05, 0.48, fill=COPPER)
tb(s, 0.76, 4.71, 8.50, 0.40,
   [[("料理人が、一日の終わりまでお客様と過ごせる場所を。", {"sz": 11, "b": True, "c": BLACK})],
    [("ゼロから新しい店をつくるより、長い時間をかけて積み上げられたものを、更に発展させ次の世代へ渡す。"
      "私たちが得意なのは、そちらです。", {"sz": 8.3, "c": NAVY})]], space=1.2)

# ═══════════ 中長期ロードマップ ═══════════
s = new_slide("中長期ロードマップ",
              f"グループ連結で売上{oku(TOT['rev'][3], 0)}億円・営業利益{oku(TOT['op'][3], 0)}億円（2029年度）へ。IPOを見据えて体制を整備",
              band_size=12)
XS = [0.45, 2.75, 5.05, 7.35]
WW = 2.20
YT = ["2026年度（予算）", "2027年度", "2028年度", "2029年度"]
YP = ["2026/10〜2027/9", "2027/10〜2028/9", "2028/10〜2029/9", "2029/10〜2030/9"]
MS = [["持分集約による組織再編", "カンテサンス取得（予定）", "Brasserie Thierry Marx\n六本木開業（2027年4月）",
       "WineBank Auction（WEB）"],
      ["ワインマイル本格展開", "私募ファンド・オークション\n年3億円規模", "Thierry Marx 札幌・大阪（2028年）"],
      ["Thierry Marx 首都圏・\n政令指定都市（2029年）", "ワインマイル会員100名超", "オークション C2C・海外展開"],
      ["Thierry Marx リゾート含め\n年2店舗（2030年〜）", "国内オーベルジュ着手", "IPOへ"]]
rect(s, 0.45, 1.22, 9.10, 0.30, fill=BEIGE)
for k, x in enumerate(XS):
    phase = ["成長戦略 PhaseⅠ", "成長戦略 PhaseⅠ", "成長戦略 PhaseⅡ", "成長戦略 PhaseⅡ"][k]
    tb(s, x, 1.27, WW, 0.20, phase, size=7.5, bold=True, color=COPPER, align=PP_ALIGN.CENTER)
for k, x in enumerate(XS):
    h = rect(s, x, 1.60, WW, 0.42, fill=NAVY2)
    boxtext(h, [[(YT[k], {"sz": 10, "b": True})], [(YP[k], {"sz": 6.5, "c": SAND})]], space=1.0)
    rect(s, x, 2.06, WW, 1.08, fill=CARD)
    tb(s, x + 0.12, 2.12, WW - 0.24, 0.18, "連結売上高", size=7, color=GREY)
    tb(s, x + 0.12, 2.28, WW - 0.24, 0.36,
       [[(oku(TOT["rev"][k]), {"sz": 20, "b": True, "c": COPPER}), (" 億円", {"sz": 9, "b": True, "c": NAVY2})]])
    tb(s, x + 0.12, 2.70, WW - 0.24, 0.18, "連結営業利益", size=7, color=GREY)
    tb(s, x + 0.12, 2.84, WW - 0.24, 0.28,
       [[(oku(TOT["op"][k]), {"sz": 14, "b": True, "c": NAVY2}), (" 億円", {"sz": 8, "b": True, "c": NAVY2}),
         (f"（{pct(TOT['op'][k] / TOT['rev'][k])}）", {"sz": 7, "c": GREY})]])
    for m, txt in enumerate(MS[k]):
        yy = 3.22 + m * 0.36
        rect(s, x, yy, WW, 0.32, fill=WHITE, line=LINE)
        rect(s, x, yy, 0.05, 0.32, fill=COPPER)
        tb(s, x + 0.12, yy + 0.02, WW - 0.18, 0.28, txt, size=7, bold=True, color=NAVY2,
           anchor=MSO_ANCHOR.MIDDLE, space=1.0)
for k in range(3):
    tb(s, XS[k] + WW, 2.40, XS[k + 1] - XS[k] - WW, 0.30, "▶", size=10, color=COPPER, align=PP_ALIGN.CENTER)
tb(s, 0.45, 4.72, 9.10, 0.50,
   "連結数値は組織再編後のプロフォーマ（組織再編の頁・事業別内訳の頁を参照）。年度＝WineBank決算期（10月〜翌9月）。"
   "カンテサンスは期初取得を仮定。",
   size=6.3, color=GREY, space=1.1)

s.notes_slide.notes_text_frame.text = (
    "元のp.44（WineBank単体 売上11→17→25億円・レストラン売上30→40→50億円）を、組織再編後のグループ連結数値に置き換えた。")

# ───────── 元スライドを削除して保存 ─────────
lst = pr.slides._sldIdLst
for el in list(lst)[:N_ORIG]:
    pr.part.drop_rel(el.get(R + "id"))
    lst.remove(el)
pr.save(OUT)
print("saved", OUT, len(pr.slides))
