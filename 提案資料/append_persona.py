"""手修正済みデッキの末尾に、build_persona.js で作った1枚を追加する。

既存スライドには一切触れず、ペルソナのスライドの図形・背景・ノートを
デッキ側の白紙レイアウトに移し替えるだけにする。

usage: python append_persona.py <手修正デッキ.pptx> <persona.pptx> <出力.pptx>
"""
import copy
import sys

from pptx import Presentation


def blank_layout(prs):
    # プレースホルダが最も少ないレイアウトを使う（余計なタイトル枠を出さない）
    return min(prs.slide_layouts, key=lambda l: len(l.placeholders))


def main(deck_path, persona_path, out_path):
    deck = Presentation(deck_path)
    src = Presentation(persona_path).slides[0]

    dst = deck.slides.add_slide(blank_layout(deck))
    for ph in list(dst.placeholders):
        ph._element.getparent().remove(ph._element)

    # 背景
    src_bg = src._element.cSld.bg
    if src_bg is not None:
        dst_csld = dst._element.cSld
        if dst_csld.bg is not None:
            dst_csld.remove(dst_csld.bg)
        dst_csld.insert(0, copy.deepcopy(src_bg))

    # 図形（画像は使っていないのでリレーションのコピーは不要）
    for el in src.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        if tag == "pic":
            raise SystemExit("画像を含むスライドには未対応")
        dst.shapes._spTree.append(copy.deepcopy(el))

    if src.has_notes_slide:
        # デッキ側のノートマスターにプレースホルダがなく、python-pptx が本文枠を
        # 作れないため、既存スライドと同じ形式の元ノートの図形ツリーを丸ごと移す
        dst_csld = dst.notes_slide._element.cSld
        dst_csld.replace(dst_csld.spTree, copy.deepcopy(src.notes_slide._element.cSld.spTree))

    deck.save(out_path)
    print(f"{len(deck.slides)} slides -> {out_path}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
