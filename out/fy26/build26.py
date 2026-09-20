# -*- coding: utf-8 -*-
"""FY2026(2026年9月期) 月次推移・着地予想
実績: 損益計算書(月次推移) 2025/10-2026/07 / 想定: 2026/08-09"""
import sys; sys.path.insert(0,'.')
from data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F='メイリオ'
BK=Font(name=F); BL=Font(name=F,color='8A6A24',bold=True); GR=Font(name=F,color='595959')
SB=Font(name=F,bold=True); HD=Font(name=F,bold=True,color='FFFFFF')
SM=Font(name=F,size=9,color='8C8C8C'); RD=Font(name=F,size=9,color='333333')
AC=Font(name=F,color='1F4E79'); ES=Font(name=F,color='8A6A24',italic=True)
HF=PatternFill('solid',fgColor='1A1A1A'); SF=PatternFill('solid',fgColor='F2F1EE')
YL=PatternFill('solid',fgColor='FAF3E2'); TF=PatternFill('solid',fgColor='F7F6F3')
OR_=PatternFill('solid',fgColor='F4EDDC'); GRP=PatternFill('solid',fgColor='EFEEEA')
GY=PatternFill('solid',fgColor='EDEDED')
TB=Border(top=Side(style='thin',color='1A1A1A'))
YEN='¥#,##0;(¥#,##0);-'; NUM='#,##0;(#,##0);-'; PCT='0.0%'
wb=Workbook(); wb.remove(wb.active)
def bar(ws,r,t,span=6):
    ws.cell(r,2,t).font=HD
    for c in range(2,2+span): ws.cell(r,c).fill=HF
    ws.cell(r,2).font=HD
FM=MO+['2026/08','2026/09']

# ================= ③前提 =================
p=wb.create_sheet('③前提')
for col,w in zip('ABCDEFG',[3,38,18,18,62,2,2]): p.column_dimensions[col].width=w
p['B1']='FY2026 着地予想　前提条件'; p['B1'].font=Font(name=F,bold=True,size=15)
p['B2']='黄色＝入力セル。2025/10-2026/07は実績（添付の損益計算書 月次推移）、2026/08-09は想定。単位：円'; p['B2'].font=SM
bar(p,4,'【1】2026年8月・9月の既存事業　※営業利益トントン（赤字を増やさない）前提',4)
a=[('2026/08 既存事業 売上',57500000,'貴社想定：8-9月の2か月で110〜120百万円。ここでは中央値115百万円を折半'),
   ('2026/09 既存事業 売上',57500000,'同上'),
   ('2026/08 販管費',25624504,'2026/04-07実績の月平均25,624,504円。貴社ご指定によりフラット'),
   ('2026/09 販管費',25624504,'同上'),
   ('2026/08 営業外費用',1323236,'2026/04-07実績の月平均'),
   ('2026/09 営業外費用',1323236,'同上')]
for i,(t,v,nt) in enumerate(a):
    r=5+i; p.cell(r,2,t).font=BK
    c=p.cell(r,3,v); c.number_format=YEN; c.font=BL; c.fill=YL
    p.cell(r,5,nt).font=SM
S8,S9,G8,G9,E8,E9=5,6,7,8,9,10
p.cell(11,2,'※営業利益トントンの前提により、各月の売上総利益＝販管費として売上原価を逆算しています。').font=RD
p.cell(12,2,'※この前提では売上レンジ（110〜120百万円）を動かしても着地の営業利益は変わりません。変わるのは売上高と粗利率の見え方だけです。').font=RD
n=14
bar(p,n,'【2】7月末までの営業利益の修正（戻し）',4)
p.cell(n+1,2,'営業利益の戻し').font=BK
c=p.cell(n+1,3,8000000); c.number_format=YEN; c.font=BL; c.fill=YL
p.cell(n+1,5,'貴社ご指摘。10か月実績▲111,980,586円に加算し、7月末時点を▲103,980,586円（≒営業赤字10,400万円）とする').font=RD
p.cell(n+2,2,'計上月（8＝2026/08・9＝2026/09）').font=BK
c=p.cell(n+2,3,8); c.number_format='0'; c.font=BL; c.fill=YL
p.cell(n+2,5,'★勘定科目（販管費の戻入か売上の追加か）はご確認ください。ここでは販管費の減少として処理しています。').font=RD
ADJ,ADJM=n+1,n+2
n=18
bar(p,n,'【3】2026年9月末までに計上する上乗せ案件',4)
for j,t in enumerate(['案件','売上','粗利']):
    c=p.cell(n+1,2+j,t); c.font=HD; c.fill=HF; c.alignment=Alignment(horizontal='center')
p.cell(n+1,5,'内容').font=HD; p.cell(n+1,5).fill=HF
up=[('① 私募（現物出資分）※売上計上なし',0,30000000,'★3億円は現物出資のため売上を計上せず、粗利30百万円（P2から10＋WineBankから20）のみ計上。売上原価のマイナスとして処理'),
    ('② 私募（売上計上分）',300000000,45000000,'3億円を売上計上。粗利45百万円（粗利率15.0%）'),
    ('③ アピシウスM&A仲介',20000000,20000000,'FY2027予定分の契約を9月末までに間に合わせる。原価なし'),
    ('④ コンサルティング料',10000000,10000000,'9月末までに計上。原価なし')]
U0=n+2
for i,(t,sl,gp,nt) in enumerate(up):
    r=U0+i; p.cell(r,2,t).font=BK
    for col,v in ((3,sl),(4,gp)):
        c=p.cell(r,col,v); c.number_format=YEN; c.font=BL; c.fill=YL
    p.cell(r,5,nt).font=SM
U1=U0+3; UT=U1+1
p.cell(UT,2,'合計').font=SB; p.cell(UT,2).border=TB
for col in (3,4):
    L=get_column_letter(col)
    c=p.cell(UT,col,f'=SUM({L}{U0}:{L}{U1})'); c.number_format=YEN; c.font=SB; c.fill=OR_; c.border=TB
p.cell(UT,5,'粗利105,000,000円は従来想定と同額。売上は630→330百万円に減少（現物出資分3億を売上計上しないため）').font=RD
n=UT+2
bar(p,n,'【4】商品在庫の見通し　※FY2027の期首在庫に直結します',4)
iv=[('2026/07末 商品（実績）',721162328,False,'添付の貸借対照表 月次推移より'),
    ('2026/08-09 商品仕入（想定）',55000000,True,'貴社ご指定：50〜60百万円。中央値55百万円'),
    ('2026/08-09 売上原価による出庫',None,False,'既存事業の8-9月＋私募（売上計上分）の原価。現物出資の評価差額は含まない'),
    ('現物出資 出資額（評価額）',300000000,True,'★現物出資した金額。売上には計上しません'),
    ('　うち 評価差額（粗利計上分）',None,False,'【3】①の粗利30,000,000円'),
    ('現物出資による在庫の振替（簿価）',None,False,'出資額−評価差額。売上原価ではなく投資その他の資産へ振替'),
    ('2026/09末 商品（見込）',None,False,'期首＋仕入−出庫−現物出資の振替')]
for i,(t,v,inp,nt) in enumerate(iv):
    r=n+1+i; p.cell(r,2,t).font=BK
    c=p.cell(r,3,v); c.number_format=YEN
    if inp: c.font=BL; c.fill=YL
    else: c.font=SB; c.fill=TF
    p.cell(r,5,nt).font=RD if inp else SM
IV0=n+1; IV_BUY=n+2; IV_OUT=n+3; IV_GK=n+4; IV_GKG=n+5; IV_GKB=n+6; IV_END=n+7
p.cell(IV_END+2,2,'※現物出資したワインは売上原価にはならず、簿価のまま投資その他の資産へ振り替わります。在庫からは出ますが損益には出ません。').font=RD
p.cell(IV_END+3,2,'※従来のFY2027モデルは期首在庫420百万円・2026/09に在庫200百万円を販売してみずほ銀行へ返済、という前提でした。').font=RD
p.cell(IV_END+4,2,'　今回の私募（現物出資300＋売上計上300）はそれより大きい取引のため、期首在庫の前提を置き直す必要があります。').font=RD
P=lambda r: f"'③前提'!$C${r}"
PD=lambda r: f"'③前提'!$D${r}"

# ================= ②月次推移表 =================
m=wb.create_sheet('②月次推移表')
m.column_dimensions['A'].width=3; m.column_dimensions['B'].width=30
for i in range(12): m.column_dimensions[get_column_letter(3+i)].width=14
m.column_dimensions['O'].width=16; m.column_dimensions['P'].width=3
m['B1']='FY2026（2025年10月-2026年9月）月次推移表　実績＋8月・9月想定'; m['B1'].font=Font(name=F,bold=True,size=14)
m['B2']='青字＝実績（2025/10-2026/07）　金色斜体＝想定（2026/08-09）　出典：損益計算書 月次推移'; m['B2'].font=SM
R=4
m.cell(R,2,'科目').font=SB; m.cell(R,2).fill=SF
for i,mm in enumerate(FM):
    c=m.cell(R,3+i,mm); c.font=SB; c.fill=SF if i<10 else OR_
    c.alignment=Alignment(horizontal='center')
c=m.cell(R,15,'通期'); c.font=SB; c.fill=SF; c.alignment=Alignment(horizontal='center')
def row(r,label,vals=None,f8=None,f9=None,bold=False,fill=None,top=False,pct=False,ind=0):
    c=m.cell(r,2,('　'*ind)+label); c.font=SB if bold else BK
    if fill: c.fill=fill
    if top: c.border=TB
    for i in range(12):
        col=3+i
        if i<10:
            v=vals[i] if vals else 0
            cc=m.cell(r,col,v); cc.font=SB if bold else AC
        else:
            cc=m.cell(r,col,(f8 if i==10 else f9)); cc.font=SB if bold else ES
            cc.fill=OR_ if not fill else fill
        cc.number_format=PCT if pct else YEN
        if fill: cc.fill=fill
        if top: cc.border=TB
    L=get_column_letter(3); Lx=get_column_letter(14)
    t=m.cell(r,15,f'=SUM({L}{r}:{Lx}{r})'); t.number_format=PCT if pct else YEN
    t.font=SB; t.fill=fill or TF
    if top: t.border=TB
r=5
m.cell(r,2,'【売上高】').font=SB; m.cell(r,2).fill=GRP
for c in range(3,16): m.cell(r,c).fill=GRP
R_SE=r+1; row(R_SE,'既存事業 売上',SALES,f'={P(S8)}',f'={P(S9)}')
R_F3=r+2; row(R_F3,'＋私募（現物出資分）※売上計上なし',None,0,f'={P(U0)}',ind=1)
R_F4=r+3; row(R_F4,'＋私募（売上計上分）',None,0,f'={P(U0+1)}',ind=1)
R_MA=r+4; row(R_MA,'＋M&A仲介',None,0,f'={P(U0+2)}',ind=1)
R_CS=r+5; row(R_CS,'＋コンサルティング料',None,0,f'={P(U0+3)}',ind=1)
R_ST=r+6
m.cell(R_ST,2,'売上高 合計').font=SB
for i in range(12):
    L=get_column_letter(3+i)
    c=m.cell(R_ST,3+i,f'=SUM({L}{R_SE}:{L}{R_CS})'); c.number_format=YEN; c.font=SB; c.fill=TF; c.border=TB
m.cell(R_ST,2).fill=TF; m.cell(R_ST,2).border=TB
c=m.cell(R_ST,15,f'=SUM(C{R_ST}:N{R_ST})'); c.number_format=YEN; c.font=SB; c.fill=TF; c.border=TB
r=R_ST+1
m.cell(r,2,'【売上原価】').font=SB; m.cell(r,2).fill=GRP
for c in range(3,16): m.cell(r,c).fill=GRP
R_CE=r+1
row(R_CE,'既存事業 売上原価',COGS,f'={P(S8)}-{P(G8)}',f'={P(S9)}-{P(G9)}')
R_CF=r+2
row(R_CF,'＋上乗せ案件 原価（売上計上分）',None,0,
    f'=SUM({P(U0+1)},{P(U0+2)},{P(U0+3)})-SUM({PD(U0+1)},{PD(U0+2)},{PD(U0+3)})',ind=1)
R_CG=r+3
row(R_CG,'＋現物出資 評価差額（原価のマイナス）',None,0,f'=-{PD(U0)}',ind=1)
R_CT=r+4
m.cell(R_CT,2,'売上原価 合計').font=SB
for i in range(12):
    L=get_column_letter(3+i)
    c=m.cell(R_CT,3+i,f'={L}{R_CE}+{L}{R_CF}+{L}{R_CG}'); c.number_format=YEN; c.font=SB; c.fill=TF; c.border=TB
m.cell(R_CT,2).fill=TF; m.cell(R_CT,2).border=TB
c=m.cell(R_CT,15,f'=SUM(C{R_CT}:N{R_CT})'); c.number_format=YEN; c.font=SB; c.fill=TF; c.border=TB
R_GP=R_CT+1
m.cell(R_GP,2,'売上総利益').font=SB; m.cell(R_GP,2).fill=OR_
for i in range(12):
    L=get_column_letter(3+i)
    c=m.cell(R_GP,3+i,f'={L}{R_ST}-{L}{R_CT}'); c.number_format=YEN; c.font=SB; c.fill=OR_
c=m.cell(R_GP,15,f'=SUM(C{R_GP}:N{R_GP})'); c.number_format=YEN; c.font=SB; c.fill=OR_
R_GM=R_GP+1
m.cell(R_GM,2,'　粗利率').font=BK
for i in range(12):
    L=get_column_letter(3+i)
    c=m.cell(R_GM,3+i,f'=IF({L}{R_ST}=0,0,{L}{R_GP}/{L}{R_ST})'); c.number_format=PCT; c.font=BK
c=m.cell(R_GM,15,f'=O{R_GP}/O{R_ST}'); c.number_format=PCT; c.font=SB; c.fill=TF
r=R_GM+1
m.cell(r,2,'【販売費及び一般管理費】').font=SB; m.cell(r,2).fill=GRP
for c in range(3,16): m.cell(r,c).fill=GRP
R_SG=r+1; row(R_SG,'販管費',SGA,f'={P(G8)}',f'={P(G9)}')
R_AJ=r+2
row(R_AJ,'7月末までの修正（戻し）',None,f'=-IF({P(ADJM)}=8,{P(ADJ)},0)',f'=-IF({P(ADJM)}=9,{P(ADJ)},0)',ind=1)
R_SGT=r+3
m.cell(R_SGT,2,'販管費 計').font=SB
for i in range(12):
    L=get_column_letter(3+i)
    c=m.cell(R_SGT,3+i,f'={L}{R_SG}+{L}{R_AJ}'); c.number_format=YEN; c.font=SB; c.fill=TF; c.border=TB
m.cell(R_SGT,2).fill=TF; m.cell(R_SGT,2).border=TB
c=m.cell(R_SGT,15,f'=SUM(C{R_SGT}:N{R_SGT})'); c.number_format=YEN; c.font=SB; c.fill=TF; c.border=TB
R_OP=R_SGT+1
m.cell(R_OP,2,'営業利益').font=SB; m.cell(R_OP,2).fill=OR_; m.cell(R_OP,2).border=TB
for i in range(12):
    L=get_column_letter(3+i)
    c=m.cell(R_OP,3+i,f'={L}{R_GP}-{L}{R_SGT}'); c.number_format=YEN; c.font=SB; c.fill=OR_; c.border=TB
c=m.cell(R_OP,15,f'=SUM(C{R_OP}:N{R_OP})'); c.number_format=YEN; c.font=SB; c.fill=OR_; c.border=TB
R_NI=R_OP+1; row(R_NI,'営業外収益',NOI,0,0)
R_NE=R_OP+2; row(R_NE,'営業外費用',NOE,f'={P(E8)}',f'={P(E9)}')
R_OR=R_OP+3
m.cell(R_OR,2,'経常利益').font=SB; m.cell(R_OR,2).fill=OR_; m.cell(R_OR,2).border=TB
for i in range(12):
    L=get_column_letter(3+i)
    c=m.cell(R_OR,3+i,f'={L}{R_OP}+{L}{R_NI}-{L}{R_NE}'); c.number_format=YEN; c.font=SB; c.fill=OR_; c.border=TB
c=m.cell(R_OR,15,f'=SUM(C{R_OR}:N{R_OR})'); c.number_format=YEN; c.font=SB; c.fill=OR_; c.border=TB
R_CO=R_OR+2
m.cell(R_CO,2,'営業利益 累計').font=SB
m.cell(R_CO+1,2,'経常利益 累計').font=SB
for k,src in ((0,R_OP),(1,R_OR)):
    for i in range(12):
        L=get_column_letter(3+i); Lp=get_column_letter(2+i)
        f_=f'={L}{src}' if i==0 else f'={Lp}{R_CO+k}+{L}{src}'
        c=m.cell(R_CO+k,3+i,f_); c.number_format=YEN; c.font=SB; c.fill=GY
m.cell(R_CO+3,2,'※2026/08-09の売上原価は「営業利益トントン」の前提から逆算（売上総利益＝販管費）。').font=RD
m.cell(R_CO+4,2,'※上乗せ4件はすべて2026/09に計上。売上330,000,000円・原価225,000,000円・粗利105,000,000円。').font=RD
m.cell(R_CO+5,2,'※「7月末までの修正（戻し）」は販管費の減少として2026年8月に計上（貴社ご指定）。').font=RD
m.cell(R_CO+6,2,'※私募のうち3億円は現物出資のため売上を計上せず、粗利30,000,000円を売上原価のマイナスとして計上しています。').font=RD
m.freeze_panes='C5'
MM=lambda r,c='O': f"'②月次推移表'!${c}${r}"
p.cell(IV_OUT,3,f"=SUM('②月次推移表'!M{R_CE}:N{R_CE})+SUM('②月次推移表'!M{R_CF}:N{R_CF})").number_format=YEN
p.cell(IV_GKG,3,f'=D{U0}').number_format=YEN
p.cell(IV_GKB,3,f'=C{IV_GK}-C{IV_GKG}').number_format=YEN
p.cell(IV_END,3,f'=C{IV0}+C{IV_BUY}-C{IV_OUT}-C{IV_GKB}').number_format=YEN
p.cell(IV_END,3).fill=OR_; p.cell(IV_END,3).font=SB

# ================= ①着地予想 =================
s=wb.create_sheet('①着地予想',0)
for col,w in zip('ABCDEFG',[3,42,20,20,58,2,2]): s.column_dimensions[col].width=w
s['B1']='FY2026（2026年9月期）着地予想'; s['B1'].font=Font(name=F,bold=True,size=15)
s['B2']='2025/10-2026/07は実績、2026/08-09は想定。単位：円（税別）'; s['B2'].font=SM
bar(s,4,'【1】営業利益のブリッジ',4)
for j,t in enumerate(['項目','金額','累計','内容']):
    c=s.cell(5,2+j,t); c.font=HD; c.fill=HF; c.alignment=Alignment(horizontal='center')
B0=6
br=[('10か月実績 営業利益（2025/10-2026/07）',f'=SUM(\'②月次推移表\'!C{R_OP}:L{R_OP})+{P(ADJ)}',
     '報告値▲111,980,586円に、7月末までの戻し8,000,000円を加算'),
    ('8月・9月 既存事業',f"=SUM('②月次推移表'!M{R_OP}:N{R_OP})-{P(ADJ)}-{PD(UT)}",
     '売上115百万円・営業利益トントン（赤字を増やさない）前提のため、貢献はゼロ'),
    ('上乗せ4件の粗利',f'={PD(UT)}','私募 現物出資30＋売上計上45＋M&A仲介20＋コンサル料10＝105百万円'),
    ('FY2026 着地 営業利益',None,'')]
for i,(lab,f_,nt) in enumerate(br):
    r=B0+i; last=(i==len(br)-1)
    s.cell(r,2,lab).font=SB if (i==0 or last) else BK
    if f_ is not None:
        c=s.cell(r,3,f_); c.number_format=YEN; c.font=SB if i==0 else BK
    c=s.cell(r,4,f'=C{B0}' if i==0 else (f'=D{r-1}' if last else f'=D{r-1}+C{r}'))
    c.number_format=YEN; c.font=SB; c.fill=OR_ if last else TF
    if last:
        c.border=TB; s.cell(r,2).fill=OR_; s.cell(r,3).fill=OR_
    s.cell(r,5,nt).font=SM
B_END=B0+3
s.cell(B_END,5,'貴社ご想定の「営業利益0円＋α」と整合します。').font=RD
n=B_END+2
bar(s,n,'【2】FY2026 着地 損益計算書',4)
for j,t in enumerate(['科目','10か月実績','8月・9月','通期 着地']):
    c=s.cell(n+1,2+j,t); c.font=HD; c.fill=HF; c.alignment=Alignment(horizontal='center')
L0=n+2
pl=[('売上高',R_ST,False),('売上原価',R_CT,False),('売上総利益',R_GP,True),
    ('販管費 計',R_SGT,False),('営業利益',R_OP,True),
    ('営業外収益',R_NI,False),('営業外費用',R_NE,False),('経常利益',R_OR,True)]
for i,(lab,rr,bold) in enumerate(pl):
    r=L0+i
    s.cell(r,2,lab).font=SB if bold else BK
    c=s.cell(r,3,f"=SUM('②月次推移表'!C{rr}:L{rr})"); c.number_format=YEN; c.font=SB if bold else BK
    c=s.cell(r,4,f"=SUM('②月次推移表'!M{rr}:N{rr})"); c.number_format=YEN; c.font=SB if bold else BK
    c=s.cell(r,5,f'=C{r}+D{r}'); c.number_format=YEN; c.font=SB; c.fill=OR_ if bold else TF
    if bold:
        for cc in range(2,6): s.cell(r,cc).border=TB
L_S,L_C,L_G,L_P,L_OP,L_NI,L_NE,L_OR=[L0+i for i in range(8)]
r=L0+8
s.cell(r,2,'　粗利率').font=BK
for col,rng in ((3,f'C{L_G}/C{L_S}'),(4,f'D{L_G}/D{L_S}'),(5,f'E{L_G}/E{L_S}')):
    c=s.cell(r,col,f'={rng}'); c.number_format=PCT; c.font=BK
r+=1
s.cell(r,2,'特別損失（和解金）').font=BK
c=s.cell(r,5,900000); c.number_format=YEN; c.font=BL; c.fill=YL
L_SL=r
r+=1
s.cell(r,2,'税引前当期純利益').font=SB
c=s.cell(r,5,f'=E{L_OR}-E{L_SL}'); c.number_format=YEN; c.font=SB; c.fill=OR_; c.border=TB
s.cell(r,2).fill=OR_; s.cell(r,2).border=TB
L_PBT=r
r+=2
s.cell(r,2,'繰越欠損金 残高（FY2026末 見込）').font=SB; s.cell(r,2).fill=TF
c=s.cell(r,3,127364212); c.number_format=YEN; c.font=BL; c.fill=YL
s.cell(r,4,'FY2025分').font=SM
c=s.cell(r,5,f'=C{r}+MAX(0,-E{L_PBT})'); c.number_format=YEN; c.font=SB; c.fill=OR_
s.cell(r,5+1,'FY2027のプランD・Eの法人税計算に使用します。').font=SM
n=r+2
bar(s,n,'【3】留意点',4)
nt=[('8-9月の売上レンジは着地に影響しない','営業利益トントンの前提では、売上が110百万円でも120百万円でも営業利益は変わりません。'
     '変わるのは売上高と粗利率の見え方だけです。⑤感応度をご参照ください。'),
    ('7月末までの戻し8,000,000円','勘定科目（販管費の戻入か売上の追加か）と計上月をご確認ください。'
     '本モデルでは販管費の減少として2026年8月に計上しています。'),
    ('商品在庫','2026/07末の商品は721,162,328円。8-9月の仕入55百万円を見込んでも、私募の売上計上分（原価255百万円）と'
     '現物出資分（簿価270百万円）が抜けるため、9月末は約187百万円まで減ります。'
     '従来のFY2027モデルの期首在庫420百万円とは大きく異なります。'),
    ('消費税','課税売上高は890百万円前後です。現物出資分3億円は売上計上しませんが、消費税法上は資産の譲渡等に'
     '該当し課税対象となる可能性があります。FY2026分の確定納付額に影響するため、顧問税理士にご確認ください。')]
for i,(k,v) in enumerate(nt):
    r=n+1+i
    s.cell(r,2,k).font=SB; s.cell(r,2).fill=OR_
    s.cell(r,3,v).font=BK
    s.merge_cells(start_row=r,start_column=3,end_row=r,end_column=5)
    s.cell(r,3).alignment=Alignment(wrap_text=True,vertical='top')
    s.row_dimensions[r].height=34

# ================= ⑤感応度 =================
z=wb.create_sheet('⑤感応度')
for col,w in zip('ABCDEFG',[3,32,18,18,18,18,3]): z.column_dimensions[col].width=w
z['B1']='8月・9月の既存事業 売上レンジ別の着地'; z['B1'].font=Font(name=F,bold=True,size=15)
z['B2']='営業利益トントンの前提のため、売上が変わっても営業利益・経常利益は変わりません。'; z['B2'].font=RD
bar(z,4,'【1】売上レンジ別',4)
for j,t in enumerate(['8-9月 既存売上','通期 売上高','通期 粗利率','通期 営業利益','通期 経常利益']):
    c=z.cell(5,2+j,t); c.font=HD; c.fill=HF; c.alignment=Alignment(horizontal='center')
base_s=f"SUM('②月次推移表'!C{R_ST}:L{R_ST})+{PD(UT)}+{P(U0)}+{P(U0+1)}+{P(U0+2)}+{P(U0+3)}-{PD(UT)}"
Z0=6
for i,v in enumerate([110000000,115000000,120000000]):
    r=Z0+i
    c=z.cell(r,2,v); c.number_format=YEN; c.font=BL; c.fill=YL
    c=z.cell(r,3,f"=SUM('②月次推移表'!C{R_ST}:L{R_ST})+B{r}+{P(U0)}+{P(U0+1)}+{P(U0+2)}+{P(U0+3)}")
    c.number_format=YEN; c.font=SB
    c=z.cell(r,4,f"=(SUM('②月次推移表'!C{R_GP}:L{R_GP})+{P(G8)}+{P(G9)}+{PD(UT)})/C{r}")
    c.number_format=PCT; c.font=BK
    c=z.cell(r,5,f"='①着地予想'!E{L_OP}"); c.number_format=YEN; c.font=SB
    c=z.cell(r,6,f"='①着地予想'!E{L_OR}"); c.number_format=YEN; c.font=SB
    if v==115000000:
        for cc in range(2,7): z.cell(r,cc).fill=OR_
z.cell(Z0+4,2,'※中央値115百万円が②月次推移表の設定値です（8月・9月に折半）。').font=SM
z.cell(Z0+5,2,'※営業利益・経常利益が一定なのは、粗利＝販管費とする「トントン」前提の必然です。売上が伸びれば粗利率が下がる形で調整されます。').font=RD
z.cell(Z0+6,2,'※8-9月に営業利益を出せる見込みがあれば、その分だけ着地は上振れします。').font=RD

# ================= ④添付ファイル検証 =================
v=wb.create_sheet('④添付ファイル検証')
for col,w in zip('ABCDEF',[3,44,20,20,58,2]): v.column_dimensions[col].width=w
v['B1']='添付ファイル（損益計算書 月次推移 2025/10-2026/07）の検証'; v['B1'].font=Font(name=F,bold=True,size=15)
v['B2']='合計欄と月次の積み上げ、各月の内部整合、棚卸高の連続性を確認しました。'; v['B2'].font=SM
bar(v,4,'【1】合計欄と月次積み上げの照合',4)
for j,t in enumerate(['項目','月次の積み上げ','報告書の合計欄','判定']):
    c=v.cell(5,2+j,t); c.font=HD; c.fill=HF; c.alignment=Alignment(horizontal='center')
items=[('売上高',SALES,'売上高'),('売上原価',COGS,'売上原価'),('売上総利益',GP,'売上総利益'),
       ('販管費',SGA,'販管費'),('営業利益',OP,'営業利益'),('営業外収益',NOI,'営業外収益'),
       ('営業外費用',NOE,'営業外費用'),('経常利益',ORD,'経常利益'),
       ('商品仕入高',BUY,'商品仕入高'),('賄費戻り金',MKN,'賄費戻り金')]
for i,(lab,arr,key) in enumerate(items):
    r=6+i
    v.cell(r,2,lab).font=BK
    c=v.cell(r,3,sum(arr)); c.number_format=YEN; c.font=BK
    c=v.cell(r,4,T[key]); c.number_format=YEN; c.font=BK
    c=v.cell(r,5,f'=IF(C{r}=D{r},"一致","差異 "&TEXT(C{r}-D{r},"#,##0"))'); c.font=SB
n=6+len(items)+1
bar(v,n,'【2】各月の内部整合（全10か月）',4)
ck=[('売上総利益 ＝ 売上高 − 売上原価','全10か月で一致'),
    ('営業利益 ＝ 売上総利益 − 販管費','全10か月で一致'),
    ('経常利益 ＝ 営業利益 ＋ 営業外収益 − 営業外費用','全10か月で一致'),
    ('売上原価 ＝ 期首棚卸 ＋ 仕入 ＋ 賄費戻り金 − 期末棚卸','全10か月で一致'),
    ('売上高 ＝ 内訳7科目の合計','全10か月で一致'),
    ('税引前 ＝ 経常利益 − 特別損失900,000（和解金）','一致。特別利益は0')]
for i,(k,r2) in enumerate(ck):
    r=n+1+i
    v.cell(r,2,k).font=BK
    v.cell(r,3,r2).font=SB
n=n+len(ck)+2
bar(v,n,'【3】確認が必要な点　※棚卸高の連続性',4)
for j,t in enumerate(['','前月末 棚卸高','当月首 棚卸高','差異']):
    c=v.cell(n+1,2+j,t); c.font=HD; c.fill=HF; c.alignment=Alignment(horizontal='center')
gaps=[(i,BEG[i]-END[i-1]) for i in range(1,10) if BEG[i]!=END[i-1]]
for i,(idx,d) in enumerate(gaps):
    r=n+2+i
    v.cell(r,2,f'{MO[idx-1]}末 → {MO[idx]}首').font=BK
    c=v.cell(r,3,END[idx-1]); c.number_format=YEN; c.font=BK
    c=v.cell(r,4,BEG[idx]); c.number_format=YEN; c.font=BK
    c=v.cell(r,5,d); c.number_format=YEN; c.font=SB; c.fill=YL
r=n+2+len(gaps)
v.cell(r,2,'純額').font=SB; v.cell(r,2).border=TB
c=v.cell(r,5,sum(d for _,d in gaps)); c.number_format=YEN; c.font=SB; c.fill=OR_; c.border=TB
r+=2
for t in ['※月をまたぐ棚卸高が2か所で一致していません（洗い替えと思われます）。各月の売上原価は月内で整合しているため、',
          '　通期の売上原価252,912,579円は月次の積み上げとしては正しい数値です。',
          '※一方、期首501,059,742円・期末721,162,328円と仕入469,514,545円から逆算すると売上原価は248,994,919円となり、',
          '　報告値との差3,917,660円は上記2件の純額と完全に一致します。通期の在庫増減を使って検算する際はご注意ください。',
          '※商品在庫は10か月で501.1百万円→721.2百万円と220.1百万円増えています。FY2027の期首在庫に直結する論点です。']:
    v.cell(r,2,t).font=RD; r+=1
r+=1
bar(v,r,'【4】転記時に判明した点',4)
r+=1
for t in ['※2026/04の売上高を32,501,576円と読むと内訳合計および売上総利益と6,000円ずれます。32,507,576円が正しい数値です。',
          '※和解金900,000円は特別損失であり、特別利益は0です（税引前当期純利益▲122,855,001円と整合）。']:
    v.cell(r,2,t).font=RD; r+=1

for ws in wb: ws.sheet_view.showGridLines=False
wb.save('WineBank_FY2026_着地予想.xlsx'); print('saved')
